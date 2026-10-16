"""Evidence-only tables for finite-controller, pure-compute MoE DSE.

Window metrics are paired before aggregation.  Seven batch sizes have equal
weight regardless of their different numbers of archived windows.
"""
from __future__ import annotations

import argparse
import csv
import json
import math
import random
from collections import defaultdict
from pathlib import Path

BATCHES = (2, 4, 8, 16, 64, 96, 128)
BUDGET = 65536
LABELS = {
    "original": "PLENA 原默认几何",
    "compact_single": "指定紧凑单核",
    "compact_homo_greedy": "指定同构 / 入口 EFT 私有队列",
    "compact_homo_mdpu": "指定同构 / M-DPU",
    "hetero_mdpu": "Dense+Stream / M-DPU",
    "hetero_hard": "Dense+Stream / 硬阈值重搜",
    "fair_single": "同能力最优单核",
    "fair_homo_greedy": "同能力同构 / 入口 EFT 私有队列",
    "fair_homo_mdpu": "同能力同构 / M-DPU",
    "fair_native_single": "BLEN≥2 单核对照",
    "fair_native_homo": "BLEN≥2 同构对照",
}


def gm(values):
    values = tuple(float(value) for value in values)
    if not values or any(not math.isfinite(x) or x < 0 for x in values):
        raise ValueError("Geometric mean requires finite nonnegative values")
    return 0.0 if 0 in values else math.exp(math.fsum(map(math.log, values)) / len(values))


def mean(values):
    values = tuple(values)
    return math.fsum(values) / len(values) if values else None


def balanced_gm(rows, field="cycles"):
    batches = {batch: [float(r[field]) for r in rows if int(r["batch"]) == batch] for batch in BATCHES}
    if any(not values for values in batches.values()):
        raise ValueError("Expected all seven batches")
    return gm(gm(values) for values in batches.values())


def balanced_mean(rows, field):
    return mean(mean(float(r[field]) for r in rows if int(r["batch"]) == b) for b in BATCHES)


def paired_speedup(reference, candidate):
    ref = {str(r["window_id"]): r for r in reference}
    cand = {str(r["window_id"]): r for r in candidate}
    if len(ref) != len(reference) or len(cand) != len(candidate) or ref.keys() != cand.keys():
        raise ValueError("Speedups require one-to-one identical window IDs")
    ratios = []
    for identifier in sorted(ref):
        a, b = ref[identifier], cand[identifier]
        if int(a["batch"]) != int(b["batch"]):
            raise ValueError("Batch mismatch")
        ratios.append({"window_id": identifier, "batch": int(a["batch"]),
                       "ratio": float(a["cycles"]) / float(b["cycles"])})
    return ratios


def bootstrap_speedup(reference, candidate, *, repeats=200, seed=20261013):
    """Paired stratified bootstrap; descriptive uncertainty on archived windows."""
    ratios = paired_speedup(reference, candidate)
    groups = [[r["ratio"] for r in ratios if r["batch"] == b] for b in BATCHES]
    if any(not group for group in groups):
        raise ValueError("Bootstrap requires all seven batches")
    rng = random.Random(seed)
    samples = sorted(gm(gm(rng.choices(group, k=len(group))) for group in groups) for _ in range(repeats))
    def quantile(q):
        p = (len(samples) - 1) * q
        lower, upper = math.floor(p), math.ceil(p)
        return samples[lower] + (p - lower) * (samples[upper] - samples[lower])
    return {"speedup": gm(gm(group) for group in groups), "ci95_low": quantile(.025),
            "ci95_high": quantile(.975), "bootstrap_repeats": repeats, "seed": seed}


def write_csv(path, rows):
    if not rows:
        return
    keys = list(dict.fromkeys(key for row in rows for key in row))
    with Path(path).open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=keys)
        writer.writeheader()
        writer.writerows(rows)


def table(rows):
    if not rows:
        return "无结果。"
    keys = tuple(rows[0])
    lines = ["| " + " | ".join(keys) + " |", "|" + "|".join("---" for _ in keys) + "|"]
    for row in rows:
        items = []
        for key in keys:
            value = row.get(key)
            if value is None:
                value = "—"
            elif isinstance(value, float):
                value = f"{value:,.6f}" if "ms" in key else f"{value:,.4f}"
            items.append(str(value).replace("|", "\\|"))
        lines.append("| " + " | ".join(items) + " |")
    return "\n".join(lines)


def physical(cores):
    return " + ".join(f"{c['pm']}×{c['pn']}×{c['groups']}" for c in cores)


def write_workbook(path, sheets):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    book = Workbook()
    book.remove(book.active)
    for name, rows in sheets.items():
        sheet = book.create_sheet(name[:31])
        if not rows:
            continue
        keys = list(rows[0])
        sheet.append(keys)
        for row in rows:
            sheet.append([json.dumps(row.get(k), ensure_ascii=False) if isinstance(row.get(k), (list, dict)) else row.get(k) for k in keys])
        sheet.freeze_panes = "C2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17365D")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[1].height = 34
        for index, key in enumerate(keys, 1):
            sheet.column_dimensions[get_column_letter(index)].width = min(55, max(16, len(str(key)) + 4))
        for cells in sheet.iter_rows(min_row=2):
            for cell in cells:
                if isinstance(cell.value, float):
                    cell.number_format = "0.000000" if "ms" in str(keys[cell.column-1]) else "0.0000"
    book.save(path)


def flatten_record(record):
    plan = record["result"]
    cores = record["cores"]
    cycles = int(plan["cycles"])
    useful, issued = int(plan["useful_MACs"]), int(plan["issued_macs"])
    busy, ends = list(plan["core_cycles"]), list(plan["core_finish_cycles"])
    if sum(int(c["macs"]) for c in cores) != BUDGET:
        raise ValueError("Unequal multiplier budget")
    if len(busy) != len(cores) or max(ends) != cycles or not 0 < useful <= issued <= BUDGET * cycles:
        raise ValueError("Invalid compute accounting")
    boundary = plan["boundary_stats"]
    row = {
        "design": record["design"], "variant": record.get("variant", "main"),
        "split": record.get("split", "heldout"), "batch": int(record["batch"]),
        "window_id": str(record["window_id"]), "policy": record.get("policy", plan["policy"]),
        "wait_best": bool(plan.get("wait_best", False)),
        "physical_PM_PN_G": physical(cores), "MAC_allocation": "+".join(str(c["macs"]) for c in cores),
        "native_frames": "+".join(f"BLEN={c['b']}, MLEN={c['r']}" for c in cores),
        "threshold": int(plan["threshold"]), "cycles": cycles, "compute_ms": cycles / 1e6,
        "array_max_busy_ms": max(busy) / 1e6,
        "useful_GFLOP_s": 2 * useful / cycles,
        "wall_MAC_utilization_pct": 100 * useful / (BUDGET * cycles),
        "padding_of_issued_pct": 100 * (issued - useful) / issued,
        "controller_and_idle_gap_cycles": cycles - max(busy),
        "core0_busy_ms": busy[0] / 1e6,
        "core1_busy_ms": busy[1] / 1e6 if len(cores) == 2 else None,
        "core0_finish_ms": ends[0] / 1e6,
        "core1_finish_ms": ends[1] / 1e6 if len(cores) == 2 else None,
        "core0_idle_pct": 100 * (cycles - busy[0]) / cycles,
        "core1_idle_pct": 100 * (cycles - busy[1]) / cycles if len(cores) == 2 else None,
        "core0_useful_FLOPs_pct": 100 * boundary["core_useful_MACs"][0] / useful,
        "core1_useful_FLOPs_pct": 100 * boundary["core_useful_MACs"][1] / useful if len(cores) == 2 else None,
        "imbalance_ratio": boundary["imbalance_ratio"],
        "zero_load_core_window": int(len(cores) == 2 and min(busy) == 0),
        "steal_count": plan["steal_count"],
        "affinity_deviation_count": plan["affinity_deviation_count"],
        "context_service_cycles": plan["context_service_cycles"],
        "context_overlap_cycles": plan["context_overlap_cycles"],
        "context_exposed_cycles": plan["context_exposed_cycles"],
        "frontend_backpressure_cycles": plan["frontend_backpressure_cycles"],
        "prediction_mae_cycles": plan["prediction_mae_cycles"],
        "prediction_worst_underestimate_cycles": plan["prediction_worst_underestimate_cycles"],
        "queue0_peak": plan["queue_max_occupancy"][0],
        "queue1_peak": plan["queue_max_occupancy"][1] if len(cores) == 2 else None,
        "threshold_realizable": int(bool(boundary["realizable_thresholds"])),
        "shared_on_core0": int(boundary["shared_core"] == 0),
        "realizable_thresholds": json.dumps(boundary["realizable_thresholds"]),
        "unique_weight_MiB": plan["unique_weight_bytes"] / 1024**2,
        "unique_weight_average_TB_s": plan["unique_weight_bytes"] / cycles / 1000,
        "unique_fetch_ms_256GB_s": plan["unique_weight_bytes"] / 256e6,
        "useful_macs": useful, "issued_macs": issued,
    }
    jobs_all = plan.get("jobs")
    if jobs_all is None:
        return row
    job_ids = [j["expert_id"] for j in jobs_all]
    if len(job_ids) != len(set(job_ids)) or sum(j["useful_macs"] for j in jobs_all) != useful:
        raise ValueError("Task conservation failed")
    for cid, end in enumerate(ends):
        jobs = sorted((j for j in jobs_all if j["core_id"] == cid), key=lambda j: j["start"])
        if jobs:
            if jobs[-1]["end"] != end or any(a["end"] > b["start"] for a, b in zip(jobs, jobs[1:])):
                raise ValueError("Overlapping jobs or wrong final completion")
            if any(j["start"] < j["context_ready"] for j in jobs):
                raise ValueError("Compute started before its context was ready")
    return row


GM_FIELDS = {"cycles", "compute_ms", "array_max_busy_ms", "useful_GFLOP_s", "wall_MAC_utilization_pct",
             "unique_weight_MiB", "unique_weight_average_TB_s", "unique_fetch_ms_256GB_s"}
MAX_FIELDS = {"prediction_worst_underestimate_cycles", "queue0_peak", "queue1_peak"}
IDENTIFIERS = {"threshold", "batch", "window_count"}


def aggregate_group(rows, batch):
    result = dict(rows[0])
    result.update(batch=batch, window_id="", window_count=len(rows))
    for key in rows[0]:
        values = [r[key] for r in rows if r.get(key) is not None]
        if key in IDENTIFIERS or not values or not all(isinstance(x, (int, float)) for x in values):
            continue
        result[key] = gm(values) if key in GM_FIELDS else max(values) if key in MAX_FIELDS else mean(values)
    result["threshold_realizable_windows"] = sum(r["threshold_realizable"] for r in rows)
    result["zero_load_core_windows"] = sum(r["zero_load_core_window"] for r in rows)
    result["realizable_thresholds"] = "按窗口见 window_metrics.csv"
    return result


def summaries(rows):
    batches = [aggregate_group([r for r in rows if r["batch"] == b], b) for b in BATCHES]
    overall = aggregate_group(batches, "7-Batch GM")
    overall["window_count"] = len(rows)
    overall["threshold_realizable_windows"] = sum(r["threshold_realizable"] for r in rows)
    overall["zero_load_core_windows"] = sum(r["zero_load_core_window"] for r in rows)
    return batches + [overall]


def presentation_main(rows, single, homo):
    result = []
    references = {(r["design"], r["batch"]): r for r in rows}
    for row in rows:
        baseline_s = references.get((single, row["batch"]))
        baseline_h = references.get((homo, row["batch"]))
        result.append({
            "Batch": f"B{row['batch']}" if isinstance(row["batch"], int) else row["batch"],
            "组织": LABELS.get(row["design"], row["design"]),
            "物理配置 PM×PN×G": row["physical_PM_PN_G"], "MAC 分配": row["MAC_allocation"],
            "派工": row["policy"], "完成时间 ms": row["compute_ms"],
            "有效 GFLOP/s": row["useful_GFLOP_s"], "整片 MAC 利用率 %": row["wall_MAC_utilization_pct"],
            "相对单核加速比": baseline_s["compute_ms"] / row["compute_ms"] if baseline_s else None,
            "相对同构加速比": baseline_h["compute_ms"] / row["compute_ms"] if baseline_h else None,
        })
    return result


def presentation_boundary(rows):
    return [{
        "Batch": f"B{r['batch']}" if isinstance(r["batch"], int) else r["batch"],
        "基准阈值 T": r["threshold"], "可由单阈值解释的窗口": f"{int(r['threshold_realizable_windows'])}/{r['window_count']}",
        "Dense 有效 FLOPs %": r["core0_useful_FLOPs_pct"], "Stream 有效 FLOPs %": r["core1_useful_FLOPs_pct"],
        "Dense 忙碌 ms": r["core0_busy_ms"], "Stream 忙碌 ms": r["core1_busy_ms"],
        "失配比 max/min": f"∞（{int(r['zero_load_core_windows'])}/{r['window_count']} 窗口有核零负载）" if r["zero_load_core_windows"] else r["imbalance_ratio"], "Dense 空闲 %": r["core0_idle_pct"],
        "Stream 空闲 %": r["core1_idle_pct"], "填充浪费占发射槽 %": r["padding_of_issued_pct"],
        "偏离基准归属任务数": r["affinity_deviation_count"], "队头旁路次数": r["steal_count"],
        "Shared 最终在 Dense 比例 %": 100 * r["shared_on_core0"],
    } for r in rows]


def read_records(path):
    import gzip
    opener = gzip.open if str(path).endswith(".gz") else open
    with opener(path, "rt", encoding="utf-8") as handle:
        return [json.loads(line) for line in handle if line.strip()]


def revision_comparison(outdir):
    """Compare preserved v1 to corrected admission, clearly separating geometry."""
    outdir = Path(outdir)
    old_path, new_path = outdir.parent / "batch_summary.csv", outdir / "batch_summary.csv"
    if not old_path.exists() or not new_path.exists():
        return
    def read_overall(path):
        with path.open(newline="", encoding="utf-8") as handle:
            return {(r["design"], r["variant"]): r for r in csv.DictReader(handle) if r["batch"] == "7-Batch GM"}
    old, new = read_overall(old_path), read_overall(new_path)
    rows, matched = [], []
    for design, variant in new:
        if variant != "main" or (design, variant) not in old:
            continue
        a, b = old[(design, variant)], new[(design, variant)]
        rows.append({"组织": LABELS.get(design, design), "旧版冻结几何": a["physical_PM_PN_G"],
                     "修订版冻结几何": b["physical_PM_PN_G"], "旧版 GM ms": float(a["compute_ms"]),
                     "修订版 GM ms": float(b["compute_ms"]),
                     "旧版时间/修订版时间": float(a["compute_ms"]) / float(b["compute_ms"])})
        if (design, "unguarded_mdpu") in new:
            unguarded = new[(design, "unguarded_mdpu")]
            noprefetch = new.get((design, "no_microop_prefetch"))
            matched.append({"组织": LABELS.get(design, design), "固定修订版几何": b["physical_PM_PN_G"],
                            "旧准入规则 GM ms": float(unguarded["compute_ms"]),
                            "允许等待 GM ms": float(b["compute_ms"]),
                            "准入修复加速比": float(unguarded["compute_ms"]) / float(b["compute_ms"]),
                            "允许等待但关预配置 GM ms": float(noprefetch["compute_ms"]) if noprefetch else None,
                            "预配置加速比": float(noprefetch["compute_ms"]) / float(b["compute_ms"]) if noprefetch else None})
    write_csv(outdir / "revision_geometry_comparison.csv", rows)
    write_csv(outdir / "revision_matched_hardware.csv", matched)
    write_workbook(outdir / "准入修订_版本与固定硬件对照.xlsx", {"两版重新选型": rows, "固定修订版硬件": matched})
    text = "\n".join(["# 准入修订：保留旧结果，区分机制与重新选型", "",
                      "上一版在最快核队列满时立即溢出到较慢核，且 shadow 预配置过早锁定任务。修订版允许等待最快队列；依旧可以旁路未发射的 FIFO 队头，不迁移已执行任务。", "",
                      "两版都只在 18 个开发窗口中选几何和阈值，再运行 135 个留出窗口。但修订版的机制来自旧留出集诊断，因此不是新的盲测。", "",
                      "## 两版各自重新选型后的结果", "", table(rows), "",
                      "上表同时可能改变几何与准入机制，不能把其全部加速归给准入。", "",
                      "## 固定修订版几何，只改准入 / 预配置", "", table(matched), "",
                      "这一表才用于单变量机制归因。所有策略同为理想供数，预配置只提前准备任务上下文。原始回退案例保留在上级 PREFETCH_DIAGNOSIS.md。", ""])
    (outdir / "REVISION_COMPARISON.md").write_text(text, encoding="utf-8")
    (outdir.parent / "CURRENT_RESULTS.md").write_text("\n".join([
        "# 当前完整结果入口", "", "当前主结果是允许等待最佳队列的准入修订版；旧版冻结数值保留。", "",
        "- [修订版完整分析与全 Batch 表](guarded/SUMMARY.md)",
        "- [修订版 Excel 完整数据表](guarded/M_DPU_完整数据表.xlsx)",
        "- [两版对照与固定硬件单变量归因](guarded/REVISION_COMPARISON.md)",
        "- [旧版完整报告](SUMMARY.md)",
        "- [旧版预配置回退诊断](PREFETCH_DIAGNOSIS.md)", "",
        "修订版遵守开发集选型；该机制已参考旧留出集问题定位，应作为修复回归 / 探索评估解读。", ""]), encoding="utf-8")
    return {"design_comparisons": len(rows), "matched_comparisons": len(matched)}


def generate(input_path, outdir, frozen_path=None):
    outdir = Path(outdir)
    outdir.mkdir(parents=True, exist_ok=True)
    records = read_records(input_path)
    frozen_path = Path(frozen_path or outdir / "FROZEN_SELECTION.json")
    frozen = json.loads(frozen_path.read_text()) if frozen_path.exists() else {}
    guarded = bool(frozen.get("contract", {}).get("wait_best_queue_admission"))
    # Runner retains full results with a flat row, and stores geometry once.
    if records and "row" in records[0]:
        records = [{**r["row"], "cores": frozen["families"][r["row"]["design"]]["cores"],
                    "result": r["result"]} for r in records]
    rows = [flatten_record(r) for r in records if r.get("split", "heldout") == "heldout"]
    groups = defaultdict(list)
    for row in rows:
        groups[(row["design"], row["variant"])].append(row)
    summary = []
    for key, values in groups.items():
        expected = {b: 27 if b <= 16 else 9 for b in BATCHES}
        actual = {b: sum(r["batch"] == b for r in values) for b in BATCHES}
        if actual != expected or len({r["window_id"] for r in values}) != 135:
            raise ValueError(f"Incomplete or duplicate heldout grid: {key}: {actual}")
        summary.extend(summaries(values))
    main = [r for r in summary if r["variant"] == "main"]
    compact_names = ("original", "compact_single", "compact_homo_greedy", "compact_homo_mdpu", "hetero_mdpu")
    fair_names = ("fair_single", "fair_homo_greedy", "fair_homo_mdpu", "hetero_mdpu")
    ordered = lambda names: sorted([r for r in main if r["design"] in names], key=lambda r: (BATCHES.index(r["batch"]) if r["batch"] in BATCHES else 7, names.index(r["design"])))
    compact = presentation_main(ordered(compact_names), "compact_single", "compact_homo_mdpu")
    fair = presentation_main(ordered(fair_names), "fair_single", "fair_homo_mdpu")
    boundary = presentation_boundary([r for r in main if r["design"] == "hetero_mdpu"])
    comparisons = []
    candidate = groups.get(("hetero_mdpu", "main"))
    bootstrap_path = outdir / "paired_bootstrap.csv"
    if bootstrap_path.exists():
        with bootstrap_path.open(newline="") as handle:
            comparisons = [{"candidate": r["candidate"], "reference": r["baseline"],
                            "speedup": float(r["speedup"]), "ci95_low": float(r["bootstrap_95_low"]),
                            "ci95_high": float(r["bootstrap_95_high"]), "bootstrap_repeats": int(r["draws"])}
                           for r in csv.DictReader(handle)]
    elif candidate:
        for reference in ("original", "compact_single", "compact_homo_greedy", "compact_homo_mdpu", "fair_single", "fair_homo_greedy", "fair_homo_mdpu", "fair_native_single", "fair_native_homo"):
            if (reference, "main") in groups:
                comparisons.append({"candidate": "hetero_mdpu", "reference": reference,
                                    **bootstrap_speedup(groups[(reference, "main")], candidate)})
    configs = [{"组织": LABELS.get(r["design"], r["design"]), "物理配置 PM×PN×G": r["physical_PM_PN_G"],
                "原生块参数": r["native_frames"], "MAC 分配": r["MAC_allocation"], "派工": r["policy"],
                "开发集冻结阈值": r["threshold"], "留出 7-Batch GM ms": r["compute_ms"]}
               for r in main if r["batch"] == "7-Batch GM"]
    ablations = [{"组织": LABELS.get(r["design"], r["design"]), "变体": r["variant"], "派工": r["policy"],
                  "7-Batch GM ms": r["compute_ms"], "核阵列忙碌最大值 GM ms": r["array_max_busy_ms"],
                  "平均上下文准备周期": r["context_service_cycles"], "平均隐藏准备周期": r["context_overlap_cycles"],
                  "平均暴露准备周期": r["context_exposed_cycles"], "平均队列背压周期": r["frontend_backpressure_cycles"],
                  "预测 MAE 周期": r["prediction_mae_cycles"], "最差低估周期": r["prediction_worst_underestimate_cycles"],
                  "平均旁路次数": r["steal_count"]}
                 for r in summary if r["batch"] == "7-Batch GM"]
    supply = [{"组织": LABELS.get(r["design"], r["design"]), "7-Batch GM ms": r["compute_ms"],
               "唯一权重容量 GM MiB": r["unique_weight_MiB"], "唯一权重平均带宽 GM TB/s": r["unique_weight_average_TB_s"],
               "唯一权重服务时间 @256 GB/s ms": r["unique_fetch_ms_256GB_s"]}
              for r in main if r["batch"] == "7-Batch GM"]
    resources_path = outdir / "rtl" / "hardware_resources.json"
    rtl_resources = json.loads(resources_path.read_text()) if resources_path.exists() else {}
    ledger = []
    if rtl_resources:
        for label, key in (("2×8 项任务 FIFO", "fifo_storage_bits"), ("两核 shadow context", "shadow_storage_bits"),
                           ("FIFO 指针", "pointer_bits"), ("FIFO 计数", "count_bits"),
                           ("上下文准备计数", "prepare_counter_bits"), ("shadow valid", "shadow_valid_bits")):
            ledger.append({"模块": label, "声明状态位数": rtl_resources[key], "字节等价": rtl_resources[key] / 8})
        if "queued_service_scoreboard_bits" in rtl_resources:
            value = rtl_resources["queued_service_scoreboard_bits"]
            ledger.append({"模块": "队列剩余工作量计分寄存器", "声明状态位数": value, "字节等价": value / 8})
        ledger.append({"模块": "上述控制 RTL 合计", "声明状态位数": rtl_resources["declared_sequential_state_bits"],
                       "字节等价": rtl_resources["declared_sequential_state_bytes"]})
        external = sum(rtl_resources["excluded_integration_state"].values())
        ledger.append({"模块": "另列集成状态 / 系数 / 模式", "声明状态位数": external, "字节等价": external / 8})
    for filename, data in (("window_metrics.csv", rows), ("batch_summary.csv", summary), ("compact_table.csv", compact),
                           ("fair_table.csv", fair), ("mdpu_boundary.csv", boundary), ("ablations.csv", ablations),
                           ("paired_speedup.csv", comparisons), ("supply_requirements.csv", supply)):
        write_csv(outdir / filename, data)
    write_workbook(outdir / "M_DPU_完整数据表.xlsx", {"配置与口径": configs, "受限家族全Batch": compact,
                   "相同能力公平对照": fair, "M-DPU分流": boundary, "调度器消融": ablations,
                   "配对置信区间": comparisons, "供数需求反推": supply, "控制RTL状态账本": ledger, "逐窗口结果": rows})
    brief = [{"比较基线": LABELS.get(r["reference"], r["reference"]), "异构加速比": r["speedup"],
              "95% CI 下界": r["ci95_low"], "95% CI 上界": r["ci95_high"]} for r in comparisons]
    overall = {(r["design"], r["variant"]): r for r in summary if r["batch"] == "7-Batch GM"}
    by_batch = {(r["design"], r["batch"]): r for r in main}
    conclusions = []
    if all((d, "main") in overall for d in ("hetero_mdpu", "compact_single", "compact_homo_mdpu", "fair_single", "fair_homo_mdpu")):
        cand_ms = overall[("hetero_mdpu", "main")]["compute_ms"]
        ratio = lambda d: overall[(d, "main")]["compute_ms"] / cand_ms
        ladder = sum(by_batch[("hetero_mdpu", b)]["compute_ms"] < by_batch[("compact_homo_mdpu", b)]["compute_ms"] < by_batch[("compact_single", b)]["compute_ms"] for b in BATCHES)
        conclusions = [
            f"指定紧凑家族：异构 GM 为 {cand_ms:.6f} ms，相对紧凑单核 {ratio('compact_single'):.4f}×，相对相同 M-DPU 的紧凑同构 {ratio('compact_homo_mdpu'):.4f}×；严格‘异构快于同构、同构快于单核’在 {ladder}/7 个 Batch 成立。",
            f"同能力控制：异构相对允许相同 PM=1/2 能力的最优单核为 {ratio('fair_single'):.4f}×，相对同一 M-DPU 的最优同构为 {ratio('fair_homo_mdpu'):.4f}×。这组对照决定能否把受限家族的结果推广为普遍架构优势。",
        ]
        cores = frozen.get("families", {}).get("hetero_mdpu", {}).get("cores", [])
        if len(cores) == 2:
            divisor = math.gcd(cores[0]["macs"], cores[1]["macs"])
            conclusions.append(f"最终 Dense:Stream 乘法器为 {cores[0]['macs']:,}:{cores[1]['macs']:,}，即 {cores[0]['macs']//divisor}:{cores[1]['macs']//divisor}。Dense/Stream 是几何形状命名，不能据此默认 Dense 的算力更大或 Shared 必须由它执行。")
        candidate_overall = overall[("hetero_mdpu", "main")]
        if candidate_overall["core0_useful_FLOPs_pct"] < 5:
            conclusions.append(f"实际分工：Dense 仅承担 {candidate_overall['core0_useful_FLOPs_pct']:.4f}% 有效 FLOPs，平均空闲 {candidate_overall['core0_idle_pct']:.4f}%；Shared 在 Dense 的比例为 {100*candidate_overall['shared_on_core0']:.2f}%。该点实际上主要依靠 Stream 一核执行，不能作为 Dense/Stream 高效互补的证据。")
        for variant, title in (("same_hw_greedy", "相同异构几何上相对入口 EFT 私有队列派工"), ("unguarded_mdpu", "相同异构几何上相对旧准入规则"), ("no_microop_prefetch", "相同异构几何上相对关闭微操作预配置"), ("no_bypass", "相同异构几何上相对关闭旁路")):
            if ("hetero_mdpu", variant) in overall:
                value = overall[("hetero_mdpu", variant)]["compute_ms"] / cand_ms
                conclusions.append(f"{title}：完整 M-DPU 为 {value:.6f}×。")
                if variant == "no_microop_prefetch" and value < 1:
                    conclusions.append("此处关闭预配置反而更快，说明当前 shadow context 的提前归属可能损害后续负载平衡；不能将‘隐藏准备周期’等同于‘整层一定加速’，具体归因还须检查逐任务时序。")
    lines = ["# M-DPU：Shared + MoE 纯计算 DSE 与有限控制器评估", "",
             "本报告数值来自当前逐窗口运行。1 GHz，65,536 个乘法器，BF16/FP32、OS、理想数据供给；完成时间包含阵列执行与有限队列/上下文控制时间。它是未用完整 BF16 PLENA RTL 校准的解析阵列模型加离散事件控制模型，不是有限 HBM 的端到端测量。", "",
             "18 个开发窗口选型后冻结，在 135 个留出窗口评估。每 Batch 先取窗口 GM，再取 7 Batch 等权 GM；过往已用过这些留出窗口，因此它们不是新的盲测集。", "",
             "## 结果先看", "", "\n\n".join(conclusions), "",
             "## 冻结配置", "", table(configs), "",
             "## 指定紧凑家族对照", "",
             "单核和同构的紧凑形状限制是本次指定实验条件。异构 Stream 允许更窄的 PM，因此这张表同时改变了可用几何，不能将所有收益都解释为调度器贡献。加速比的同构参照使用同一 M-DPU。greedy 为入口 EFT 后进入私有有限队列，不是无约束的全局共享任务池。", "", table(compact), "",
             "## 相同能力对照", "",
             "单核和同构也允许候选使用的 PM=1/2 扩展。这里回答允许相同硬件能力后，异构是否仍有优势。", "", table(fair), "",
             "## 配对比较", "", table(brief), "",
             f"置信区间按 Batch 内窗口配对重采样 {comparisons[0]['bootstrap_repeats'] if comparisons else 0} 次。它反映归档窗口差异，不能代替阵列时序模型误差或芯片校准。加速比 >1 表示候选更快。", "",
             "## M-DPU 实际分流", "", table(boundary), "",
             "阈值是亲和性偏好；队列预计完成时刻和空间会改变最终归属。部分窗口无法由一个严格 Me 阈值解释，已明确计数，不能把它们伪写成某个 T_dyn。核忙碌时间、空闲率、任务数量取窗口算术平均，再做 Batch 等权平均；主时间仍为 GM。", "",
             "## 固定硬件调度与预配置消融", "", table(ablations), "",
             "只有同一组织相同冻结尺寸的变体能用于归因。理想数据环境中的预取仅指下一个专家上下文预配置；没有 HBM 延迟可隐藏。模型内部的时长预测误差不是实芯片预测器精度。队列背压表示新描述符尚未准入，期间核心可能仍满负荷计算，因此背压周期不能直接加到整层延迟。", "",
             "## 硬件控制结构与资源", "", "```text\nRouter 完成后的专家描述符（Shared 优先，再按 expert ID）\n             ↓\n编译器系数 × ceil(Me / PM) + 固定尾延迟\n             ↓\n比较两核队列预计完成时刻 + 阈值亲和性 + FIFO 空间\n             ↓\n       DTQ 8项               STQ 8项\n             ↘  未发射队头旁路  ↙\n      Core0 shadow         Core1 shadow\n       4词预配置              4词预配置\n             ↓                    ↓\n       Dense 阵列             Stream 阵列\n             └── 实际完成 / 剩余时长反馈 ──┘\n```", "",
             "跨轨旁路只处理尚未发射的队头任务，不迁移正在执行的专家。每个核一个 current 和一个 shadow；上下文未就绪不能启动。相同控制能力也对单核与同构开放。本轮是编译器形状系数加实时排队状态的估计器，不是经过历史训练的 BTB，也不预测已经由 Router 给出的 Expert ID。通用 RTL 中系数乘法是 16×32 位乘法；只有把系数固化后才可能综合为常数移位加法，不能预先按‘只有几个移位器’估算面积。", "", table(ledger), "",
             "## 供数条件反推", "", table(supply), "",
             "唯一权重带宽只给必要的平均供数下界，未包含操作数重复读取、激活和瞬时端口峰值，也不是可实现 SRAM 容量或供数性能承诺。", "",
             "## 物理与文献边界", "",
             "PLENA v3 Table III 允许 BLEN 从 2 开始；系统评估为 BLEN32/MLEN2048。原论文 PE 使用 MX 输入与整数累加，本轮 BF16/FP32 是拓扑研究扩展。原论文已有 OS、M_SUM 与独立预取。不能把 BLEN≥16 写成论文强制约束，也不能把新增 BF16 标量扩展仅授予异构来证明普遍优势。[PLENA 原文](https://arxiv.org/html/2509.09505v3)", "",
             "双队列、预取和资源分区本身不是首次提出：[Gemmini](https://github.com/ucb-bar/gemmini) 已采用解耦的访存/执行队列；[Herald](https://research.nvidia.com/publication/2021-02_heterogeneous-dataflow-accelerators-multi-dnn-workloads) 已联合优化异构分区与调度。本轮是可运行控制机制与其适用范围的验证，尚不能宣称顶会级首创新颖性。", "",
             "控制 RTL 的功能测试、位宽与状态资源见 rtl/validation.json 和 rtl/hardware_resources.json。没有 PDK 综合、布局布线和活动率功耗结果时，不填写 mm²、mW 或 1 GHz 时序通过结论。队列 RTL 验证不等于整套计算阵列完成物理实现。", ""]
    if guarded:
        lines[1:1] = ["", "此目录是允许等待的准入修订版。先比较两核预计完成时刻，再判断最佳核的队列空间；若最佳队列已满，描述符等待，不因另一个慢核有空位就强制溢出。两核已有任务仍继续执行。", "",
                      "该机制由上一版留出集的提前绑定回退问题推动。几何和阈值仍只在 18 个开发窗口重新选择，但同一留出集已经参与问题定位，本轮属于修复回归与探索验证，不能宣称全新盲测泛化证据。旧版冻结结果保留在上级目录。", ""]
    diagnosis = outdir / "PREFETCH_DIAGNOSIS.md"
    if diagnosis.exists():
        lines += ["## 已定位的预配置回退原因", "",
                  "当前 FIFO 满时会把任务派给较慢但仍有空位的核；预配置又会把该任务过早移入不可撤销的 shadow。B2 示例中，一个冷专家在 Dense 需 4,414 周期、在 Stream 只需 314 周期。提前绑定使总时间达到 8,843 周期；关闭预配置后任务保留在可旁路 FIFO，总时间为 4,587 周期。节省的 40 个准备周期抵不过错误核归属多出的计算。", "",
                  "这是当前准入 / 绑定策略的实际缺陷，不是预配置原则无效。改为‘允许等待、预解码与最终绑定解耦、可撤销未发射 shadow’是下一版建议，本轮冻结结果未做事后替换。完整例子和深度 16 的诊断对照见 [PREFETCH_DIAGNOSIS.md](PREFETCH_DIAGNOSIS.md)。", ""]
    audit_path = outdir / "INDEPENDENT_AUDIT.json"
    if audit_path.exists():
        audit = json.loads(audit_path.read_text())
        lines += ["## 独立审计", "",
                  f"审计状态 {audit['status']}：检查 {audit['all_rows_checked']:,} 个点及其第二次运行凭据；独立回放 {audit['main_rows_independently_replayed']:,} 个主结果、检查 {audit['main_whole_expert_jobs_checked']:,} 个整专家任务；另用独立逐周期参考模型核对 {audit['independent_tick_reference_cases']} 个场景。核预算、任务守恒、有限队列、上下文就绪、非重叠执行和吞吐下界均通过。该审计证明模型内部一致性，不代替真实阵列或 PPA 校准。", ""]
    validation_path = outdir / "rtl" / "validation.json"
    if validation_path.exists():
        verification = json.loads(validation_path.read_text())
        if verification.get("full_runtime_cycle_equivalence"):
            lines += ["## 控制 RTL 与事件模型回放", "",
                      f"Verilator 已通过 {verification['runtime_replay_cases']:,} 个完整控制回放场景，逐专家归属、开始 / 完成时刻和整层 makespan 与 Python 一致。具体覆盖域：{verification['runtime_replay_scope']}。阵列运算时间仍由已声明的解析服务模型提供，故这不能证明实际矩阵数据通路时序或 1 GHz 综合收敛。", ""]
    (outdir / "SUMMARY.md").write_text("\n".join(lines), encoding="utf-8")
    if guarded:
        revision_comparison(outdir)
    return {"records": len(rows), "groups": len(groups), "comparisons": comparisons,
            "workbook": str(outdir / "M_DPU_完整数据表.xlsx")}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--out", required=True)
    parser.add_argument("--frozen")
    args = parser.parse_args()
    print(json.dumps(generate(args.input, args.out, args.frozen), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
