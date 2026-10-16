"""Render screenshot-style tables solely from the frozen full MoE outputs."""
from __future__ import annotations
import argparse
import csv
import gzip
import hashlib
import json
from pathlib import Path
from statistics import mean
from collections import defaultdict

LABELS={'original':'PLENA原几何','extended_single':'DSE单核','extended_homogeneous':'DSE同构',
        'extended_heterogeneous':'DSE异构'}
PRIMARY=('extended_single','extended_homogeneous','extended_heterogeneous')
BATCHES=('2','4','8','16','64','96','128','7-Batch GM')


def load(path):
    with path.open() as stream:
        return list(csv.DictReader(stream))


def main(source,output,historical_directory=None):
    source,output=Path(source),Path(output)
    output.mkdir(parents=True,exist_ok=True)
    inputs=('FROZEN_SELECTION.json','workload_summary.csv','compute_summary.csv',
            'dataflow_tradeoff.csv','resource_ledger.csv','frozen_timing_sensitivity.csv',
            'paired_bootstrap.csv','DISPATCH_QUEUE_SUMMARY.json','split_diagnostic.csv')
    hashes={name:hashlib.sha256((source/name).read_bytes()).hexdigest() for name in inputs}
    compute=load(source/'compute_summary.csv')
    current={(r['design'],r['batch']):r for r in compute if r['split']=='heldout'}
    workloads=load(source/'workload_summary.csv')
    flows=load(source/'dataflow_tradeoff.csv')
    resources=load(source/'resource_ledger.csv')
    timing=load(source/'frozen_timing_sensitivity.csv')
    chosen=json.loads((source/'FROZEN_SELECTION.json').read_text())['families']
    tables=[]
    def add(name,headers,rows,note=''):
        tables.append({'title':name,'headers':headers,'rows':rows,'note':note})
    def number(row,key):
        return float(row[key])
    add('1. 硬件配置与实验约束',
        ['组织','物理尺寸 PM×PN×G','BLEN / MLEN','乘法器分配','计算与供数口径'],
        [[LABELS[name],' + '.join(f"{c['pm']}×{c['pn']}×{c['groups']}" for c in chosen[name]['cores']),
          ' + '.join(f"{c['b']} / {c['r']}" for c in chosen[name]['cores']),
          ' + '.join(str(c['macs']) for c in chosen[name]['cores']),
          'BF16/FP32；1GHz；理想操作数；整专家精确分配']
         for name in LABELS],
        '每组总计65,536乘法器。G=MLEN/BLEN，不是逻辑GEMM的K。几何在18个开发窗口选定后冻结。')
    add('2. Shared 与路由专家的实际矩阵维度',
        ['专家','阶段','输入 X','权重 W','输出 Y','M 定义'],
        [['Shared','Gate/Up','[B,2048]','[2048,2816]','[B,2816]','M=B'],
         ['Shared','Down','[B,2816]','[2816,2048]','[B,2048]','M=B'],
         ['路由专家','Gate/Up','[Me,2048]','[2048,1408]','[Me,1408]','实际Top-6路由确定'],
         ['路由专家','Down','[Me,1408]','[1408,2048]','[Me,2048]','实际Top-6路由确定']],
        '每个活跃专家包含Gate、Up、Down三个GEMM；Shared和全部活跃路由专家均计入。')
    rows=[]
    for batch in BATCHES[:-1]:
        records=[r for r in workloads if r['split']=='heldout' and r['batch']==batch]
        hist=[json.loads(r['routed_me_histogram']) for r in records]
        hot=mean(sum(v for k,v in h.items() if int(k)>2) for h in hist)
        cold=mean(sum(v for k,v in h.items() if int(k) in (1,2)) for h in hist)
        inactive=mean(h.get('0',0) for h in hist)
        assert abs(hot+cold+inactive-64)<1e-10
        rows.append(['B'+batch,len(records),hot,cold,inactive,hot+cold,1])
    add('3. 真实负载统计：路由专家激活分布',
        ['Batch','留出窗口','热路由专家 Me>2','冷路由专家 Me=1/2','未激活 Me=0','活跃路由专家合计','Shared专家'],rows,
        'DeepSeek-V2-Lite，64个路由专家，Top-6；专家数量为逐窗口算术平均，Shared另计。')
    rows=[]
    for batch in BATCHES:
        group=[current[(design,batch)] for design in PRIMARY]
        rows.append(['B'+batch if batch.isdigit() else '七Batch等权GM']+
          [number(r,'compute_ms') for r in group]+[number(r,'useful_GFLOP_s') for r in group]+
          [number(group[0],'compute_ms')/number(group[2],'compute_ms'),
           number(group[1],'compute_ms')/number(group[2],'compute_ms')])
    add('4. 纯计算 DSE 主结果',
        ['Batch','单核Compute ms','同构Compute ms','异构Compute ms','单核有效GFLOP/s',
         '同构有效GFLOP/s','异构有效GFLOP/s','异构/单核加速比','异构/同构加速比'],rows,
        '理想操作数及端口；完整Shared+MoE矩阵任务的完成时间。每Batch先取窗口GM，综合再取七Batch等权GM。有效FLOPs=2×有用MAC。')
    rows=[]
    for batch in BATCHES:
        group=[current[(design,batch)] for design in PRIMARY]
        rows.append(['B'+batch if batch.isdigit() else '七Batch等权GM',number(group[0],'unique_weight_MiB'),
            number(group[0],'unique_weight_fetch_ms_at256')]+
            [number(r,'unique_weight_required_GB_s') for r in group])
    add('5. 理论取数服务与维持纯计算的带宽需求',
        ['Batch','唯一BF16权重 MiB','权重/256GB/s服务 ms','单核所需权重GB/s',
         '同构所需权重GB/s','异构所需权重GB/s'],rows,
        '唯一权重=33MiB Shared+16.5MiB×活跃路由专家；本表数值为窗口GM。所需带宽=W/Compute，是必要平均供数率，不是实测HBM吞吐。')
    rows=[]
    for batch in BATCHES:
        group=[current[(design,batch)] for design in PRIMARY]
        rows.append(['B'+batch if batch.isdigit() else '全留出窗口最大值']+
          [number(r,'resident_full_payload_upper_MiB') for r in group])
    add('6. OS全驻留容量上界',
        ['Batch','单核 MiB','同构 MiB','异构 MiB'],rows,
        '容量取各组窗口最大值，包含活跃权重、中间结果、固定寄存器/树/操作数缓冲；这是保守上界，不是已实现SRAM宏容量。')
    rows=[]
    for r in resources:
        if r['design'] not in LABELS:
            continue
        rows.append([LABELS[r['design']]]+[number(r,k)/1024 for k in
          ('PE_FP32_bytes','shadow_FP32_bytes','M_SUM_tree_FP32_register_bytes',
           'full_operand_double_buffer_bytes','static_storage_bytes')])
    add('7. 固定计算硬件的寄存器与缓冲账本',
        ['组织','PE部分和 KiB','Shadow部分和 KiB','M_SUM树寄存器 KiB','操作数双缓冲 KiB','固定存储合计 KiB'],rows,
        'PE/Shadow/tree是寄存器状态；不能将本表所有容量统称为SRAM。面积和新几何Fmax尚未综合校准。')
    rows=[]
    for design in PRIMARY:
        for r in flows:
            if r['split']=='heldout' and r['design']==design:
                rows.append([LABELS[design],r['dataflow'],number(r,'compute_ms'),
                  number(r,'pool_weight_refill_required_GB_s')/1000,
                  number(r,'pool_X_refill_required_GB_s'),number(r,'pool_context_required_GB_s')/1000,
                  number(r,'local_scratch_tree_upper_KiB')/1024,
                  number(r,'no_pool_weight_refetch_factor')])
    assert len(rows)==21
    add('8. 全数据流组合：冻结几何，七Batch等权GM',
        ['组织','核0/核1驻留','Compute ms','池W TB/s','池X GB/s','池上下文 TB/s',
         '本地/上下文/树容量上界 MiB','无驻留池权重重读倍率'],rows,
        'PE仍为OS；WS/IS为操作数循环驻留扩展。理想无限端口下Compute相同，流量/容量不同。池内refill不是HBM流量；上下文读写已收费。')
    rows=[]
    for design in LABELS:
        values=[next(r for r in timing if r['design']==design and r['split']=='heldout' and r['profile']==profile)
          for profile in ('pipelined_tree1','pipelined_tree2','serialized_output')]
        rows.append([LABELS[design]]+[number(r,'GM_compute_ms') for r in values])
    add('9. 同一冻结硬件的计算时序敏感性',
        ['组织','归约树1周期/级 ms','归约树2周期/级 ms','输出串行退休 ms'],rows,
        '这是计算流水线敏感性，三个模式均无访存端口等待。不能将它们改名为旧Pipelined/Port_Tight系统模式。')
    rows=[]
    for r in load(source/'paired_bootstrap.csv'):
        if r['split']=='heldout' and r['candidate']=='extended_heterogeneous':
            rows.append([LABELS[r['baseline']],number(r,'batch_balanced_GM_speedup'),
                         number(r,'bootstrap_95pct_lower'),number(r,'bootstrap_95pct_upper'),
                         int(r['candidate_win_windows']),int(r['candidate_loss_windows'])])
    add('10. 异构相对优化后基线的配对比较',
        ['基线','异构加速比','95%区间下限','95%区间上限','胜窗口','负窗口'],rows,
        '2000次按Batch分层配对bootstrap；区间仅反映历史捕获窗口的波动，不包含解析模型/硬件校准误差。')
    add('11. 截图中的系统与预测器项目：本轮测试状态',
        ['项目','65,536平台状态','可报告内容'],
        [['256GB/s Pipelined端到端','未测试','只有唯一权重/256GB/s的服务等价时间'],
         ['Port_Tight端到端','未测试','当前只有理想端口模型；树2周期并非Port_Tight'],
         ['1024乘法器敏感性','未测试','截图是旧12,288预算下的另一个实验'],
         ['nominal/random/static/BTB/EMA/ours/oracle','未测试','本轮使用整专家离线精确分配；没有预测器MAE/命中率'],
         ['原生Rust/RTL校准','未完成','以上Compute来自PLENA拓扑一致的BF16解析模型']],
        '旧系统/预测器数值属于12,288乘法器实验，不并入65,536主表。')
    historical_hashes={}
    if historical_directory is not None:
        historical=Path(historical_directory)
        historic_paths=[historical/'E4/selected_designs.json',historical/'E5/predictor/predictor_table.csv',
                        historical/'E5/predictor/task_predictions.csv.gz']
        historical_hashes={str(p):hashlib.sha256(p.read_bytes()).hexdigest() for p in historic_paths}
        details=json.loads(historic_paths[0].read_text())
        old_names={'pipelined':'Pipelined','port_tight':'Port Tight'}
        old_designs={mode:('B1','B2',details['best_hetero_by_mode'][mode]) for mode in old_names}
        rows=[]
        for mode,label in old_names.items():
            values=[details['design_details'][mode]['C0'][d]['geometry'] for d in old_designs[mode]]
            budgets=[sum(__import__('math').prod(map(int,c.split('x'))) for c in value.split('+')) for value in values]
            assert budgets==[12288]*3
            rows.append([label,*values,'WS / WS+WS / WS+WS',12288])
        add('历史附表 A. 预测器实验配置（12,288乘法器）',
            ['模式','单核','同构','异构','驻留','总乘法器'],rows,
            '这是第三轮历史系统实验，256GB/s、135留出窗口、控制窗口8records；与本轮65,536纯计算表独立。')
        errors=defaultdict(lambda:[0.0,0])
        with gzip.open(historic_paths[2],'rt') as stream:
            for row in csv.DictReader(stream):
                if float(row['bw_GBps'])!=256:
                    continue
                key=(row['onchip_mode'],row['design'],row['method'])
                errors[key][0]+=abs(float(row['predicted_cycles'])-float(row['actual_cycles']))
                errors[key][1]+=1
        old_latency={(r['onchip_mode'],r['design'],r['method']):float(r['all_geomean_ms'])
                     for r in load(historic_paths[1]) if float(r['bw_GBps'])==256}
        for mode,label in old_names.items():
            rows=[]
            for method in ('nominal','random','static','btb','ema','ours','oracle'):
                keys=[(mode,d,method) for d in old_designs[mode]]
                names={'btb':'BTB','ema':'EMA','oracle':'oracle回放'}
                rows.append([names.get(method,method)]+[old_latency[k] for k in keys]+
                            [errors[k][0]/errors[k][1] for k in keys])
            add(f'历史附表 {label}. 预测方法对照（12,288乘法器）',
                ['方法','单核 ms','同构 ms','异构 ms','单核MAE 周期','同构MAE 周期','异构MAE 周期'],rows,
                '时间为135窗口GM，并非七Batch等权GM；MAE为逐任务算术平均。random随机估计时长；oracle固定ours计划回放，不是全局最优调度。')
    json_path=output/'tables.json'
    json_path.write_text(json.dumps({'source':str(source),'tables':tables},ensure_ascii=False,indent=2)+'\n')
    for i,table in enumerate(tables,1):
        with (output/f'table_{i:02d}.csv').open('w',newline='') as stream:
            writer=csv.writer(stream);writer.writerow(table['headers']);writer.writerows(table['rows'])
    from openpyxl import Workbook,load_workbook
    from openpyxl.styles import Alignment,Font,PatternFill,Border,Side
    from openpyxl.utils import get_column_letter
    book=Workbook(); sheet=book.active;sheet.title='按截图格式汇总'
    widths=[25,31,27,27,29,29,28,31,29]
    for ci,width in enumerate(widths,1): sheet.column_dimensions[get_column_letter(ci)].width=width
    sheet.freeze_panes='B5';row=1
    sheet.merge_cells(start_row=row,start_column=1,end_row=row,end_column=9)
    sheet.cell(row,1,'65,536乘法器 · Shared + 真实Top-6 MoE · 2026-10-12').font=Font(size=16,bold=True,color='17365D')
    row+=2;expected=[]
    for table in tables:
        sheet.merge_cells(start_row=row,start_column=1,end_row=row,end_column=9)
        sheet.cell(row,1,table['title']).font=Font(size=13,bold=True,color='FFFFFF')
        for ci in range(1,10):sheet.cell(row,ci).fill=PatternFill('solid',fgColor='17365D')
        row+=1
        for ci,value in enumerate(table['headers'],1):
            cell=sheet.cell(row,ci,value);cell.font=Font(bold=True);cell.fill=PatternFill('solid',fgColor='DBE5F1');cell.alignment=Alignment(wrap_text=True,vertical='center')
            expected.append((row,ci,value))
        sheet.row_dimensions[row].height=34;row+=1
        for values in table['rows']:
            for ci,value in enumerate(values,1):
                cell=sheet.cell(row,ci,value);cell.alignment=Alignment(vertical='center',wrap_text=True)
                if isinstance(value,float):cell.number_format='#,##0.000000'
                expected.append((row,ci,value))
            sheet.row_dimensions[row].height=34;row+=1
        sheet.merge_cells(start_row=row,start_column=1,end_row=row,end_column=9)
        sheet.cell(row,1,table['note']).alignment=Alignment(wrap_text=True,vertical='center')
        sheet.row_dimensions[row].height=44;row+=3
    source_sheet=book.create_sheet('口径与来源')
    source_sheet.append(['源目录',str(source)]);source_sheet.append(['源文件','SHA256'])
    for name,digest in hashes.items():source_sheet.append([name,digest])
    for name,digest in historical_hashes.items():source_sheet.append([name,digest])
    source_sheet.column_dimensions['A'].width=44;source_sheet.column_dimensions['B'].width=100
    file=output/'Shared_MoE_按截图格式完整表.xlsx';book.save(file)
    readback=load_workbook(file,data_only=True)
    for ri,ci,value in expected:
        actual=readback['按截图格式汇总'].cell(ri,ci).value
        if isinstance(value,float):assert abs(actual-value)<=1e-10*max(1,abs(value))
        else:assert actual==value
    assert hashes=={name:hashlib.sha256((source/name).read_bytes()).hexdigest() for name in inputs}
    assert historical_hashes=={name:hashlib.sha256(Path(name).read_bytes()).hexdigest() for name in historical_hashes}
    receipt={'status':'PASS','tables':len(tables),'data_rows':sum(len(t['rows']) for t in tables),
             'workbook_readback_cells':len(expected),'source_sha256':hashes,'workbook':str(file),
             'original_results_modified':False,'new_experiments_claimed':False,
             'historical_source_sha256':historical_hashes,'historical_predictor_budget':12288 if historical_hashes else None}
    (output/'TABLE_AUDIT.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({k:receipt[k] for k in ('status','tables','data_rows','workbook_readback_cells','workbook')},ensure_ascii=False))


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--source',type=Path,required=True);parser.add_argument('--out',type=Path,required=True)
    parser.add_argument('--historical-directory',type=Path)
    args=parser.parse_args();main(args.source,args.out,args.historical_directory)
