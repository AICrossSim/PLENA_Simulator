"""Render completed compute-ceiling DSE and reverse-supply results.

This is a read-only result consumer.  It never launches a simulation, changes
frozen selections, or waits for running work.  Run after COMPLETION.json exists.
The reporter uses only the standard library and, for XLSX, openpyxl.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import statistics
from collections import defaultdict
from pathlib import Path
from typing import Any

BATCHES = (2, 4, 8, 16, 64, 96, 128)
FAMILIES = ("single", "homogeneous", "heterogeneous")
LABELS = {
    "original_single": "原始单核",
    "original_homogeneous": "原始同构双核",
    "original_heterogeneous": "原始异构双核",
    "best_single": "DSE 最优单核",
    "best_homogeneous": "DSE 最优同构",
    "best_heterogeneous": "DSE 最优异构",
    "robust_single": "时序稳健单核",
    "robust_homogeneous": "时序稳健同构",
    "robust_heterogeneous": "时序稳健异构",
}
MAIN_LABELS = ("original_single", "best_single", "best_homogeneous", "best_heterogeneous")
ORDER = tuple(LABELS)


def _json(path: Path, *, required: bool = True) -> Any:
    if not path.exists():
        if required:
            raise FileNotFoundError(path)
        return None
    return json.loads(path.read_text())


def _csv(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    with path.open(newline="") as stream:
        return [_coerce(row) for row in csv.DictReader(stream)]


def _coerce(row: dict[str, str]) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for key, value in row.items():
        if value == "":
            result[key] = None
        elif value == "True":
            result[key] = True
        elif value == "False":
            result[key] = False
        else:
            try:
                result[key] = int(value)
            except (ValueError, TypeError):
                try:
                    result[key] = float(value)
                except (ValueError, TypeError):
                    result[key] = value
    return result


def _write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    keys = list(dict.fromkeys(key for row in rows for key in row))
    with path.open("w", newline="") as stream:
        if not keys:
            return
        writer = csv.DictWriter(stream, fieldnames=keys, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def _nested(value: dict[str, Any], *keys: str) -> Any:
    for key in keys:
        if not isinstance(value, dict) or key not in value:
            return None
        value = value[key]
    return value


def _gm(values: list[float]) -> float | None:
    if not values or any(value < 0 for value in values):
        return None
    if any(value == 0 for value in values):
        return 0.0
    return math.exp(statistics.mean(math.log(value) for value in values))


def _balanced(rows: list[dict[str, Any]], field: str, batch: int | str) -> float | None:
    groups: dict[int, list[float]] = defaultdict(list)
    for row in rows:
        value = row.get(field)
        if value is not None and (batch == "all" or row["batch"] == batch):
            groups[row["batch"]].append(float(value))
    if batch == "all":
        if not groups:
            return None
        means = [_gm(values) for _, values in sorted(groups.items())]
        return _gm([value for value in means if value is not None]) if all(value is not None for value in means) else None
    return _gm(groups.get(batch, []))


def _shape(core: dict[str, Any]) -> str:
    return "×".join(str(core[key]) for key in ("pm", "pn", "pk"))


def _format(value: Any, places: int = 3) -> str:
    if value is None:
        return "未提供"
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, float):
        return f"{value:,.{places}f}"
    return str(value).replace("|", "\\|").replace("\n", " ")


def _table(headers: list[str], rows: list[list[Any]], places: int = 3) -> str:
    if not rows:
        return "该项没有对应的完成结果。\n"
    output = ["| " + " | ".join(headers) + " |", "| " + " | ".join("---" for _ in headers) + " |"]
    output.extend("| " + " | ".join(_format(value, places) for value in row) + " |" for row in rows)
    return "\n".join(output) + "\n"


def _flatten_supply(source: dict[str, Any], compute: dict[str, Any]) -> dict[str, Any]:
    cycles = source["compute_cycles"]
    assert cycles == compute["cycles"]
    frequency = source["frequency_hz"]
    row = {
        "label": source["label"], "geometry": source["geometry"],
        "window_id": source["window_id"], "batch": source["batch"],
        "compute_cycles": cycles, "compute_ms": source["compute_latency_ms"],
        "useful_macs": compute["useful_macs"], "issued_macs": compute["issued_macs"],
        "padding_macs": compute["padding_macs"],
        "spatial_utilization": compute["spatial_utilization"],
        "wall_mac_utilization": compute["wall_mac_utilization"],
        "useful_GFLOP_s": 2 * compute["useful_macs"] * frequency / cycles / 1e9,
        "core_completion_gap_cycles": compute["finish_gap_cycles"],
    }
    external = source.get("external", {})
    capacity = source.get("capacity", {})
    for key in (
        "selected_unique_bf16_weight_bytes", "cold_weight_average_required_GB_s",
        "fully_resident_weight_HBM_bytes_per_replay", "compulsory_global_input_bf16_bytes",
        "combined_layer_output_fp32_bytes", "cold_weights_plus_input_and_output_bytes",
        "cold_weights_plus_input_and_output_average_GB_s",
    ):
        row[key] = external.get(key)
    for key in (
        "selected_layer_weights_bytes", "all_routed_and_shared_weight_bytes",
        "global_input_X_bf16_bytes", "combined_output_Y_fp32_bytes",
        "sum_one_expert_Z_per_core_bytes", "all_selected_experts_Z_bf16_bytes_if_retained",
        "sum_declared_per_core_buffer_peaks_bytes",
        "fully_resident_selected_weight_os_working_capacity_bytes",
        "fully_resident_all_model_weight_os_working_capacity_bytes",
    ):
        row[key] = capacity.get(key)
    onchip = source.get("onchip", {})
    for boundary, nested_key in (
        ("array", "os_kmajor_array_average_rates"),
        ("pool", "os_kmajor_shared_pool_average_rates"),
    ):
        for key, value in onchip.get(nested_key, {}).items():
            row[f"{boundary}_{key}"] = value
    for key, value in source.get("peak_refill", {}).items():
        if isinstance(value, (int, float)):
            row[key] = value
    cores = source.get("per_core", [])
    for key in (
        "peak_partial_records_full_bytes", "x_k_panel_double_buffer_bytes",
        "weight_tile_double_buffer_bytes", "pipeline_output_values_inflight_full_bytes",
    ):
        row[f"sum_core_{key}"] = sum(core.get("capacity", {}).get(key, 0) for core in cores)
    row["per_core_shapes"] = " + ".join(_shape(core["shape"]) for core in cores)
    row["per_core_partial_peaks_KiB"] = " / ".join(
        f"{core.get('capacity', {}).get('peak_partial_records_full_bytes', 0)/1024:.6f}" for core in cores)
    row["per_core_completion_cycles"] = " / ".join(str(value) for value in compute["core_cycles"])
    return row


def _aggregate_supply(rows: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    rates, capacities = [], []
    rate_fields = (
        "compute_ms", "cold_weight_average_required_GB_s", "useful_GFLOP_s",
        "pool_weight_full_slice_average_GB_s", "pool_x_full_slice_average_GB_s",
        "array_weight_full_slice_average_GB_s", "array_x_full_slice_average_GB_s",
        "array_partial_read_full_record_average_GB_s", "array_partial_write_full_record_average_GB_s",
    )
    capacity_fields = (
        "selected_unique_bf16_weight_bytes", "all_routed_and_shared_weight_bytes",
        "sum_core_peak_partial_records_full_bytes", "sum_core_x_k_panel_double_buffer_bytes",
        "sum_core_weight_tile_double_buffer_bytes", "sum_one_expert_Z_per_core_bytes",
        "global_input_X_bf16_bytes", "combined_output_Y_fp32_bytes",
        "fully_resident_selected_weight_os_working_capacity_bytes",
        "fully_resident_all_model_weight_os_working_capacity_bytes",
    )
    for label in ORDER:
        label_rows = [row for row in rows if row["label"] == label]
        if not label_rows:
            continue
        for batch in (*BATCHES, "all"):
            selected = [row for row in label_rows if batch == "all" or row["batch"] == batch]
            if not selected:
                continue
            rates_row = {"label": label, "geometry": selected[0]["geometry"], "batch": batch, "windows": len(selected)}
            for field in rate_fields:
                rates_row[f"{field}_GM"] = _balanced(selected, field, batch)
            for field in (
                "cold_weight_average_required_GB_s", "simultaneous_pool_weight_GB_s_envelope_with_M_reuse",
                "simultaneous_pool_X_panel_steady_state_GB_s_envelope",
                "simultaneous_array_weight_full_GB_s_upper_envelope",
                "simultaneous_array_X_full_GB_s_upper_envelope",
                "simultaneous_partial_read_full_GB_s_upper_envelope",
                "simultaneous_partial_write_full_GB_s_upper_envelope",
            ):
                values = [row[field] for row in selected if row.get(field) is not None]
                rates_row[f"{field}_max"] = max(values) if values else None
            rates.append(rates_row)
            capacities_row = {"label": label, "geometry": selected[0]["geometry"], "batch": batch, "windows": len(selected)}
            for field in capacity_fields:
                values = [row[field] for row in selected if row.get(field) is not None]
                capacities_row[f"{field}_max"] = max(values) if values else None
                capacities_row[f"{field}_AM"] = statistics.mean(values) if values else None
            capacities.append(capacities_row)
    return rates, capacities


def _config_rows(frozen: dict[str, Any]) -> list[dict[str, Any]]:
    robust = {row["family"]: row for row in frozen.get("robust", [])}
    rows = []
    for label in ORDER:
        cores = frozen.get("configs", {}).get(label)
        if not cores:
            continue
        macs = [core["pm"] * core["pn"] * core["pk"] for core in cores]
        if sum(macs) != 12288:
            raise ValueError(f"Unexpected MAC budget for {label}")
        rows.append({
            "label": label, "organization": LABELS[label],
            "physical_M_N_K": " + ".join(_shape(core) for core in cores),
            "per_core_multipliers": " + ".join(str(value) for value in macs),
            "total_multipliers": sum(macs),
            "MAC_ratio": ":".join(str(value // math.gcd(*macs)) for value in macs),
            "selection": "固定原始配置" if label.startswith("original_") else (
                "开发集跨三个时序假设的最小最坏遗憾" if label.startswith("robust_") else "开发集主时序最优"),
            "worst_development_timing_regret": robust.get(label.removeprefix("robust_"), {}).get("worst_development_timing_regret") if label.startswith("robust_") else None,
        })
    return rows


def _scope_rows(contract: dict[str, Any], completion: dict[str, Any]) -> list[dict[str, Any]]:
    return [
        {"项目": "主模型", "设置": "DeepSeek-V2-Lite；64 routed experts，top-K 6，Shared 合并宽度 2816"},
        {"项目": "GEMM", "设置": "Routed Gate/Up [Me,2048]×[2048,1408]；Down [Me,1408]×[1408,2048]；Shared 将 1408 替换为 2816，Me=batch"},
        {"项目": "开发与评估", "设置": f"{completion.get('development_windows', '未提供')} 开发窗口选硬件；冻结后 {completion.get('primary_windows', '未提供')} 历史留出窗口评估；7 个 batch 等权"},
        {"项目": "物理资源", "设置": "每个设计主乘法器总数 12,288；BF16 X/W，FP32 部分和；等乘法器不等于等面积"},
        {"项目": "时钟与主计算时序", "设置": "假定 1 GHz，1 周期=1 ns；dot 20 周期+提交 1 周期；II=1；不同输出可重叠，相同输出 K 依赖必须满足"},
        {"项目": "数组几何", "设置": f"M 为全部合法正整数因子，N 为 {contract.get('N_alignment_primary', '未提供')} 的倍数；K={min(contract.get('PK_values', [16]))}…{max(contract.get('PK_values', [2048]))}，步长 16；总乘法器严格相等"},
        {"项目": "主数据流", "设置": "Gate/Up 共同 K-major、输出部分和全程驻留；每个 K 段的 X panel 跨 N 复用，W tile 跨 M 复用；不是只按 FLOPs 除算力"},
        {"项目": "派工", "设置": "整专家最优离线分配；Shared 可分给任一核，但主表不拆 Shared；不计动态控制服务成本"},
        {"项目": "排除成本", "设置": "HBM、SRAM 端口/容量等待、DMA、router、SiLU、combine、派工服务；表中是三个 GEMM 的 compute makespan，不是完整 FFN 或模型 E2E"},
        {"项目": "反推供数", "设置": "分别统计外部冷加载、共享池→本地、阵列操作数/部分和流量；容量与带宽分开；平均带宽不保证每块按时就绪"},
        {"项目": "校准与泛化", "设置": "本轮分析模型，尚未原生 Rust/RTL 校准；历史留出已多次使用；Qwen 冻结硬件补充评估不再搜索"},
    ]


def _ready_pool_section(rows: list[dict[str, Any]]) -> str:
    """Explain the optional stronger issue-policy diagnostic separately."""
    if not rows:
        return ""
    output = [
        "## 跨专家交织发射的代表窗口诊断\n",
        "在同一套冻结硬件和同一份整专家归属下，允许各核从多个专家的就绪输出记录中交织发射。每个输出的K段依赖及每个专家的Gate/Up到Down依赖保持不变。下面每个batch只使用预先指定的一个代表窗口，比较原串行专家调度和更好的可行调度；主表与搜索证书未被替换。\n",
    ]
    labels = ("best_single", "best_homogeneous", "best_heterogeneous")
    for model in sorted({row["model"] for row in rows}):
        selected = [row for row in rows if row["model"] == model]
        by = {(row["design"], row["batch"]): row for row in selected}
        table = []
        for batch in BATCHES:
            if not all((label, batch) in by for label in labels):
                continue
            ids = {by[(label, batch)]["window_id"] for label in labels}
            if len(ids) != 1:
                raise ValueError("Ready-pool organizations do not use the same representative window")
            table.append([f"B{batch}", *[
                value
                for label in labels
                for value in (by[(label, batch)]["whole_expert_cycles"] / 1e6,
                              by[(label, batch)]["best_feasible_ms"])
            ]])
        output.extend([
            f"**{model}**\n",
            _table(["Batch", "单核原串行 ms", "单核可行调度 ms", "同构原串行 ms", "同构可行调度 ms", "异构原串行 ms", "异构可行调度 ms"], table, 6),
        ])
    descriptor_counts = [row["pooled_issue_descriptor_count"] for row in rows
                         if row.get("pooled_issue_descriptor_count") is not None]
    partial_peaks = [row["pool_partial_peak_KiB"] for row in rows
                     if row.get("pool_partial_peak_KiB") is not None]
    if descriptor_counts and partial_peaks:
        output.append(
            f"这组交织方案预创建 {min(descriptor_counts):,}～{max(descriptor_counts):,} 个输出记录描述符，"
            f"部分和存活峰值为 {min(partial_peaks):,.2f}～{max(partial_peaks):,.2f} KiB。"
            "各窗口的原始/交织部分和容量和描述符数量在工作簿“跨专家交织代表窗口”页中逐项列出。\n"
        )
    output.append(
        "可行调度列允许每个核在交织方案与原串行方案之间回退，因此不会把更差的交织方案强加给基线。"
        "它不是全局最优调度、没有重新搜索几何，也不是全部窗口上的稳健结论。"
        "跨专家交织会增加同时存活的上下文并改变W/X复用，"
        "不能沿用主表OS程序的供数带宽、缓存容量或流量估算；需要为这套程序重新计算供数与硬件控制成本。\n"
    )
    return "\n".join(output)


def _load_pooled(out: Path) -> dict[str, Any] | None:
    """Read the independent pooled campaign only after both receipts exist."""
    directory = out / "pooled"
    if not all((directory / name).exists() for name in
               ("COMPLETION.json", "COMPLETION_REFERENCE.json")):
        return None
    completion = _json(directory / "COMPLETION.json")
    reference = _json(directory / "COMPLETION_REFERENCE.json")
    if completion.get("heldout_windows") != 135 or completion.get("repeat_exactly_equal") is not True:
        raise ValueError("Pooled campaign is not a completed repeated 135-window evaluation")
    if reference.get("windows") != 135 or reference.get("all_repeated_results_identical") is not True:
        raise ValueError("Pooled original reference has not completed repeated replay")
    frozen = _json(directory / "FROZEN_SELECTION.json")
    original = _json(directory / "ORIGINAL_REFERENCE.json")
    configurations = {**original["configs"], **frozen["configs"]}
    configs = []
    names = {"original_single": "原始单核", "single": "pooled 最优单核",
             "homogeneous": "pooled 最优同构", "heterogeneous": "pooled 最优异构"}
    for label in ("original_single", *FAMILIES):
        cores = configurations[label]
        macs = [c["pm"] * c["pn"] * c["pk"] for c in cores]
        if sum(macs) != 12288:
            raise ValueError("Pooled configuration violates the fixed multiplier budget")
        configs.append({"label": label, "organization": names[label],
                       "physical_M_N_K": " + ".join(_shape(c) for c in cores),
                       "per_core_multipliers": " + ".join(str(v) for v in macs),
                       "total_multipliers": sum(macs),
                       "MAC_ratio": ":".join(str(v // math.gcd(*macs)) for v in macs)})
    compute = _csv(directory / "original_single_by_batch.csv") + _csv(directory / "compute_by_batch.csv")
    comparisons = _csv(directory / "comparisons.csv")
    certificates = _csv(directory / "development_certificates.csv")
    if len(compute) != 32 or len(comparisons) != 24 or len(certificates) != 3:
        raise ValueError("Incomplete pooled compute, comparison or certificate tables")
    raw = []
    sources = [directory / name for name in (
        "COMPLETION.json", "COMPLETION_REFERENCE.json", "CONTRACT.json", "FROZEN_SELECTION.json",
        "ORIGINAL_REFERENCE.json", "original_single_by_batch.csv", "compute_by_batch.csv",
        "comparisons.csv", "development_certificates.csv", "development_all_geometry_bounds.csv",
        "development_near_optimal_bounds.csv", "heldout_per_window.csv",
    ) if (directory / name).exists()]
    for label in ("original_single", *FAMILIES):
        path = directory / f"heldout_{label}.json"
        results = _json(path)
        if len(results) != 135 or len({r["window_id"] for r in results}) != 135:
            raise ValueError("Pooled per-window output does not contain 135 distinct windows")
        raw.extend({"label": label, **row} for row in results)
        sources.append(path)
    supply = []
    supply_per_window = []
    if (directory / "REVERSE_SUPPLY.json").exists():
        receipt = _json(directory / "REVERSE_SUPPLY.json")
        if receipt.get("status") != "PASS" or receipt.get("sources_unchanged") is not True:
            raise ValueError("Pooled reverse-supply receipt is not verified")
        supply = _csv(directory / "reverse_supply_by_batch.csv")
        supply_per_window = _csv(directory / "reverse_supply_per_window.csv")
        if len(supply) != receipt.get("by_batch_rows") or len(supply_per_window) != receipt.get("window_rows"):
            raise ValueError("Pooled supply receipt row counts disagree with output CSVs")
        sources.extend(directory / name for name in (
            "reverse_supply_by_batch.csv", "reverse_supply_per_window.csv", "REVERSE_SUPPLY.json"))
    global_ceiling = []
    if (directory / "GLOBAL_MAC_CEILING.json").exists():
        receipt = _json(directory / "GLOBAL_MAC_CEILING.json")
        if receipt.get("status") != "PASS" or receipt.get("window_count") != 135:
            raise ValueError("Global MAC ceiling receipt is not a verified 135-window bound")
        global_ceiling = _csv(directory / "global_MAC_ceiling.csv")
        if len(global_ceiling) != 24 or any(
            row.get("bound_is_not_a_feasible_schedule") is not True for row in global_ceiling):
            raise ValueError("Incomplete or mislabeled global MAC ceiling")
        sources.extend(directory / name for name in ("global_MAC_ceiling.csv", "GLOBAL_MAC_CEILING.json"))
    return {"completion": completion, "reference": reference,
            "contract": _json(directory / "CONTRACT.json"), "configs": configs,
            "compute": compute, "comparisons": comparisons, "certificates": certificates,
            "per_window": raw, "supply": supply, "supply_per_window": supply_per_window,
            "global_ceiling": global_ceiling,
            "source_paths": sources}


def _pooled_section(pooled: dict[str, Any] | None) -> str:
    if pooled is None:
        return ""
    by = {(r["label"], r["batch"]): r for r in pooled["compute"]}
    comp = {(r["candidate"], r["baseline"], r["batch"]): r for r in pooled["comparisons"]}
    labels = ("original_single", *FAMILIES)
    text = [
        "## 更强执行契约：全部窗口跨专家交织与全域几何搜索\n",
        "这一组允许同一核在不同专家的就绪输出记录间交织发射，不同输出可以填补K依赖等待；全部操作数驻留，允许足够多的上下文。仍然保持每个专家归一个核、每个输出的K段按依赖提交、Gate/Up完成后才能开始该专家Down；不进行同次发射跨专家拼行，也不拆Shared列。原始单核与三种搜索组织均获得相同交织能力。\n",
        "因此，上方整专家依次执行契约的收益比例不能作为无限上下文下的普遍结论。两套契约的硬件都在各自18个开发窗口上选定、135个历史留出窗口前冻结，不能把旧单核与新异构交叉比较。新组只采用flat20时序，尚未原生校准。\n",
        _table(["组织", "每核 M×N×K", "每核乘法器", "算力比例"],
               [[r["organization"], r["physical_M_N_K"], r["per_core_multipliers"], r["MAC_ratio"]]
                for r in pooled["configs"]]),
    ]
    hetero = next(r for r in pooled["configs"] if r["label"] == "heterogeneous")
    if hetero["MAC_ratio"] == "1:1":
        text.append("这里选中的异构是两个相同乘法器数量、不同N/K几何的核；它是形状异构，不是大核算力多、小核算力少。不能把这一结果称为4+2大小算力分配的证明。\n")
    hs = comp[("heterogeneous", "single", "all")]
    hh = comp[("heterogeneous", "homogeneous", "all")]
    def relative(row: dict[str, Any]) -> str:
        reduction = row["latency_reduction_pct"]
        return f"快{reduction:.4f}%" if reduction >= 0 else f"慢{-reduction:.4f}%"
    text.append(f"新执行契约下，冻结形状异构相对最优单核{relative(hs)}，相对最优同构{relative(hh)}；未达到相对两个基线都快5%的进入校准门槛。\n")
    rows = []
    for batch in (*BATCHES, "all"):
        cs = comp.get(("heterogeneous", "single", batch), {})
        ch = comp.get(("heterogeneous", "homogeneous", batch), {})
        rows.append([f"B{batch}" if batch != "all" else "七batch等权GM",
                     *[by[(label, batch)]["compute_ms_geomean"] for label in labels],
                     cs.get("speedup"), ch.get("speedup")])
    text.extend([
        _table(["Batch", "原始单核 ms", "新最优单核 ms", "新最优同构 ms", "新最优异构 ms",
                "异构/单核加速比", "异构/同构加速比"], rows, 6),
        "上表加速比按同一窗口配对、七batch等权计算；大于1才表示异构更快。全部4组每个留出点进行了两次真实事件重放、结果一致。这里是纯计算makespan，不是完整FFN或模型端到端耗时。\n",
        _table(["比较", "七batch等权加速比", "延迟减少 %", "窗口bootstrap 95%区间"],
               [[f"{r['candidate']} 相对 {r['baseline']}", r["speedup"], r["latency_reduction_pct"],
                 f"[{r['speedup_CI95_low']:.6f}, {r['speedup_CI95_high']:.6f}]"]
                for r in pooled["comparisons"] if r["batch"] == "all"], 6),
        "部分区间退化成一个数，是因为本契约的PM=1与固定top-K使总发射量主要由总token数决定，当前各窗口的比例相同；这不表示硬件误差为零。窗口bootstrap不包含端口、时钟、面积或模型失配，不能代替原生校准。\n",
        "### 全域下界与实际可行执行的差距\n",
        _table(["组织", "域内几何数", "精确issue分配点", "实际事件重放几何数",
                "开发集可行上界 cycles", "开发集已认证下界 cycles", "上界/下界−1 %"],
               [[r["family"], r["domain_geometries"], r["exact_issue_floor_points"],
                 r["actual_event_geometries"], r["best_feasible_score_cycles"],
                 r["family_min_valid_LB_cycles"], 100 * r["certified_relative_gap"]]
                for r in pooled["certificates"]], 6),
        "先对完整尺寸域求合法的分数化issue负载下界，再用整数归属精确求解认证该组织的最小issue-floor。任何实际执行还至少需要最后一次发射后20周期排空；实际RR事件重放提供可行上界。若某组织开发目标的上界与全域认证下界相等，便在该有限尺寸域、整专家归属和flat20契约下达到下界；这不表示所有候选调度都已精确求解，也不认证任意数据流、Shared拆分或有限SRAM实现的全局最优。\n",
    ])
    if pooled.get("global_ceiling"):
        text.extend([
            "### 若Shared和路由专家都可任意跨核拆分，最多还能快多少\n",
            "再放宽整专家归属限制，允许任意拆Shared和路由专家，甚至假设全部padding消失。在相同12,288 MAC/周期、II=1、L=21、1GHz下，任何程序仍至少需要 `ceil(有效MAC总数/12288)+20` 周期。这个下界故意给予调度最有利的条件，用来限制剩余空间，不代表已经实现的执行程序。\n",
            _table(["参照冻结设计", "七batch等权计算 ms", "无归属限制下界 ms",
                    "理论最大加速比", "最多可降低延迟 %"],
                   [[row["family"], row["frozen_design_compute_ms"],
                     row["same_L21_II1_lower_bound_ms"], row["best_possible_speedup_at_same_peak"],
                     row["maximum_possible_latency_reduction_pct"]]
                    for row in pooled["global_ceiling"] if row["batch"] == "all"], 6),
            "因此，Shared共同使用两核可能有帮助，但当前真实窗口、同计算峰值下，连消除所有浪费的理想程序也无法相对形状最优单核或最优同构达到5%延迟收益；这并非HBM导致。完整各batch下界见工作簿“无限拆分算力上限”。该限制不外推到改变精度、计算峰值、时钟或算术流水的架构。\n",
        ])
    text.append("### 跨专家交织的存储与供数要求\n")
    store_rows = []
    for batch in BATCHES:
        store_rows.append([f"B{batch}", *[by[(label, batch)]["FP32_partial_peak_KiB"] for label in labels]])
    text.append(_table(["Batch", "原始单核部分和峰值 KiB", "新单核部分和峰值 KiB",
                        "新同构部分和峰值 KiB", "新异构部分和峰值 KiB"], store_rows, 2))
    raw = pooled["per_window"]
    descriptors = [sum(r.get("installed_record_descriptors_per_core", [])) for r in raw]
    text.append(f"事件重放需预创建 {min(descriptors):,}～{max(descriptors):,} 个输出记录描述符；上述部分和容量为实际选定RR/串行回退程序的逐核峰值之和、再取窗口最大值。描述符元数据容量、调度端口和切换电路尚未计入，不能宣称这是几个寄存器即可实现的控制器。\n")
    if pooled["supply"]:
        supply_rows = []
        for r in pooled["supply"]:
            if r["batch"] != "all":
                continue
            supply_rows.append([r["family"],
                r.get("cold_external_unique_weight_necessary_average_GB_s_GM"),
                r.get("array_weight_full_slice_average_GB_s_GM"),
                r.get("array_X_full_slice_average_GB_s_GM"),
                r.get("FP32_partial_read_full_record_average_GB_s_GM"),
                r.get("FP32_partial_write_full_record_average_GB_s_GM"),
                r.get("pool_to_local_weight_refill_GB_s")])
        text.append(_table(["组织", "冷加载所需平均 GB/s", "阵列W交付 GB/s", "阵列X交付 GB/s",
                            "部分和读 GB/s", "部分和写 GB/s", "共享池→本地实际复用带宽"], supply_rows, 2))
        supply_by = {(r["family"], r["batch"]): r for r in pooled["supply"]}
        demand = []
        for batch in BATCHES:
            entries = [supply_by[(family, batch)] for family in FAMILIES]
            demand.append([f"B{batch}", entries[0]["selected_unique_bf16_weight_MiB_AM"],
                *[r["cold_external_unique_weight_necessary_average_GB_s_GM"] for r in entries],
                *[r["fully_resident_selected_weight_operand_and_context_bytes_conservative_max_window"] / 2**20
                  for r in entries]])
        text.append(_table(["Batch", "选中权重AM MiB", "单核冷加载GB/s", "同构冷加载GB/s", "异构冷加载GB/s",
                            "单核选中权重+工作区峰值MiB", "同构选中权重+工作区峰值MiB", "异构选中权重+工作区峰值MiB"], demand, 2))
        text.append("这一组独立按RR记录依赖重新计算不变的有效权重字节、逐发射W/X交付、部分和读写和存活容量；共享池到本地的复用带宽仍待实际RR供数程序验证。不能直接沿用前面K-major OS循环的W/X重用率或供数值。完整逐窗口和分batch需求见pooled/reverse_supply_per_window.csv及pooled/reverse_supply_by_batch.csv。\n")
        text.append("工作区容量采用同时保留所有专家X/Z、选定程序部分和上下文、本地W/X双缓冲及合并输出的保守分配；不是最小容量证明，也不保证两槽即可支撑任意交织。描述符、元数据、banking与流水线内部状态另算。表中需要外部平均GB/s只有在每次冷加载时成立；预驻留可把重放的权重HBM字节变为0，但共享池交付吞吐还必须验证。\n")
    else:
        text.append("独立RR供数反推尚未完成；此处只报告事件结果已记录的部分和峰值及描述符数量，不移用主表OS循环的容量或带宽。\n")
    text.append("充分的驻留容量不会自动保证交付带宽。跨专家交织需要可接受的上下文容量、就绪检测/发射机制、SRAM/互连端口以及真正满足II=1的电路。当前结果能判断理想算力空间，不能据此保证形状异构或大小算力异构在实际芯片上胜出。\n")
    return "\n".join(text)


def _markdown(
    out: Path, contract: dict[str, Any], completion: dict[str, Any], frozen: dict[str, Any],
    config_rows: list[dict[str, Any]], scope: list[dict[str, Any]], compute: list[dict[str, Any]],
    comparisons: list[dict[str, Any]], rates: list[dict[str, Any]], capacity: list[dict[str, Any]],
    certificates: list[dict[str, Any]], timing: list[dict[str, Any]], qwen: list[dict[str, Any]],
    split: list[dict[str, Any]], ready_pool: list[dict[str, Any]] | None = None,
    pooled: dict[str, Any] | None = None,
) -> str:
    by = {(row["label"], row["batch"]): row for row in compute}
    comp = {(row["candidate"], row["baseline"], row["batch"]): row for row in comparisons}
    rate_by = {(row["label"], row["batch"]): row for row in rates}
    cap_by = {(row["label"], row["batch"]): row for row in capacity}
    overall = comp.get(("best_heterogeneous", "best_single", "all"))
    overall_hom = comp.get(("best_heterogeneous", "best_homogeneous", "all"))
    def result_phrase(row: dict[str, Any] | None) -> str:
        if row is None:
            return "未提供配对结果"
        reduction = row["latency_reduction_pct"]
        return (f"快 {reduction:.2f}%" if reduction >= 0 else f"慢 {-reduction:.2f}%")
    text = [
        "# MoE 纯计算几何搜索与供数需求\n",
        "本报告把两套执行契约分开：原主表每个核依次执行整专家；pooled补充主结果允许跨专家就绪输出交织，两套都排除供数与控制等待。各契约使用各自公平搜索并冻结的基线；供数容量和带宽分别反推，不混合比较。\n",
        f"在整专家依次执行契约、主时序下，开发集选定的异构设计相对 DSE 最优单核{result_phrase(overall)}，相对 DSE 最优同构{result_phrase(overall_hom)}。这两个比例只属于该执行契约，不能直接推广到无限上下文的更强执行机制。比较采用七个 batch 等权的配对几何平均，硬件在评估前冻结。\n",
        _pooled_section(pooled),
        "## 实验范围\n",
        _table(["项目", "设置"], [[row["项目"], row["设置"]] for row in scope]),
        "整层计算先把同一专家的 Gate/Up 输出记录交织执行，再执行 Down。每个后续 K 段至少等待前段提交；输出记录数量充足时，独立输出使流水线保持每周期一次发射。单投影 q 个输出记录、s 个 K 段、完成延迟 L 的周期为 `(s−1)×max(q,L)+(q−1)+L`。这里的 L=21，不把 tile 波次直接当作完成周期。\n",
        "## 冻结硬件配置\n",
        _table(["组织", "每核 M×N×K", "每核乘法器", "算力比例", "选择口径"],
               [[row["organization"], row["physical_M_N_K"], row["per_core_multipliers"], row["MAC_ratio"], row["selection"]] for row in config_rows]),
        "M×N×K 表示物理一次发射的 token 行、输出列、点积深度；与专家完整 Me×N×K 不同。大小核按乘法器总数定义，不由 M 高度单独定义。每种设计只有一套冻结形状用于全部 batch；不是按测试窗口换配置。\n",
        "## 主时序的计算延迟\n",
    ]
    main_rows = []
    for batch in (*BATCHES, "all"):
        entry = [f"B{batch}" if batch != "all" else "七 batch 等权 GM"]
        entry += [by.get((label, batch), {}).get("compute_ms_geomean") for label in MAIN_LABELS]
        cs = comp.get(("best_heterogeneous", "best_single", batch), {})
        ch = comp.get(("best_heterogeneous", "best_homogeneous", batch), {})
        entry += [cs.get("speedup"), ch.get("speedup")]
        main_rows.append(entry)
    text += [_table(["Batch", "原始单核 ms", "最优单核 ms", "最优同构 ms", "最优异构 ms", "异构/最优单核加速比", "异构/最优同构加速比"], main_rows, 6),
             "加速比=基线延迟÷异构延迟，>1 表示异构更快。七 batch 等权 GM 是先在每个 batch 内取窗口几何平均，再对七个 batch 的结果取几何平均；它不等于把 135 个窗口直接混合。\n",
             "## 有效计算和浪费\n"]
    all_compute = []
    for label in ORDER:
        row = by.get((label, "all"))
        if row:
            r = rate_by.get((label, "all"), {})
            all_compute.append([LABELS[label], row["compute_ms_geomean"], r.get("useful_GFLOP_s_GM"),
                                row["spatial_utilization"] * 100, row["wall_mac_utilization"] * 100,
                                row["FP32_partial_peak_KiB"], row["finish_gap_cycles_mean"]])
    text += [_table(["组织", "计算 ms", "有效 GFLOP/s", "发射空间利用率 %", "全阵列时间利用率 %", "部分和容量峰值 KiB", "双核完成差均值 周期"], all_compute, 4),
             "发射空间利用率只统计有效 MAC÷发射 MAC；全阵列时间利用率还包含核空闲和流水线等待。聚合利用率按有效 MAC 总量/对应容量总量统计，容量列取所有窗口的最大值；不与等 batch 权重的延迟 GM 混为一个统计口径。\n",
             "## 配对统计与时序稳健性\n"]
    statistical_rows = []
    for candidate in ("best_heterogeneous", "robust_heterogeneous"):
        for baseline in ("original_single", "best_single", "best_homogeneous"):
            row = comp.get((candidate, baseline, "all"))
            if row:
                statistical_rows.append([LABELS[candidate], LABELS[baseline], row["speedup"],
                                         row["latency_reduction_pct"],
                                         f"[{row['speedup_CI95_low']:.5f}, {row['speedup_CI95_high']:.5f}]"])
    text += [_table(["候选", "基线", "加速比", "延迟减少 %", "窗口 bootstrap 95% 区间"], statistical_rows, 5),
             "置信区间仅表示在 batch 内重采样窗口的抽样波动，不包含计算模型误差、硬件时序误差及历史留出复用产生的选择偏差。未原生校准前不能宣布真实硬件胜出。\n"]
    sensitivity_rows = []
    for row in timing:
        if row["batch"] == "all":
            sensitivity_rows.append([row["profile"], LABELS.get(row["label"], row["label"]), row["geometry"],
                                     row["compute_ms_geomean"], row["spatial_utilization"] * 100])
    text += [_table(["时序假设", "冻结设计", "M×N×K", "七 batch 等权计算 ms", "空间利用率 %"], sensitivity_rows, 6),
             "flat20 继承 PK=512 时 dot20 的研究假设，并固定全部 K 的 dot 延迟为20；log_stage1/log_stage2 以 PK512 为同一锚点，每增加一层加法树分别加1/2周期，非二次幂 K 按 ceil(log2 K) 计层数。稳健设计在开发集上最小化三个时序假设下的最大相对遗憾；这是时序假设稳健性，不是已经证明的物理频率一致性。\n",
             "## 外部冷加载需要多少带宽\n",
             "反推带宽=这层选中专家的 BF16 唯一权重字节÷纯计算延迟。它回答在每次都冷加载权重时，至少需要多少平均供数吞吐；它不是本轮测到的 HBM 带宽。若全部权重持续驻留共享池，本次重放的权重 HBM 字节为0，但首次加载和跨层换入仍需另算。\n"]
    supply_main = []
    for batch in BATCHES:
        rr = [rate_by.get((label, batch), {}) for label in MAIN_LABELS[1:]]
        cc = cap_by.get(("best_single", batch), {})
        am_bytes = cc.get("selected_unique_bf16_weight_bytes_AM")
        supply_main.append([f"B{batch}", am_bytes/2**20 if am_bytes is not None else None,
                            *[row.get("cold_weight_average_required_GB_s_GM") for row in rr],
                            *[row.get("cold_weight_average_required_GB_s_max") for row in rr]])
    text += [_table(["Batch", "选中权重 AM MiB", "单核冷加载平均 GB/s", "同构冷加载平均 GB/s", "异构冷加载平均 GB/s", "单核最大窗口 GB/s", "同构最大窗口 GB/s", "异构最大窗口 GB/s"], supply_main, 2),
             "其中“平均 GB/s”是逐窗口字节/时间比值的几何平均；最大窗口用于观察更严格的供数需求。冷加载要求很高时，合理解释是需要权重预驻留的高带宽共享池，而不是宣称现有 HBM 可以达到该速度。\n",
             "## 共享池到核和核内端口需求\n"]
    onchip_rows = []
    for label in MAIN_LABELS[1:]:
        row = rate_by.get((label, "all"), {})
        onchip_rows.append([LABELS[label], row.get("pool_weight_full_slice_average_GB_s_GM"),
                            row.get("pool_x_full_slice_average_GB_s_GM"),
                            row.get("array_weight_full_slice_average_GB_s_GM"),
                            row.get("array_x_full_slice_average_GB_s_GM"),
                            row.get("array_partial_read_full_record_average_GB_s_GM"),
                            row.get("array_partial_write_full_record_average_GB_s_GM")])
    text += [_table(["组织", "共享池 W GB/s", "共享池 X/Z GB/s", "阵列 W GB/s", "阵列 X GB/s", "部分和读 GB/s", "部分和写 GB/s"], onchip_rows, 2),
             "共享池只在复用期开始搬入，阵列操作数表示每次发射的交付量，驻留寄存器可减少真正 SRAM 读取。表中 W、X、部分和流量属于不同接口，不能直接把它们相加当作一个 HBM 带宽。物理完整 slice 含尾部补齐；masked 字节是可以抑制无效地址时的下界，逐窗口 CSV 同时保留。\n"]
    envelope_rows = []
    for label in MAIN_LABELS[1:]:
        row = rate_by.get((label, "all"), {})
        envelope_rows.append([LABELS[label], row.get("simultaneous_pool_weight_GB_s_envelope_with_M_reuse_max"),
                             row.get("simultaneous_pool_X_panel_steady_state_GB_s_envelope_max"),
                             row.get("simultaneous_array_weight_full_GB_s_upper_envelope_max"),
                             row.get("simultaneous_array_X_full_GB_s_upper_envelope_max"),
                             row.get("simultaneous_partial_read_full_GB_s_upper_envelope_max"),
                             row.get("simultaneous_partial_write_full_GB_s_upper_envelope_max")])
    text += [_table(["组织", "共享池 W 包络 GB/s", "共享池 X 稳态包络 GB/s", "阵列 W 包络 GB/s", "阵列 X 包络 GB/s", "部分和读包络 GB/s", "部分和写包络 GB/s"], envelope_rows, 2),
             "这些包络是同时活跃核最大每发射需求的和，供端口初步设计使用；不是实际峰值测量，也不是最小共享池带宽证明。第一组 X panel 必须提前加载；若内存延迟为 L ns、目标吞吐为 B GB/s，在途 payload 至少 ceil(B×L) B，32 B 请求信用至少 ceil(B×L/32)，还需消费中的存活数据和落地空间。双缓冲本身不保证隐藏任意延迟。\n",
             "## 需要多少容量\n"]
    capacity_rows = []
    for batch in BATCHES:
        for label in MAIN_LABELS[1:]:
            row = cap_by.get((label, batch), {})
            if not row:
                continue
            def div(field: str, denominator: int) -> Any:
                value = row.get(f"{field}_max")
                return value / denominator if value is not None else None
            capacity_rows.append([f"B{batch}", LABELS[label], div("sum_core_peak_partial_records_full_bytes", 1024),
                                  div("sum_core_x_k_panel_double_buffer_bytes", 1024),
                                  div("sum_core_weight_tile_double_buffer_bytes", 1024),
                                  div("sum_one_expert_Z_per_core_bytes", 1024),
                                  div("fully_resident_selected_weight_os_working_capacity_bytes", 2**20)])
    text += [_table(["Batch", "组织", "部分和峰值 KiB", "X panel 双缓冲 KiB", "W tile 双缓冲 KiB", "每核一个专家 Z 合计 KiB", "全部选中权重驻留加工作区 MiB"], capacity_rows, 2),
             "容量取各窗口的最大值，并保守合计各核峰值。它包含物理补齐的全部输出部分和、完整 X panel 双缓冲、W tile 双缓冲、共享 X/Y 和每核当前专家 Z；不包含 SRAM banking、元数据、对齐和算术流水线内部寄存器。全部 Gate/Up 部分和保留到该专家 GU 完成后才产生 Z，不能凭循环名字省去这部分容量。\n",
             _table(["驻留范围", "BF16 容量", "包含什么"], [
                 ["一个 routed expert", "16.5 MiB", "Gate、Up、Down 三张权重"],
                 ["Shared 合并专家", "33 MiB", "合并中间宽度2816的三张权重"],
                 ["一整层全部专家", "1,089 MiB ≈ 1.0635 GiB", "64 routed expert + Shared；工作区另加"],
                 ["26 个 MoE 层全部专家", "27.6504 GiB", "只包含MoE FFN权重，不含attention、dense、router等"],
             ]),
             "足以装下一层的缓存不等于能永久装下全模型。即便1 GB级共享池解决了这层的容量，也必须验证数值表中的共享池读取、互连和本地端口吞吐。容量不能由 GB/s 直接推出，带宽也不能由 GB 容量直接推出。\n",
             "## Shared 同时使用两个核的补充诊断\n",
             _table(["Batch", "窗口", "组织", "整专家最优周期", "Shared 拆列可行周期", "Gate/Up 给核0的列数", "Down 给核0的列数"],
                    [[f"B{row['batch']}", row["window_id"], LABELS.get(row["label"], row["label"]),
                      row["whole_expert_optimal_cycles"], row["cycles"], row.get("GateUp_columns_core0"), row.get("Down_columns_core0")] for row in split], 0),
             "拆列诊断对同构、异构使用相同的4列切分粒度，保持每个输出全部 K 段在原核累加。它只提供每个 batch 的一个预定义代表窗口的 Shared-first 可行调度；不证明所有窗口的拆分最优，也不能代替主表整专家契约的全域证书。诊断更慢时，可以回退到整专家调度。\n",
             "## 冻结硬件在另一模型上的补充评估\n"]
    qwen_by = {(row["label"], row["batch"]): row for row in qwen}
    qwen_rows = []
    for batch in (*BATCHES, "all"):
        if ("best_single", batch) in qwen_by:
            qwen_rows.append([f"B{batch}" if batch != "all" else "七 batch 等权 GM",
                              qwen_by[("best_single", batch)]["windows"],
                              *[qwen_by.get((label, batch), {}).get("compute_ms_geomean") for label in MAIN_LABELS]])
    text += [_table(["Batch", "窗口数", "原始单核 ms", "DS冻结单核 ms", "DS冻结同构 ms", "DS冻结异构 ms"], qwen_rows, 6),
             "Qwen3-30B-A3B 使用真实单token路由结果进行微批重放：H2048、F768、128专家、top-K8，无Shared。请求按来源ID分开，随后按捕获样本顺序合成不同batch；没有新执行一次批量模型推理。测试沿用DS开发集冻结几何，不对Qwen重新选择。不同batch共享部分源请求，因此窗口有相关性。\n",
             _ready_pool_section(ready_pool or []),
             "## 搜索域和最优性证书\n",
             _table(["时序", "N对齐", "组织", "域内几何数", "精确求解点", "有效下界剪枝点", "开发集选择", "域内已认证"],
                    [[row["profile"], row["N_alignment"], row["family"], row["domain_geometries"],
                      row["exact_points"], row["pruned_by_valid_LB"], row["selected"],
                      row["within_declared_domain_certified"]] for row in certificates], 0),
             "所有配置具有相同12,288乘法器。每个候选先计算允许任务分数化的合法下界；只有可能优于当前结果的候选才进行整专家整数分配精确求解。被合法下界剪掉的配置仍属于已覆盖的搜索域。‘最优’只针对明确的BF16时序、循环、整专家归属和尺寸范围；不包含动态可重构阵列、任意多核、任意层流水或全部可能编译器循环。\n",
             "## 可以据此设计什么\n",
             "本轮可以确定：在统一计算模型下哪些形状更有效、哪些batch能从异构受益，以及维持该速度的外部平均供数、共享池读取、阵列交付和工作区容量。下一步应按胜出点实际需求验证共享SRAM/3D互连端口、有限预取队列及动态派工，再综合时序与面积；不能把纯计算优势直接写成当前HBM系统或完整模型的加速。\n",
             "完整逐窗口数据见 `reverse_supply.csv`，分batch供数见 `required_bandwidth_by_batch.csv`，容量峰值见 `capacity_by_batch.csv`。工作簿包含全部冻结结果、比较、时序敏感性和来源清单。\n",
    ]
    return "\n".join(text)


def _workbook(path: Path, tables: dict[str, list[dict[str, Any]]], summary: str) -> dict[str, Any]:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    workbook = Workbook()
    workbook.remove(workbook.active)
    cover = workbook.create_sheet("阅读说明")
    notes = [
        ["MoE纯计算DSE与供数反推", "单位和口径"],
        ["主数据", "本轮完成结果；不混入旧表；纯计算排除HBM/SRAM等待"],
        ["尺寸", "M×N×K是物理tile尺寸；每个组织总乘法器12288"],
        ["延迟", "毫秒；1GHz；七batch等权GM；不是完整FFN/模型E2E"],
        ["带宽", "十进制GB/s；cold带宽是需求反推，非测到的HBM吞吐"],
        ["容量", "原始字节字段B；MiB=2^20B；KiB=2^10B；峰值不含metadata/银行/流水内部寄存器"],
        ["加速比", "基线延迟÷候选延迟；>1更快；小数利用率按百分比显示"],
        ["验证", "旧heldout已反复使用；bootstrap不包含模型误差；未原生/RTL校准"],
        ["详细讲解", "SUMMARY.md"],
    ]
    for row in notes:
        cover.append(row)
    for name, rows in tables.items():
        sheet = workbook.create_sheet(name[:31])
        if not rows:
            sheet.append(["此项没有对应完成结果"])
            continue
        fields = list(dict.fromkeys(key for row in rows for key in row))
        sheet.append(fields)
        for row in rows:
            sheet.append([json.dumps(row.get(field), ensure_ascii=False, sort_keys=True) if isinstance(row.get(field), (dict, list)) else row.get(field) for field in fields])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column, field in enumerate(fields, 1):
            sheet.column_dimensions[get_column_letter(column)].width = min(56, max(14, len(field) * 0.85))
            if field in ("spatial_utilization", "wall_mac_utilization"):
                for cells in sheet.iter_rows(min_row=2, min_col=column, max_col=column):
                    cells[0].number_format = "0.00%"
            else:
                for cells in sheet.iter_rows(min_row=2, min_col=column, max_col=column):
                    if isinstance(cells[0].value, float):
                        cells[0].number_format = "0.000000"
    for sheet in workbook:
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="17365D")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.row_dimensions[1].height = 36
    cover.column_dimensions["A"].width = 24
    cover.column_dimensions["B"].width = 112
    for row in cover:
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    workbook.save(path)
    readback = load_workbook(path, read_only=True, data_only=True)
    checked = 0
    for name, rows in tables.items():
        sheet = readback[name[:31]]
        if not rows:
            continue
        fields = list(dict.fromkeys(key for row in rows for key in row))
        data = sheet.iter_rows(values_only=True)
        assert list(next(data)) == fields
        for original, values in zip(rows, data):
            for field, actual in zip(fields, values):
                expected = original.get(field)
                if isinstance(expected, (dict, list)):
                    expected = json.dumps(expected, ensure_ascii=False, sort_keys=True)
                if isinstance(expected, float):
                    assert actual is not None and math.isclose(expected, actual, rel_tol=1e-13, abs_tol=1e-10)
                else:
                    assert actual == expected, (name, field, actual, expected)
                checked += 1
        assert sheet.max_row == len(rows) + 1
    sheets = len(readback.sheetnames)
    readback.close()
    return {"workbook_sheets": sheets, "cells_readback_checked": checked}


def generate_report(out: str | Path) -> dict[str, Any]:
    """Create report artifacts from completed frozen outputs, without rerunning."""
    out = Path(out)
    completion = _json(out / "COMPLETION.json")
    contract = _json(out / "CONTRACT.json")
    frozen = _json(out / "FROZEN_SELECTION.json")
    compute = _csv(out / "compute_by_batch.csv")
    comparisons = _csv(out / "comparisons.csv")
    certificates = _csv(out / "search_certificates.csv")
    sensitivity = _csv(out / "timing_sensitivity.csv")
    qwen = _csv(out / "qwen_frozen_hardware_by_batch.csv")
    shared_split = _csv(out / "shared_split_diagnostic.csv")
    ready_pool = []
    ready_receipt = None
    if (out / "ready_pool_diagnostic.csv").exists() and (out / "READY_POOL_DIAGNOSTIC.json").exists():
        ready_pool = _csv(out / "ready_pool_diagnostic.csv")
        ready_receipt = _json(out / "READY_POOL_DIAGNOSTIC.json")
        if ready_receipt.get("case_count") != len(ready_pool):
            raise ValueError("Ready-pool completion receipt disagrees with CSV case count")
        if any(row.get("two_repeats_identical") is not True for row in ready_pool):
            raise ValueError("Ready-pool diagnostic has not passed its repeated replay checks")
    if not compute or not comparisons or not certificates:
        raise ValueError("Completion receipt exists but required completed CSV tables are missing")
    pooled = _load_pooled(out)
    source_paths = [out / name for name in (
        "COMPLETION.json", "CONTRACT.json", "FROZEN_SELECTION.json", "compute_by_batch.csv",
        "comparisons.csv", "search_certificates.csv", "timing_sensitivity.csv",
        "qwen_frozen_hardware_by_batch.csv", "shared_split_diagnostic.csv", "QWEN_PROVENANCE.json",
    ) if (out / name).exists()]
    if ready_receipt is not None:
        source_paths.extend((out / "ready_pool_diagnostic.csv", out / "READY_POOL_DIAGNOSTIC.json"))
    if pooled is not None:
        source_paths.extend(pooled["source_paths"])
    source_hashes = {str(path): hashlib.sha256(path.read_bytes()).hexdigest() for path in source_paths}
    supplies: list[dict[str, Any]] = []
    core_rows: list[dict[str, Any]] = []
    compute_rows: list[dict[str, Any]] = []
    loop_rows: list[dict[str, Any]] = []
    for label in ORDER:
        path = out / f"supply_{label}.json"
        raw_path = out / f"heldout_{label}.json"
        if not path.exists() or not raw_path.exists():
            continue
        raw = _json(raw_path)
        raw_by = {row["window_id"]: row for row in raw}
        if len(raw_by) != len(raw):
            raise ValueError("Duplicate per-window compute IDs")
        source_paths.extend((path, raw_path))
        source_hashes[str(path)] = hashlib.sha256(path.read_bytes()).hexdigest()
        source_hashes[str(raw_path)] = hashlib.sha256(raw_path.read_bytes()).hexdigest()
        source = _json(path)
        if {row["window_id"] for row in source} != set(raw_by):
            raise ValueError("Supply/compute window IDs do not match")
        for row in source:
            flat = _flatten_supply(row, raw_by[row["window_id"]])
            supplies.append(flat)
            cr = raw_by[row["window_id"]]
            compute_rows.append({"label": label, **{key: value for key, value in cr.items() if key != "assignments"},
                                 "assignments_json": json.dumps(cr["assignments"], ensure_ascii=False, sort_keys=True)})
            for core in row.get("per_core", []):
                core_rows.append({"label": label, "window_id": row["window_id"], "batch": row["batch"],
                                  "shape_M_N_K": _shape(core["shape"]),
                                  **{key: value for key, value in core.items() if key != "shape"}})
            for expert in row.get("per_expert", []):
                for projection in expert.get("projections", []):
                    loop_rows.append({"label": label, "window_id": row["window_id"], "batch": row["batch"],
                                      "expert_id": expert["expert_id"], "owner": expert["owner"],
                                      "phase": projection["phase"], "M": projection["M"], "N": projection["N"], "K": projection["K"],
                                      "loop_alternatives_json": json.dumps(projection.get("loop_alternatives", {}), ensure_ascii=False, sort_keys=True)})
    if not supplies:
        raise ValueError("No completed reverse-supply JSON files found")
    rates, capacities = _aggregate_supply(supplies)
    configs = _config_rows(frozen)
    scope = _scope_rows(contract, completion)
    _write_csv(out / "reverse_supply.csv", supplies)
    _write_csv(out / "required_bandwidth_by_batch.csv", rates)
    _write_csv(out / "capacity_by_batch.csv", capacities)
    _write_csv(out / "frozen_configs_readable.csv", configs)
    summary = _markdown(out, contract, completion, frozen, configs, scope, compute, comparisons,
                        rates, capacities, certificates, sensitivity, qwen, shared_split, ready_pool, pooled)
    (out / "SUMMARY.md").write_text(summary)
    source_rows = [{"path": path, "sha256": digest} for path, digest in sorted(source_hashes.items())]
    tables = {
        "实验范围": scope, "冻结配置": configs, "计算分batch": compute,
        "配对加速与置信区间": comparisons, "外部与片上带宽": rates,
        "容量分batch": capacities, "逐窗口反推供数": supplies,
        "逐窗口计算与分配": compute_rows, "每核接口和容量": core_rows,
        "时序敏感性": sensitivity, "Shared拆列诊断": shared_split,
        "Qwen冻结硬件结果": qwen, "精确搜索证书": certificates,
        "数据流备选流量": loop_rows, "结果来源": source_rows,
    }
    if ready_pool:
        tables["跨专家交织代表窗口"] = ready_pool
    if pooled is not None:
        tables.update({
            "交织全域冻结配置": pooled["configs"],
            "交织全窗口计算": pooled["compute"],
            "交织配对统计": pooled["comparisons"],
            "交织下界证书": pooled["certificates"],
            "交织逐窗口和上下文": pooled["per_window"],
        })
        if pooled["supply"]:
            tables["交织独立供数分batch"] = pooled["supply"]
            tables["交织逐窗口供数和容量"] = pooled["supply_per_window"]
        if pooled.get("global_ceiling"):
            tables["无限拆分算力上限"] = pooled["global_ceiling"]
    workbook_audit = _workbook(out / "MoE_纯计算DSE与供数需求.xlsx", tables, summary)
    assert all(hashlib.sha256(Path(path).read_bytes()).hexdigest() == digest for path, digest in source_hashes.items())
    audit = {
        "status": "PASS", "scope": "completed-result reporting; no fresh simulation/native calibration",
        "source_sha256": source_hashes, "source_files_unchanged": True,
        "reverse_supply_windows": len(supplies), "per_core_supply_rows": len(core_rows),
        "loop_alternative_projection_rows": len(loop_rows),
        "compute_summary_rows": len(compute), "comparison_rows": len(comparisons),
        "ready_pool_diagnostic_cases": len(ready_pool),
        "ready_pool_scope": ready_receipt.get("scope") if ready_receipt else None,
        "ready_pool_supply_not_reused": True if ready_pool else None,
        "pooled_campaign_present": pooled is not None,
        "pooled_source_windows": len(pooled["per_window"]) if pooled else 0,
        "pooled_development_lower_bound_gaps": {
            row["family"]: row["certified_relative_gap"]
            for row in pooled["certificates"]} if pooled else None,
        "pooled_supply_independent_from_serial_OS": True if pooled and pooled["supply"] else None,
        "pooled_supply_window_rows": len(pooled["supply_per_window"]) if pooled else 0,
        "global_MAC_ceiling_rows": len(pooled.get("global_ceiling", [])) if pooled else 0,
        "global_MAC_ceiling_is_not_a_schedule": True if pooled and pooled.get("global_ceiling") else None,
        "all_GM_weighting": "equal weight per observed batch; primary all has exactly seven batches",
        "external_bandwidth_is_reverse_requirement_not_measurement": True,
        "native_calibrated": completion.get("native_calibrated", False),
        **workbook_audit,
    }
    (out / "REPORT_AUDIT.json").write_text(json.dumps(audit, indent=2, sort_keys=True) + "\n")
    return audit


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    print(json.dumps(generate_report(args.out), ensure_ascii=False, sort_keys=True))


if __name__ == "__main__":
    main()
