"""Complete frozen-hardware Shared-only cross-dataflow result tables.

Every OS/WS/IS retention combination uses the same ideal-port compute timing.
Native PLENA PE dataflow stays OS: WS/IS here are declared loop/operand-retention
extensions with fully charged grouped PE context spill.  No numeric result is
represented as a calibrated native HBM/RTL measurement.
"""
from __future__ import annotations

import argparse
import csv
import gzip
import hashlib
import itertools
import json
import math
from pathlib import Path
from typing import Any

from .dataflow import FLOWS, add_rates, plena_shared_traffic
from .geometry import BUDGET, Core
from .timing import BATCHES, F, H, PROFILES, partition, regions, resource_ledger

PROFILE = PROFILES[0]
MIB = 1024 ** 2
KIB = 1024


def gm(values):
    values = list(values)
    if not values:
        raise ValueError("cannot aggregate an empty collection")
    if any(value < 0 for value in values):
        raise ValueError("geometric average requires nonnegative values")
    return 0.0 if any(value == 0 for value in values) else math.exp(sum(math.log(x) for x in values) / len(values))


def canonical(value):
    return json.dumps(value, sort_keys=True, ensure_ascii=False, separators=(",", ":"))


def sha256(path):
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def write_csv(path, rows):
    rows = list(rows)
    if not rows:
        raise ValueError(f"no rows for {path}")
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def write_json(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def physical_string(cores):
    return " + ".join(f"{c.pm}×{c.pn}×{c.groups}" for c in cores)


def frame_string(cores):
    return " + ".join(f"b={c.b},R={c.r},rM={c.rm},rN={c.rn}" for c in cores)


def frozen_designs(selection):
    designs = [("paper_original", "native", (Core(32, 2048),), "paper_original")]
    for scope in ("native", "extended"):
        for family in ("single", "homogeneous", "heterogeneous"):
            name = scope + "_" + family
            chosen = selection[name][PROFILE.name]
            cores = tuple(Core(**{key: core[key] for key in ("b", "r", "rm", "rn")})
                          for core in chosen["cores"])
            designs.append((name, scope, cores, PROFILE.name))
    witness = (Core(2, 256, 1, 96), Core(2, 256, 1, 32))
    witness_score = gm(partition(witness, batch, PROFILE)["cycles"] for batch in BATCHES)
    optimum = selection["extended_heterogeneous"][PROFILE.name].get("GM_cycles", witness_score)
    if not math.isclose(witness_score, optimum, rel_tol=1e-12):
        raise AssertionError("3:1 witness does not match frozen heterogeneous optimum")
    designs.append(("extended_3to1_cooptimal_witness", "extended", witness,
                    "verified_3to1_cooptimal_witness"))
    for family in ("single", "homogeneous", "heterogeneous"):
        chosen = selection["extended_" + family]["timing_robust"]
        cores = tuple(Core(**{key: core[key] for key in ("b", "r", "rm", "rn")})
                      for core in chosen["cores"])
        designs.append(("timing_robust_" + family, "extended_robust", cores, "timing_robust"))
    return designs


def point(cores, batch, flows):
    plan = partition(cores, batch, PROFILE)
    owned = {key: [region for region in value
                   if region.m_stop > region.m_start and region.n_stop > region.n_start]
             for key, value in regions(cores, batch, plan).items()}
    traffic = plena_shared_traffic(batch, cores, flows, owned)
    rates = add_rates(traffic, plan["cycles"])
    ledger = resource_ledger(cores)
    if ledger["multipliers"] != BUDGET:
        raise AssertionError("frozen hardware violates the multiplier budget")
    if traffic["useful_macs"] != 3 * batch * H * F:
        raise AssertionError("partition lost useful computation")
    if rates["spatial_utilization_pct"] > 100.000000001:
        raise AssertionError("invalid spatial utilization")
    # Fixed PE, shadow, operand-staging and M_SUM hardware is charged once per
    # chip, even when Gate and Up share it.  Their off-PE scratch contexts can
    # both be live, so sum GU scratch and compare against the Down scratch.
    capacity = traffic["capacity"]
    static_bytes = (ledger["PE_FP32_bytes"] + ledger["shadow_FP32_bytes"]
                    + ledger["full_operand_double_buffer_bytes"]
                    + ledger["M_SUM_tree_FP32_register_bytes"])
    scratch_peak = 0
    for core_id in range(len(cores)):
        own = [item for item in traffic["projections"] if item["core_id"] == core_id]
        gu = sum(item["capacity"]["partial_scratch_full_bytes"]
                 for item in own if item["projection"] in ("gate", "up"))
        down = sum(item["capacity"]["partial_scratch_full_bytes"]
                   for item in own if item["projection"] == "down")
        scratch_peak += max(gu, down)
    local_upper = static_bytes + scratch_peak
    return {"plan": plan, "traffic": traffic, "rates": rates, "ledger": ledger,
            "static_compute_staging_tree_bytes": static_bytes,
            "extra_grouped_partial_scratch_upper_bytes": scratch_peak,
            "local_scratch_tree_upper_bytes": local_upper,
            "resident_shared_payload_upper_bytes": capacity["all_shared_weights_bytes"]
                                                  + capacity["activation_liveness_upper_bytes"] + local_upper}


def result_row(design, scope, source, cores, batch, flows, got):
    t, r, plan, ledger = (got["traffic"], got["rates"], got["plan"], got["ledger"])
    pool, external = t["pool_to_local"], t["external"]
    return {
        "design": design, "scope": scope, "selection_source": source,
        "batch": batch, "row_type": "batch", "physical_PM_PN_G": physical_string(cores),
        "native_frames": frame_string(cores), "multiplier_allocation": "+".join(str(c.macs) for c in cores),
        "dataflow": "/".join(flows), "profile": PROFILE.name,
        "partition": plan["kind"], "gu_split": plan["gu_split"],
        "down_split": plan["down_split"], "row_split": plan["row_split"],
        "compute_cycles": plan["cycles"], "compute_ms": r["compute_ms"],
        "useful_GFLOP_s": r["useful_GFLOP_s"],
        "spatial_utilization_pct": r["spatial_utilization_pct"],
        "wall_peak_utilization_pct": 100 * t["useful_macs"] / (plan["cycles"] * BUDGET),
        "unique_weight_MiB": external["cold_weight_only_unique_bytes"] / MIB,
        "unique_weight_required_GB_s": r["unique_weight_required_GB_s"],
        "unique_weight_fetch_ms_at256": r["unique_weight_fetch_ms"],
        "cold_all_required_GB_s": r["cold_all_required_GB_s"],
        "cold_all_fetch_ms_at256": r["cold_all_fetch_ms"],
        "pool_weight_refill_MiB": pool["weight_masked_bytes"] / MIB,
        "pool_X_refill_MiB": pool["x_masked_bytes"] / MIB,
        "pool_context_transfer_MiB": (pool["partial_read_masked_bytes"]
                                      + pool["partial_write_masked_bytes"]) / MIB,
        "pool_weight_refill_required_GB_s": r["pool_to_local"]["weight_masked_GB_s"],
        "pool_X_refill_required_GB_s": r["pool_to_local"]["x_masked_GB_s"],
        "pool_context_required_GB_s": r["pool_to_local"]["partial_read_masked_GB_s"]
                                      + r["pool_to_local"]["partial_write_masked_GB_s"],
        "array_weight_delivery_required_GB_s": r["array_delivery"]["weight_full_GB_s"],
        "array_X_delivery_required_GB_s": r["array_delivery"]["x_full_GB_s"],
        "unified_pool_external_weight_refetch_factor": external["unified_pool_external_weight_refetch_factor"],
        "no_pool_weight_refetch_factor": external["no_pool_stream_weight_refetch_factor"],
        "no_pool_stream_total_MiB": external["no_pool_stream_total_bytes"] / MIB,
        "no_pool_stream_required_GB_s": r["no_pool_stream_required_GB_s"],
        "no_pool_stream_fetch_ms_at256": r["no_pool_stream_fetch_ms"],
        "local_scratch_tree_upper_KiB": got["local_scratch_tree_upper_bytes"] / KIB,
        "resident_shared_payload_upper_MiB": got["resident_shared_payload_upper_bytes"] / MIB,
        "PE_partial_KiB": ledger["PE_FP32_bytes"] / KIB,
        "shadow_partial_KiB": ledger["shadow_FP32_bytes"] / KIB,
        "M_SUM_adders": ledger["M_SUM_adders"],
        "M_SUM_tree_register_KiB": ledger["M_SUM_tree_FP32_register_bytes"] / KIB,
        "repeat_identical": True,
    }


def gm_row(rows):
    first = rows[0]
    result = dict(first)
    result.update(batch="7-Batch GM", row_type="GM", partition="adaptive rows/columns",
                  gu_split=None, down_split=None, row_split=None)
    capacities = {"local_scratch_tree_upper_KiB", "resident_shared_payload_upper_MiB"}
    metadata = {"batch", "row_type", "gu_split", "down_split", "row_split"}
    for key, value in first.items():
        if key in metadata or isinstance(value, bool):
            continue
        if isinstance(value, (int, float)):
            result[key] = max(row[key] for row in rows) if key in capacities else gm(row[key] for row in rows)
    return result


def format_cell(value, precision=6):
    if value is None:
        return "—"
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, float):
        return f"{value:,.{precision}f}"
    return str(value).replace("|", "\\|")


def markdown_table(rows, columns):
    lines = ["| " + " | ".join(label for _, label in columns) + " |",
             "| " + " | ".join("---" for _ in columns) + " |"]
    for row in rows:
        lines.append("| " + " | ".join(format_cell(row[key]) for key, _ in columns) + " |")
    return "\n".join(lines)


def crossover_rows(designs):
    output = []
    samples = sorted(set(BATCHES) | {256, 512, 1024, 2048, 4096})
    for design, scope, cores, source in designs:
        # Unique-weight service is constant, so compute crossing is monotone
        # under these row/column partition contracts and can be located exactly.
        target_cycles = 33 * MIB / 256
        lower, upper = 2, 4096
        if partition(cores, upper, PROFILE)["cycles"] < target_cycles:
            unique_cross = None
        else:
            while lower < upper:
                mid = (lower + upper) // 2
                if partition(cores, mid, PROFILE)["cycles"] >= target_cycles:
                    upper = mid
                else:
                    lower = mid + 1
            unique_cross = lower
        for flows in itertools.product(FLOWS, repeat=len(cores)):
            points = [(batch, point(cores, batch, flows)) for batch in samples]
            hits = [batch for batch, got in points
                    if got["rates"]["compute_ms"] >= got["rates"]["no_pool_stream_fetch_ms"]]
            first_stream = min(hits) if hits else None
            for batch, got in points:
                r = got["rates"]
                output.append({
                    "design": design, "scope": scope, "selection_source": source,
                    "dataflow": "/".join(flows), "batch": batch,
                    "compute_ms": r["compute_ms"],
                    "unique_weight_fetch_ms_at256": r["unique_weight_fetch_ms"],
                    "unique_weight_compute_service_larger": r["compute_ms"] >= r["unique_weight_fetch_ms"],
                    "exact_first_unique_cross_batch_in_2_4096": unique_cross,
                    "no_pool_stream_fetch_ms_at256": r["no_pool_stream_fetch_ms"],
                    "no_pool_compute_service_larger": r["compute_ms"] >= r["no_pool_stream_fetch_ms"],
                    "first_no_pool_cross_on_reported_sample_grid": first_stream,
                    "no_pool_cross_not_exhaustive": True,
                    "not_full_E2E_bound_classification": True,
                })
    return output


def make_workbook(out, tables):
    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        raise SystemExit("Workbook requires openpyxl; rerun using /tmp/plena-report-venv/bin/python")
    path = out / "PLENA65536_Shared_完整DSE与数据流.xlsx"
    workbook = Workbook()
    workbook.remove(workbook.active)
    about = workbook.create_sheet("阅读说明")
    notes = [
        ["口径", "说明"],
        ["主时序", "论文拓扑一致解析模型，1GHz，非原生Rust/RTL已校准结果"],
        ["物理第三维", "G=R/b为并行K组数；R为逻辑K块宽度，b为每段时序步数"],
        ["数据流", "PE始终OS；WS/IS是操作数循环驻留扩展，收费完整G部分和上下文"],
        ["Cross Dataflow", "同硬件理想端口计算时间相同，实际供数和上下文流量不同"],
        ["唯一权重", "Shared三矩阵33MiB；共享池冷取一次，不由循环顺序虚构HBM重读"],
        ["No-pool", "另一个无共享池系统，操作数重读与中间结果及上下文溢写算外部流量"],
        ["Fetch Time", "字节/256GB/s的服务等价，不是原生HBM计时，也不与Compute直接相加"],
        ["GM", "七个Batch等权GM；容量列的GM行报告七个Batch最大值"],
        ["所有运行", "每个主点计算/流量重算两次，规范化结果完全一致"],
        ["交叉点", "唯一W交叉点在2..4096整数Batch精确搜索；无池交叉仅报告列出采样点"],
        ["范围", "Shared-only仅七种不同形状；不将路由窗口复用当作135个独立Shared样本"],
    ]
    for row in notes:
        about.append(row)
    audit_count = 0
    for name, rows in tables.items():
        sheet = workbook.create_sheet(name[:31])
        headers = list(rows[0])
        sheet.append(headers)
        for row in rows:
            sheet.append([row[key] for key in headers])
    for sheet in workbook:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="183B56")
            cell.font = Font(bold=True, color="FFFFFF")
            cell.alignment = Alignment(wrap_text=True)
        sheet.row_dimensions[1].height = 45
        for column in sheet.columns:
            length = max(len(str(cell.value or "")) for cell in column[:100])
            sheet.column_dimensions[column[0].column_letter].width = min(72, max(12, length + 2))
    workbook.save(path)
    saved = load_workbook(path, read_only=True, data_only=True)
    for name, rows in tables.items():
        cells = list(saved[name[:31]].values)
        headers = list(rows[0])
        if list(cells[0]) != headers or len(cells) != len(rows) + 1:
            raise AssertionError("workbook row/header readback mismatch")
        for source, actual in zip(rows, cells[1:]):
            for key, got in zip(headers, actual):
                expected = source[key]
                if isinstance(expected, float):
                    if not math.isclose(expected, got, rel_tol=2e-14, abs_tol=1e-12):
                        raise AssertionError((name, key, expected, got))
                elif expected != got and not (expected == "" and got is None):
                    raise AssertionError((name, key, expected, got))
                audit_count += 1
    saved.close()
    return path, audit_count


def main(out):
    selection_path = out / "FROZEN_SELECTION.json"
    if not selection_path.exists():
        raise SystemExit("Run search first: FROZEN_SELECTION.json is required")
    selection = json.loads(selection_path.read_text())
    designs = frozen_designs(selection)
    all_rows, tradeoff, detailed, resources, sensitivity = [], [], [], [], []
    repeats = 0
    for design, scope, cores, source in designs:
        ledger = resource_ledger(cores)
        resources.append({"design": design, "scope": scope, "physical_PM_PN_G": physical_string(cores),
                          "native_frames": frame_string(cores),
                          "multiplier_allocation": "+".join(str(c.macs) for c in cores), **ledger})
        for profile in PROFILES:
            sensitivity.append({"design": design, "profile": profile.name,
                                "GM_compute_ms": gm(partition(cores, batch, profile)["cycles"] for batch in BATCHES) / 1e6,
                                "physical_PM_PN_G": physical_string(cores), "native_frames": frame_string(cores)})
        for flows in itertools.product(FLOWS, repeat=len(cores)):
            batch_rows = []
            for batch in BATCHES:
                got = point(cores, batch, flows)
                repeat = point(cores, batch, flows)
                if canonical(got) != canonical(repeat):
                    raise AssertionError("repeated point differed")
                repeats += 1
                row = result_row(design, scope, source, cores, batch, flows, got)
                batch_rows.append(row)
                detailed.append({"design": design, "scope": scope, "batch": batch,
                                 "dataflows": flows, "result_sha256": hashlib.sha256(canonical(got).encode()).hexdigest(),
                                 **got})
            combined = gm_row(batch_rows)
            all_rows.extend(batch_rows + [combined])
            tradeoff.append({key: combined[key] for key in (
                "design", "scope", "physical_PM_PN_G", "native_frames", "multiplier_allocation", "dataflow",
                "compute_ms", "useful_GFLOP_s", "unique_weight_required_GB_s", "unique_weight_fetch_ms_at256",
                "pool_weight_refill_required_GB_s", "pool_X_refill_required_GB_s", "pool_context_required_GB_s",
                "no_pool_weight_refetch_factor", "no_pool_stream_required_GB_s", "no_pool_stream_fetch_ms_at256",
                "local_scratch_tree_upper_KiB", "resident_shared_payload_upper_MiB")})
    crossings = crossover_rows(designs)
    write_csv(out / "cross_dataflow.csv", all_rows)
    write_csv(out / "dataflow_tradeoff.csv", tradeoff)
    write_csv(out / "resource_ledger.csv", resources)
    write_csv(out / "frozen_timing_sensitivity.csv", sensitivity)
    write_csv(out / "crossover.csv", crossings)
    supply = [{key: row[key] for key in (
        "design", "scope", "batch", "dataflow", "compute_ms", "unique_weight_required_GB_s",
        "unique_weight_fetch_ms_at256", "pool_weight_refill_required_GB_s", "pool_X_refill_required_GB_s",
        "pool_context_required_GB_s", "array_weight_delivery_required_GB_s", "array_X_delivery_required_GB_s",
        "no_pool_stream_required_GB_s", "no_pool_stream_fetch_ms_at256")} for row in all_rows]
    write_csv(out / "supply_bw.csv", supply)
    with gzip.open(out / "cross_dataflow_details.json.gz", "wt", encoding="utf-8") as handle:
        json.dump(detailed, handle, ensure_ascii=False, separators=(",", ":"))
    columns = [("batch", "Batch"), ("design", "组织"), ("physical_PM_PN_G", "物理PM×PN×G"),
               ("multiplier_allocation", "乘法器分配"), ("dataflow", "驻留循环"),
               ("compute_ms", "Compute ms"), ("useful_GFLOP_s", "有效GFLOP/s"),
               ("unique_weight_required_GB_s", "唯一W所需GB/s"), ("unique_weight_fetch_ms_at256", "Fetch ms @256"),
               ("local_scratch_tree_upper_KiB", "本地/上下文/树容量上界KiB"),
               ("spatial_utilization_pct", "有效/发射MAC利用率%"),
               ("wall_peak_utilization_pct", "含时序峰值利用率%")]
    lines = ["# PLENA 65,536：Shared-only 完整交叉数据流表", "",
             "这是论文拓扑一致的BF16/FP32解析模型，非原生Rust/RTL校准结果。物理第三维G=R/b，不是逻辑归约宽度R。", "",
             "PE本身始终OS。WS/IS表示操作数循环驻留扩展：其G个PE部分和上下文的存取与容量均单独收费。理想端口下同几何计算时间相同，不能把交通差异伪装成已实测的计算加速。", "",
             "Shared权重固定33MiB，共享池冷取一次，256GB/s等价Fetch固定0.135168ms；无池重读是另列的不同系统。GM行的容量取七个Batch最大值，不取GM。", ""]
    for scope in ("native", "extended", "extended_robust"):
        lines += [f"## {scope}：全配置×驻留循环×七Batch及GM", "",
                  markdown_table([row for row in all_rows if row["scope"] == scope], columns), ""]
    lines += ["## 数据流权衡：同冻结几何", "",
              markdown_table(tradeoff, [("design", "组织"), ("dataflow", "驻留循环"),
                  ("compute_ms", "GM Compute ms"), ("pool_weight_refill_required_GB_s", "池W GB/s"),
                  ("pool_X_refill_required_GB_s", "池X GB/s"), ("pool_context_required_GB_s", "池上下文GB/s"),
                  ("no_pool_weight_refetch_factor", "无池W重读GM"),
                  ("local_scratch_tree_upper_KiB", "本地容量上界KiB")]), "",
              "## 硬件资源账本", "", markdown_table(resources, [
                  ("design", "组织"), ("native_frames", "b/R/复制"), ("multipliers", "乘法器"),
                  ("PE_FP32_bytes", "PE部分和B"), ("shadow_FP32_bytes", "影子部分和B"),
                  ("M_SUM_adders", "M_SUM加法器"), ("M_SUM_tree_FP32_register_bytes", "树寄存器B")]), "",
              "交叉点详细表在crossover.csv。唯一权重交叉点在B=2..4096整数域内精确查找；无池只报告表中采样点，未声称找出全域首次翻转。计算/Fetch大小比较是必要服务约束比较，不是带真实HBM延迟与端口冲突的E2E证明。", ""]
    (out / "REPORT.md").write_text("\n".join(lines), encoding="utf-8")
    workbook, cells = make_workbook(out, {
        "Native全数据流": [row for row in all_rows if row["scope"] == "native"],
        "Extended全数据流": [row for row in all_rows if row["scope"] == "extended"],
        "Robust全数据流": [row for row in all_rows if row["scope"] == "extended_robust"],
        "GM数据流权衡": tradeoff, "反推带宽": supply,
        "硬件资源": resources, "时序敏感性": sensitivity, "带宽交叉点": crossings,
    })
    sources = {path.name: sha256(path) for path in (
        selection_path, Path(__file__), Path(__file__).with_name("dataflow.py"),
        Path(__file__).with_name("timing.py"), Path(__file__).with_name("geometry.py"))}
    write_json(out / "REPORT_AUDIT.json", {
        "status": "PASS", "cross_dataflow_rows": len(all_rows),
        "distinct_repeated_points": repeats, "repetitions_each": 2,
        "point_comparisons_identical": True, "tradeoff_rows": len(tradeoff),
        "crossover_rows": len(crossings), "workbook_sheets": 9,
        "workbook_readback_cells_checked": cells,
        "workbook": str(workbook), "source_sha256": sources,
        "native_timing_calibrated": False, "WS_IS_original_hardware_support_claimed": False,
        "all_cross_flow_compute_times_identical_by_declared_ideal_ports": True,
        "capacity_rows_are_conservative_upper_bounds": True,
    })
    print(canonical({"status": "PASS", "rows": len(all_rows), "repeat_points": repeats,
                     "workbook_cells": cells, "workbook": str(workbook)}), flush=True)
    for row in tradeoff:
        if row["dataflow"] in ("OS", "OS/OS"):
            print(canonical(row), flush=True)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", type=Path, required=True)
    main(parser.parse_args().out)
