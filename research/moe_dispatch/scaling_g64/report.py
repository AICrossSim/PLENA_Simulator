"""Reproducible tables for the fixed G=64 scaling experiment.

All paired comparisons are formed per archived window before batch-balanced
aggregation.  No configuration is selected using these held-out results.
"""
from __future__ import annotations

import argparse
import csv
import json
import math
from collections import defaultdict
from pathlib import Path

import numpy as np

BATCHES = (2, 4, 8, 16, 64, 96, 128)
COUNTS = {2: 27, 4: 27, 8: 27, 16: 27, 64: 9, 96: 9, 128: 9}
MAIN_PROFILE = "pipelined_tree1"
POLICY_LABELS = {
    "static": "静态亲和绑定",
    "naive": "旧准入（满队列强制溢出）",
    "no_preconfig": "修正 M-DPU／无预配置",
    "full": "修正 M-DPU／完整",
}
ROLE_LABELS = {
    "single": "单核",
    "homo": "同构双核",
    "hetero": "异构 P+E",
    "fair_homo_stream": "同能力对照／同构流式核",
    "fair_stream_cluster": "同能力对照／独立流式簇",
}
VARIANT_LABELS = {
    "main": "指定主配置",
    "alt_p": "备选 P 核",
    "split_e": "E 拆为独立执行单元",
    "alt_p_split_e": "备选 P＋独立 E 单元",
}


def gm(values):
    values = tuple(float(v) for v in values)
    if not values or any(v <= 0 or not math.isfinite(v) for v in values):
        raise ValueError("Geometric mean requires finite positive values")
    return math.exp(math.fsum(math.log(v) for v in values) / len(values))


def mean(values):
    values = tuple(float(v) for v in values if v is not None)
    return math.fsum(values) / len(values) if values else None


def _dict(configs):
    if isinstance(configs, dict):
        if "configs" in configs:
            return _dict(configs["configs"])
        if "designs" in configs:
            return _dict(configs["designs"])
        return configs
    return {c.get("design_id", c.get("id")): c for c in configs}


def physical(cores, groups):
    def shape(c):
        return f"{c['pm']}×{c['pn']}×{c['groups']}"
    return " + ".join("(" + " ∥ ".join(shape(cores[i]) for i in group) + ")"
                      if len(group) > 1 else shape(cores[group[0]]) for group in groups)


def _core_useful(result, ncores):
    stats = result.get("boundary_stats", {})
    if "core_useful_MACs" in stats:
        return list(stats["core_useful_MACs"])
    for key in ("core_useful_MACs", "core_useful_macs"):
        if key in result:
            return list(result[key])
    totals = [0] * ncores
    if "jobs" not in result:
        raise ValueError("Per-core useful MACs or complete jobs are required")
    for job in result["jobs"]:
        totals[int(job["core_id"])] += int(job["useful_macs"])
    return totals


def flatten_record(record):
    result, cores = record["result"], record["cores"]
    scale = int(record["scale"])
    groups = record.get("groups", [[i] for i in range(len(cores))])
    if sorted(i for group in groups for i in group) != list(range(len(cores))):
        raise ValueError("Every physical worker must appear in exactly one group")
    budget = sum(int(c["macs"]) for c in cores)
    if budget != scale:
        raise ValueError(f"Budget mismatch: {budget} != {scale}")
    cycles, useful = int(result["cycles"]), int(result["useful_MACs"])
    issued = int(result["issued_macs"])
    busy, ends = list(result["core_cycles"]), list(result["core_finish_cycles"])
    if (len(busy) != len(cores) or len(ends) != len(cores) or max(ends) != cycles
            or not 0 < useful <= issued <= scale * cycles):
        raise ValueError("Invalid timing or MAC accounting")
    core_useful = _core_useful(result, len(cores))
    if len(core_useful) != len(cores) or sum(core_useful) != useful:
        raise ValueError("Per-core useful MACs do not conserve work")
    group_ends = [max(ends[i] for i in group) for group in groups]
    group_work = [sum(core_useful[i] for i in group) for group in groups]
    group_caps = [sum(cores[i]["macs"] for i in group) for group in groups]
    group_idle = [100 * (1 - sum(busy[i] * cores[i]["macs"] for i in group)
                               / (cycles * cap)) for group, cap in zip(groups, group_caps)]
    jobs = result.get("jobs")
    dimension_capacity = [0, 0, 0]
    if jobs is not None:
        ids = [j["expert_id"] for j in jobs]
        if len(ids) != len(set(ids)) or sum(j["useful_macs"] for j in jobs) != useful:
            raise ValueError("Expert jobs missing, duplicated, or work changed")
        cd = lambda n, d: (n + d - 1) // d
        for job in jobs:
            core = cores[job["core_id"]]
            m, h, f = job["M"], job["H"], job["F"]
            dimension_capacity[0] += 3 * cd(m, core["pm"]) * core["pm"] * h * f
            dimension_capacity[1] += (2 * m * h * cd(f, core["pn"]) * core["pn"]
                                      + m * f * cd(h, core["pn"]) * core["pn"])
            dimension_capacity[2] += (2 * m * cd(h, core["r"]) * core["r"] * f
                                      + m * cd(f, core["r"]) * core["r"] * h)
        for cid in range(len(cores)):
            owned = sorted((j for j in jobs if j["core_id"] == cid), key=lambda j: j["start"])
            if any(a["end"] > b["start"] for a, b in zip(owned, owned[1:])):
                raise ValueError("Two arithmetic jobs overlap on one worker")
            if owned and (owned[-1]["end"] != ends[cid]
                          or sum(j["end"] - j["start"] for j in owned) != busy[cid]):
                raise ValueError("Worker trace does not match service counters")
    two = len(groups) == 2
    balance = max(group_ends) / min(group_ends) if two and min(group_ends) > 0 else None
    weights = result.get("unique_weight_bytes")
    row = {
        "design_id": record["design_id"], "scale": scale,
        "role": record["role"], "variant": record.get("variant", "main"),
        "profile": record["profile"], "policy": record["policy"],
        "window_id": record["window_id"], "batch": int(record["batch"]),
        "physical_PM_PN_G": physical(cores, groups),
        "native_BLEN_MLEN": " + ".join(f"{c['b']}×{c['r']}" for c in cores),
        "MAC_allocation": " + ".join(str(n) for n in group_caps),
        "physical_workers": len(cores), "logical_groups": len(groups),
        "cycles": cycles, "compute_ms": cycles / 1e6,
        "effective_GFLOP_s": 2 * useful / cycles,
        "MAC_utilization_pct": 100 * useful / (scale * cycles),
        "padding_waste_pct": 100 * (issued - useful) / issued,
        "M_only_fill_pct": 100 * useful / dimension_capacity[0] if dimension_capacity[0] else None,
        "N_only_fill_pct": 100 * useful / dimension_capacity[1] if dimension_capacity[1] else None,
        "K_only_fill_pct": 100 * useful / dimension_capacity[2] if dimension_capacity[2] else None,
        "max_worker_busy_ms": max(busy) / 1e6,
        "P_useful_FLOPs_pct": 100 * group_work[0] / useful if two else None,
        "E_useful_FLOPs_pct": 100 * group_work[1] / useful if two else None,
        "P_idle_pct": group_idle[0] if two else None,
        "E_idle_pct": group_idle[1] if two else None,
        "P_finish_ms": group_ends[0] / 1e6 if two else None,
        "E_finish_ms": group_ends[1] / 1e6 if two else None,
        "finish_imbalance": balance,
        "unused_group_window": int(two and min(group_work) == 0),
        "useful_macs": useful, "issued_macs": issued,
        "core_busy_cycles": json.dumps(busy), "core_finish_cycles": json.dumps(ends),
        "core_useful_macs": json.dumps(core_useful),
        "context_exposed_cycles": result.get("context_exposed_cycles"),
        "context_overlap_cycles": result.get("context_overlap_cycles"),
        "context_service_cycles": result.get("context_service_cycles"),
        "steal_count": result.get("steal_count"),
        "affinity_deviation_count": result.get("affinity_deviation_count"),
        "frontend_backpressure_cycles": result.get("frontend_backpressure_cycles"),
        "prediction_mae_cycles": result.get("prediction_mae_cycles"),
        "unique_weight_MiB": weights / 1024**2 if weights is not None else None,
        "unique_weight_average_TB_s": weights / cycles / 1000 if weights is not None else None,
    }
    return row


GEOMETRIC = {"cycles", "compute_ms", "effective_GFLOP_s", "MAC_utilization_pct",
             "max_worker_busy_ms", "unique_weight_MiB", "unique_weight_average_TB_s",
             "M_only_fill_pct", "N_only_fill_pct", "K_only_fill_pct"}
ARITHMETIC = {"padding_waste_pct", "P_useful_FLOPs_pct", "E_useful_FLOPs_pct", "P_idle_pct",
              "E_idle_pct", "P_finish_ms", "E_finish_ms", "finish_imbalance",
              "context_exposed_cycles", "context_overlap_cycles", "context_service_cycles",
              "steal_count", "affinity_deviation_count", "frontend_backpressure_cycles",
              "prediction_mae_cycles"}


def aggregate(rows, batch):
    identifiers = {"design_id", "scale", "role", "variant", "profile", "policy",
                   "physical_PM_PN_G", "native_BLEN_MLEN", "MAC_allocation",
                   "physical_workers", "logical_groups"}
    result = {k: v for k, v in rows[0].items() if k in identifiers}
    result.update(batch=batch, window_count=len(rows))
    for key in GEOMETRIC | ARITHMETIC:
        values = [r[key] for r in rows if r.get(key) is not None]
        result[key] = gm(values) if key in GEOMETRIC and values else mean(values)
    result["unused_group_windows"] = sum(r.get("unused_group_window", 0) for r in rows)
    result["imbalance_defined_windows"] = sum(r.get("finish_imbalance") is not None for r in rows)
    return result


def summarize(rows):
    per_batch = [aggregate([r for r in rows if r["batch"] == b], b) for b in BATCHES]
    overall = aggregate(per_batch, "7-Batch GM")
    overall["window_count"] = len(rows)
    overall["unused_group_windows"] = sum(r["unused_group_window"] for r in rows)
    overall["imbalance_defined_windows"] = sum(r["finish_imbalance"] is not None for r in rows)
    return per_batch + [overall]


def speedup(reference, candidate, bootstrap=2000, seed=20261015):
    ref = {r["window_id"]: r for r in reference}
    cand = {r["window_id"]: r for r in candidate}
    if len(ref) != len(reference) or len(cand) != len(candidate) or ref.keys() != cand.keys():
        raise ValueError("Paired speedups require exactly matching window IDs")
    logs = [[] for _ in BATCHES]
    for identifier, row in cand.items():
        other = ref[identifier]
        if row["batch"] != other["batch"] or row["useful_macs"] != other["useful_macs"]:
            raise ValueError("Paired comparison workload changed")
        logs[BATCHES.index(row["batch"])].append(math.log(other["cycles"] / row["cycles"]))
    if any(not group for group in logs):
        raise ValueError("All batches required")
    rng = np.random.default_rng(seed)
    means = np.zeros(bootstrap)
    for group in logs:
        values = np.asarray(group)
        indices = rng.integers(len(values), size=(bootstrap, len(values)))
        means += values[indices].mean(axis=1) / len(BATCHES)
    samples = np.exp(means)
    result = {"speedup": math.exp(mean(mean(group) for group in logs)),
              "ci95_low": float(np.quantile(samples, .025)),
              "ci95_high": float(np.quantile(samples, .975)),
              "bootstrap_repeats": bootstrap, "seed": seed}
    result["by_batch"] = {str(b): math.exp(mean(group)) for b, group in zip(BATCHES, logs)}
    return result


def write_csv(path, rows):
    if not rows:
        return
    fields = list(dict.fromkeys(key for row in rows for key in row))
    with Path(path).open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def table(rows):
    if not rows:
        return "无此组结果。"
    fields = list(rows[0])
    lines = ["| " + " | ".join(fields) + " |", "|" + "|".join("---" for _ in fields) + "|"]
    for row in rows:
        cells = []
        for key in fields:
            value = row.get(key)
            if value is None:
                text = "—"
            elif isinstance(value, float):
                text = f"{value:.6f}" if "ms" in key else f"{value:.3f}"
            else:
                text = str(value)
            cells.append(text.replace("|", "\\|").replace("\n", " "))
        lines.append("| " + " | ".join(cells) + " |")
    return "\n".join(lines)


def workbook(path, sheets):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    book = Workbook()
    book.remove(book.active)
    for name, rows in sheets.items():
        sheet = book.create_sheet(name[:31])
        if not rows:
            continue
        fields = list(rows[0])
        sheet.append(fields)
        for row in rows:
            sheet.append([json.dumps(row.get(f), ensure_ascii=False)
                          if isinstance(row.get(f), (dict, list)) else row.get(f) for f in fields])
        sheet.freeze_panes = "D2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17365D")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[1].height = 42
        for i, field in enumerate(fields, 1):
            sheet.column_dimensions[get_column_letter(i)].width = min(58, max(17, len(str(field)) * 1.3))
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = "0.000000" if "ms" in str(fields[cell.column - 1]) else "0.000"
    book.save(path)


def generate_report(records, configs, out):
    out = Path(out)
    out.mkdir(parents=True, exist_ok=True)
    configs = _dict(configs)
    flat = [flatten_record(record) for record in records]
    grouped = defaultdict(list)
    for row in flat:
        grouped[(row["design_id"], row["profile"], row["policy"])].append(row)
    if configs and {key[0] for key in grouped} != set(configs):
        raise ValueError("Record designs do not exactly match frozen configurations")
    summaries = {}
    for key, rows in grouped.items():
        counts = {b: sum(r["batch"] == b for r in rows) for b in BATCHES}
        if counts != COUNTS or len({r["window_id"] for r in rows}) != 135:
            raise ValueError(f"Incomplete or duplicated held-out set for {key}: {counts}")
        summaries[key] = summarize(rows)

    def identify(row):
        role = ROLE_LABELS.get(row["role"], row["role"])
        variant = VARIANT_LABELS.get(row["variant"], row["variant"])
        return f"{role}／{variant}" if row["role"] == "hetero" else role

    def baseline(scale, profile, role):
        matches = [rows for (design, p, policy), rows in grouped.items()
                   if p == profile and policy == "full" and rows[0]["scale"] == scale
                   and rows[0]["role"] == role and rows[0]["variant"] == "main"]
        if len(matches) != 1:
            raise ValueError(f"Expected unique {role} baseline for {scale}, {profile}")
        return matches[0]

    comparison_cache = {}
    def comparison(key, role):
        cache_key = key, role
        if cache_key not in comparison_cache:
            rows = grouped[key]
            comparison_cache[cache_key] = speedup(baseline(rows[0]["scale"], key[1], role), rows)
        return comparison_cache[cache_key]

    keys = sorted(grouped, key=lambda key: (grouped[key][0]["scale"],
                  {"single": 0, "homo": 1, "hetero": 2}.get(grouped[key][0]["role"], 3),
                  {"main": 0, "alt_p": 1, "split_e": 2, "alt_p_split_e": 3}.get(grouped[key][0]["variant"], 4),
                  key[0], key[1], {"static": 0, "naive": 1, "no_preconfig": 2, "full": 3}.get(key[2], 4)))
    main_keys = [key for key in keys if key[1:] == (MAIN_PROFILE, "full")]
    main_table, batch_table, balance_table, ablation_table = [], [], [], []
    fair_table, sensitivity_table, comparisons = [], [], []
    for key in main_keys:
        summary = summaries[key][-1]
        cs, ch = comparison(key, "single"), comparison(key, "homo")
        row = {"规模 MACs": summary["scale"], "组织／变体": identify(summary),
               "配置 ID": key[0], "物理尺寸 PM×PN×G": summary["physical_PM_PN_G"],
               "P/E 乘法器分配": summary["MAC_allocation"],
               "7-Batch GM 耗时 ms": summary["compute_ms"],
               "有效 GFLOP/s": summary["effective_GFLOP_s"],
               "整片有效 MAC 利用率 %": summary["MAC_utilization_pct"],
               "仅 M 维行填充率 %": summary["M_only_fill_pct"],
               "相对单核加速比": cs["speedup"], "相对同构加速比": ch["speedup"],
               "相对同构 95% CI": f"[{ch['ci95_low']:.4f}, {ch['ci95_high']:.4f}]"}
        (fair_table if summary["role"].startswith("fair_") else main_table).append(row)
        comparisons.append({"design_id": key[0], "scale": summary["scale"],
                            "vs_single": cs, "vs_homo": ch})
        if summary["role"] in ("single", "homo", "hetero"):
            for item in summaries[key]:
                batch = item["batch"]
                batch_table.append({"规模 MACs": item["scale"], "Batch": batch,
                                    "组织／变体": identify(item), "配置 ID": key[0],
                                    "计算耗时 ms": item["compute_ms"],
                                    "有效 GFLOP/s": item["effective_GFLOP_s"],
                                    "整片有效 MAC 利用率 %": item["MAC_utilization_pct"],
                                    "相对单核加速比": cs["speedup"] if isinstance(batch, str) else cs["by_batch"][str(batch)],
                                    "相对同构加速比": ch["speedup"] if isinstance(batch, str) else ch["by_batch"][str(batch)]})
        if summary["role"] == "hetero":
            for item in summaries[key]:
                balance_table.append({"规模 MACs": item["scale"], "Batch": "7-Batch 等权均值" if isinstance(item["batch"], str) else item["batch"],
                                      "组织／变体": identify(item), "配置 ID": key[0],
                                      "P 有效 FLOPs %": item["P_useful_FLOPs_pct"],
                                      "E 有效 FLOPs %": item["E_useful_FLOPs_pct"],
                                      "P 平均空闲率 %": item["P_idle_pct"],
                                      "E 容量加权空闲率 %": item["E_idle_pct"],
                                      "P 完成时间均值 ms": item["P_finish_ms"],
                                      "E 完成时间均值 ms": item["E_finish_ms"],
                                      "完成时间失配比均值": item["finish_imbalance"],
                                      "失配比有定义窗口数": item["imbalance_defined_windows"],
                                      "一侧零任务窗口数": item["unused_group_windows"]})
    for key in keys:
        item = summaries[key][-1]
        if item["role"] == "hetero" and key[1] == MAIN_PROFILE:
            ref = grouped[(key[0], key[1], "full")]
            rel = speedup(grouped[key], ref)
            ablation_table.append({"规模 MACs": item["scale"], "组织／变体": identify(item),
                                   "配置 ID": key[0], "调度策略": POLICY_LABELS.get(key[2], key[2]),
                                   "7-Batch GM 耗时 ms": item["compute_ms"],
                                   "整片有效 MAC 利用率 %": item["MAC_utilization_pct"],
                                   "完整 M-DPU 相对此策略加速比": rel["speedup"],
                                   "P 有效 FLOPs %": item["P_useful_FLOPs_pct"],
                                   "E 有效 FLOPs %": item["E_useful_FLOPs_pct"],
                                   "暴露上下文周期均值": item["context_exposed_cycles"],
                                   "队首借调次数均值": item["steal_count"]})
        if key[2] == "full":
            cs, ch = comparison(key, "single"), comparison(key, "homo")
            sensitivity_table.append({"规模 MACs": item["scale"], "组织／变体": identify(item),
                                      "配置 ID": key[0], "时序模型": key[1],
                                      "7-Batch GM 耗时 ms": item["compute_ms"],
                                      "MAC 利用率 %": item["MAC_utilization_pct"],
                                      "相对同模型单核加速比": cs["speedup"],
                                      "相对同模型同构加速比": ch["speedup"]})

    config_table = []
    seen = set()
    for key in keys:
        if key[0] in seen:
            continue
        seen.add(key[0])
        row = grouped[key][0]
        note = "固定配置；未以留出集选型"
        if row["scale"] == 262144 and row["role"] == "hetero" and "split_e" not in row["variant"]:
            note += "；单体 E 的 PM=4，不满足窄行 PM≤2 条件"
        elif row["role"] == "hetero" and "split_e" in row["variant"]:
            note += "；E 内有两个独立专家执行单元，额外上下文／仲裁必须计入"
        if row["role"].startswith("fair_"):
            note += "；检验流式执行能力是否必须依赖异构"
        config_table.append({"配置 ID": key[0], "规模 MACs": row["scale"], "组织／变体": identify(row),
                             "物理尺寸 PM×PN×G": row["physical_PM_PN_G"],
                             "BLEN×MLEN（逐 worker）": row["native_BLEN_MLEN"],
                             "MAC 分配（逻辑组）": row["MAC_allocation"],
                             "独立执行单元数": row["physical_workers"], "口径": note})
    all_aggregates = [row for key in keys for row in summaries[key]]
    sheets = {"00_配置与口径": config_table, "01_规模GM": main_table,
              "02_逐Batch": batch_table, "03_PE负载": balance_table,
              "04_调度消融": ablation_table, "05_同能力对照": fair_table,
              "06_时序敏感性": sensitivity_table, "07_全部聚合": all_aggregates,
              "08_逐窗口": flat}
    csv_names = {"00_配置与口径": "configurations", "01_规模GM": "table1_scaling",
                 "02_逐Batch": "table2_by_batch", "03_PE负载": "table3_balance",
                 "04_调度消融": "table4_ablation", "05_同能力对照": "fair_controls",
                 "06_时序敏感性": "timing_sensitivity", "07_全部聚合": "all_aggregates",
                 "08_逐窗口": "window_metrics"}
    for name, rows in sheets.items():
        write_csv(out / (csv_names[name] + ".csv"), rows)
    workbook(out / "G64_算力扩展_完整数据.xlsx", sheets)

    summary_json = {"conditions": len(grouped), "records": len(flat), "unique_windows": 135,
                    "batch_counts": COUNTS, "comparisons": comparisons,
                    "table1": main_table, "fair_controls": fair_table,
                    "table4": ablation_table, "configurations": config_table}
    (out / "report_summary.json").write_text(json.dumps(summary_json, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    intro = (
        "# G=64、1:1 P/E 算力分配：固定配置扩展实验\n\n"
        "本报告使用全部 135 个归档真实路由窗口，执行 Shared 加全部活跃 Routed 专家。"
        "B2/B4/B8/B16 各 27 个窗口，B64/B96/B128 各 9 个。硬件配置预先固定，未根据留出结果挑选赢家。\n\n"
        "这些留出窗口已在前序研究中反复使用；本轮不是新增的盲测数据集。\n\n"
        "**这是理想供数下的解析计算与有限控制事件仿真，不是实测芯片端到端性能。** "
        "HBM、SRAM 端口和数据搬运延迟均为零；保留阵列流水线、任务上下文配置和队列控制时间。"
        "1 GHz 是折算假设，不是本轮综合得到的频率。未报告经过综合验证的面积或功耗。\n\n"
        "统计口径：先对每个 Batch 的窗口取几何平均，再对 7 个 Batch 等权几何平均。"
        "有效 GFLOP/s 仅计真实矩阵运算；利用率分母是整片乘法器总数乘整层完成周期。"
        "“仅 M 维行填充率”只将 token 行补齐到所属 worker 的 PM，N/K 使用真实尺寸；"
        "它单独量化行 padding，不混入控制停顿或 N/K padding。逐窗口 CSV 另列 N、K 单维填充率。"
        "加速比为相同窗口的基线时间／候选时间，超过 1 才表示候选更快。"
        "95% 区间采用 2,000 次按 Batch 分层、窗口配对 bootstrap；它不包含时序模型误差。"
        "P/E 工作量占比、空闲率等比例量先按窗口算术平均，再对 Batch 等权算术平均。\n\n"
        "**必须更正原 Prompt 的三个物理断言。** G 是 MLEN/BLEN，实际 reduction 分块为 MLEN=BLEN×G，"
        "因此固定 G=64 不保证 K 方向零 padding。64 路平衡归约树有 6 个逻辑层级，"
        "但不能据此证明 BF16/FP32 在 1 GHz 单周期完成。紧凑核 PN≤256、流式核 PN=512 "
        "是不同组织的假设，需要综合互连／寄存器开销后才能谈等面积优势。\n\n"
        "单体 `4×512×64` 与两个独立 `2×512×64` 不等价：前者一次处理一个专家的四行，"
        "后者可以同时处理两个专家。以下分别报告；`∥` 表示独立执行单元。"
        "256K 单体 E 的 PM=4 不满足原 Prompt 的 PM≤2 限制，不能用它作为窄行实现的证明。\n\n"
    )
    sections = [intro, "## 配置账本\n\n" + table(config_table),
                "## 表 1：跨规模 7-Batch GM\n\n" + table(main_table),
                "## 表 2：逐 Batch 时间与加速比\n\n" + table(batch_table),
                "## 表 3：P/E 分工与负载均衡\n\n"
                "E 空闲率按其所有独立单元的 MAC 容量加权；E 完成时间取这些单元完成时刻的最大值。"
                "这里的空闲指未执行算术服务的时间，执行期间的 padding 浪费另计，不能把 0% 空闲理解为 100% 有效 MAC 利用率。"
                "完成时间失配比先逐窗口算 max/min，再等权平均；一侧没有工作时比值未定义，"
                "不把它记为 1，而单独计数。该比值描述完成时刻，不替代工作量占比。\n\n" + table(balance_table),
                "## 表 4：同一硬件上的调度消融\n\n"
                "四种策略只改变派工／预配置，不改变任务、几何与乘法器总数。"
                "上下文预配置并非权重预取；当前数据零延迟模型不能证明 HBM/SRAM 预取收益。\n\n" + table(ablation_table),
                "## 同能力对照\n\n"
                "让对照设计同样使用流式／独立执行能力，检验收益是否来源于异构本身。\n\n" + table(fair_table),
                "## 时序敏感性\n\n"
                "每一行只与同一时序假设下的单核和同构比较；不将不同假设的数字混在一起。\n\n" + table(sensitivity_table),
                "## 可复现文件\n\n"
                "`window_metrics.csv` 保留每个条件、每个窗口的所有统计；`all_aggregates.csv` 保留所有 Batch 聚合。"
                "工作簿包含配置、四张主表、同能力对照、时序敏感性与逐窗口结果。"
                "测试／RTL 重放证据由本轮验证文件单独提供；本报告不会将解析模型的数值当作综合或芯片测量。\n"]
    (out / "REPORT.md").write_text("\n\n".join(sections), encoding="utf-8")
    for name, rows in (("TABLE1.md", main_table), ("TABLE2.md", batch_table),
                       ("TABLE3.md", balance_table), ("TABLE4.md", ablation_table)):
        (out / name).write_text(table(rows) + "\n", encoding="utf-8")
    return summary_json


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--records", type=Path)
    parser.add_argument("--configs", type=Path)
    parser.add_argument("--out", "--output", dest="out", type=Path, required=True)
    args = parser.parse_args()
    records_path = args.records or args.out / "records.jsonl"
    configs_path = args.configs or args.out / "configs.json"
    receipt_path = records_path.parent / "run_receipt.json"
    receipt = json.loads(receipt_path.read_text())
    if not receipt.get("completed") or not receipt.get("all_repetitions_byte_identical"):
        raise ValueError("Report requires a completed, repeat-identical run receipt")
    with records_path.open() as handle:
        records = (json.loads(line) for line in handle if line.strip())
        summary = generate_report(records, json.loads(configs_path.read_text()), args.out)
    if summary["records"] != receipt["unique_points"]:
        raise ValueError("Run receipt and report record counts disagree")
    print(json.dumps({"conditions": summary["conditions"], "records": summary["records"],
                      "out": str(args.out)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
