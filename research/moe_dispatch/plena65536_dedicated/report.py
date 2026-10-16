"""Captured-route tables and independent audits for dedicated compute DSE."""
from __future__ import annotations

import argparse
import csv
import gzip
import hashlib
import json
import math
from pathlib import Path

from research.moe_dispatch.plena65536_moe.workloads import (
    BATCHES, load_workloads, profile_window, write_workload_artifacts,
)

BUDGET = 65_536
MIB = 1024**2
DEFAULT_OUT = Path("/scratch/shared/mcl123/plena/outputs/plena65536_dedicated_20261012")
LABELS = {
    "original": "原论文单核",
    "strict_single": "DSE最优单核",
    "strict_homogeneous": "DSE最优同构",
    "strict_dedicated_largeShared": "DSE专职异构（Shared大核）",
    "strict_dedicated_free": "专职异构（允许Routed资源更大）",
    "extended_single": "DSE最优单核",
    "extended_homogeneous": "DSE最优同构",
    "extended_dedicated_largeShared": "DSE专职异构（Shared大核）",
    "extended_dedicated_free": "专职异构（允许Routed资源更大）",
}


def gm(values):
    values = tuple(float(x) for x in values)
    if not values or any(not math.isfinite(x) or x < 0 for x in values):
        raise ValueError("GM requires finite nonnegative nonempty values")
    return 0.0 if any(x == 0 for x in values) else math.exp(math.fsum(math.log(x) for x in values) / len(values))


def canonical(value):
    return json.dumps(value, sort_keys=True, ensure_ascii=False, separators=(",", ":"), allow_nan=False)


def sha256(path):
    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def write_json(path, value):
    Path(path).write_text(json.dumps(value, indent=2, ensure_ascii=False, allow_nan=False) + "\n", encoding="utf-8")


def write_csv(path, rows):
    if not rows:
        raise ValueError(f"No rows for {path}")
    with Path(path).open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def physical_string(cores):
    return " + ".join(f"{c.pm}×{c.pn}×{c.groups}" for c in cores)


def frame_string(cores):
    return " + ".join(f"BLEN={c.b},MLEN={c.r}" for c in cores)


def scope_for(name):
    return "PM1_extension" if name.startswith("extended_") else "paper_Blen_ge2"


def family_for(name):
    if name == "original":
        return "original"
    if name.endswith("single"):
        return "single"
    if name.endswith("homogeneous"):
        return "homogeneous"
    return "dedicated"


def _ceildiv(a, b):
    return (a + b - 1) // b


def expert_issued_macs(core, expert):
    """Independent count of physically issued feed slots, including padding."""
    m, h, f = (int(expert[key]) for key in ("Me", "H", "F"))
    feed = (2 * _ceildiv(m, core.pm) * _ceildiv(f, core.pn) * _ceildiv(h, core.r) * core.b
            + _ceildiv(m, core.pm) * _ceildiv(h, core.pn) * _ceildiv(f, core.r) * core.b)
    return feed * core.macs


def checked_point(cores, window, scheduler, profile=None):
    """Audit affinity, all experts, arithmetic, and deterministic replay."""
    plan = scheduler(cores, window) if profile is None else scheduler(cores, window, profile)
    repeat = scheduler(cores, window) if profile is None else scheduler(cores, window, profile)
    if canonical(plan) != canonical(repeat):
        raise AssertionError("Independent schedule replay changed")
    if sum(c.macs for c in cores) != BUDGET:
        raise AssertionError("Every design requires exactly 65,536 MACs")
    observed = profile_window(window)
    if observed["routed_useful_macs"] != 3 * observed["shared_useful_macs"]:
        raise AssertionError("All Top-6 Routed MACs must equal 3x Shared")
    times = [int(x) for x in plan["core_cycles"]]
    if len(times) != len(cores) or int(plan["cycles"]) != max(times):
        raise AssertionError("Completion must equal max of independent core times")
    experts = {int(e["id"]): e for e in window["experts"]}
    if sorted(int(job["expert_id"]) for job in plan["jobs"]) != sorted(experts):
        raise AssertionError("All active experts must execute exactly once")
    issued, useful = 0, 0
    ends = [0] * len(cores)
    for job in plan["jobs"]:
        cid, eid = int(job["core_id"]), int(job["expert_id"])
        if len(cores) == 2 and cid != (0 if eid == -1 else 1):
            raise AssertionError("Hard affinity violated: Shared=Core0, Routed=Core1")
        if not 0 <= cid < len(cores):
            raise AssertionError("Invalid execution core")
        if int(job["start"]) != ends[cid] or int(job["end"]) <= int(job["start"]):
            raise AssertionError("Core queue overlaps, stalls, or loses tasks")
        ends[cid] = int(job["end"])
        e = experts[eid]
        issued += expert_issued_macs(cores[cid], e)
        useful += 3 * int(e["Me"]) * int(e["H"]) * int(e["F"])
    if ends != times:
        raise AssertionError("Independent queue ends differ from core totals")
    if useful != observed["total_useful_macs"] or issued < useful:
        raise AssertionError("Useful MAC conservation or padded-slot count failed")
    if "issued_macs" in plan and int(plan["issued_macs"]) != issued:
        raise AssertionError("Issued arithmetic differs from independent slot count")
    if int(plan["useful_MACs"]) != useful or int(plan["unique_weight_bytes"]) != observed["total_unique_weight_bytes"]:
        raise AssertionError("Model arithmetic or unique weight ledger differs from captured descriptors")
    if any(int(plan[key]) != 0 for key in ("migration_count", "steal_count", "supply_delay_cycles")):
        raise AssertionError("Migration, stealing, or supply latency violates the compute contract")
    return plan, observed, issued


def window_row(name, cores, split, window, plan, observed, issued):
    cycles, dual = int(plan["cycles"]), len(cores) == 2
    times = [int(t) for t in plan["core_cycles"]]
    weights = observed["total_unique_weight_bytes"]
    row = {
        "design": name, "organization": LABELS.get(name, name), "scope": scope_for(name),
        "family": family_for(name), "split": split, "batch": int(window["batch"]),
        "row_type": "window", "window_id": str(window["id"]), "window_count": 1,
        "physical_PM_PN_G": physical_string(cores), "native_frames": frame_string(cores),
        "MAC_allocation": "+".join(str(c.macs) for c in cores),
        "shared_core_ms": times[0] / 1e6 if dual else None,
        "routed_core_ms": times[1] / 1e6 if dual else None,
        "compute_cycles": cycles, "compute_ms": cycles / 1e6,
        "useful_GFLOP_s": 2 * observed["total_useful_macs"] / cycles,
        "wall_peak_utilization_pct": 100 * observed["total_useful_macs"] / (BUDGET * cycles),
        "spatial_utilization_pct": 100 * observed["total_useful_macs"] / issued,
        "speedup_vs_optimal_single": 0.0, "speedup_vs_optimal_homogeneous": 0.0,
        "core0_duty_pct": 100 * times[0] / cycles,
        "core1_duty_pct": 100 * times[1] / cycles if dual else None,
        "core0_wait_ms": (cycles - times[0]) / 1e6,
        "core1_wait_ms": (cycles - times[1]) / 1e6 if dual else None,
        "routed_over_shared_time": times[1] / times[0] if dual else None,
        "slow_fast_ratio": max(times) / min(times) if dual else None,
        "unique_weight_MiB": weights / MIB,
        "unique_weight_fetch_ms_at256": weights / 256e6,
        "equivalent_unique_weight_BW_GB_s": weights / cycles,
        "unique_weight_service_over_compute": weights / (256 * cycles),
        "shared_useful_macs": observed["shared_useful_macs"],
        "routed_useful_macs": observed["routed_useful_macs"],
        "total_useful_macs": observed["total_useful_macs"], "issued_macs": issued,
        "active_routed_experts": observed["active_routed_experts"],
        "shared_jobs_core0": 1,
        "routed_jobs_core0": observed["active_routed_experts"] if not dual else 0,
        "shared_jobs_core1": 0,
        "routed_jobs_core1": observed["active_routed_experts"] if dual else 0,
        "hard_affinity_pass": True, "repeat_identical": True,
    }
    if row["wall_peak_utilization_pct"] > 100 + 1e-9 or row["spatial_utilization_pct"] > 100 + 1e-9:
        raise AssertionError("Useful utilization exceeds physical capacity")
    return row


MEAN_FIELDS = {"active_routed_experts", "shared_jobs_core0", "routed_jobs_core0",
               "shared_jobs_core1", "routed_jobs_core1", "core0_wait_ms", "core1_wait_ms"}


def aggregate(rows, batch, row_type):
    result = dict(rows[0])
    result.update(batch=batch, row_type=row_type, window_id="", window_count=sum(r["window_count"] for r in rows))
    for key, value in rows[0].items():
        if key in {"batch", "window_count"} or isinstance(value, bool) or value is None:
            continue
        if isinstance(value, (int, float)):
            result[key] = (math.fsum(r[key] for r in rows) / len(rows) if key in MEAN_FIELDS
                           else gm(r[key] for r in rows))
    return result


def fill_speedups(rows):
    lookup = {(r["design"], r["split"], r["row_type"], str(r["batch"]), r["window_id"]): r for r in rows}
    for r in rows:
        prefix = "extended" if r["scope"] == "PM1_extension" else "strict"
        for family, field in (("single", "speedup_vs_optimal_single"),
                              ("homogeneous", "speedup_vs_optimal_homogeneous")):
            key = (prefix + "_" + family, r["split"], r["row_type"], str(r["batch"]), r["window_id"])
            r[field] = lookup[key]["compute_ms"] / r["compute_ms"]


def user_table1(rows, *, original_baselines=None):
    table = []
    for r in rows:
        single, hom = r["speedup_vs_optimal_single"], r["speedup_vs_optimal_homogeneous"]
        if r["design"] == "original" and original_baselines is not None:
            base = original_baselines[(r["split"], str(r["batch"]))]
            single, hom = base["single"] / r["compute_ms"], base["homogeneous"] / r["compute_ms"]
        table.append({
            "Batch": f"B{r['batch']}" if isinstance(r["batch"], int) else r["batch"],
            "组织形态": r["organization"],
            "物理尺寸配置 (Core 0 + Core 1)": r["physical_PM_PN_G"],
            "算力分配 (MACs)": r["MAC_allocation"],
            "大核纯计算 T_大 (ms)": r["shared_core_ms"],
            "小核纯计算 T_小 (ms)": r["routed_core_ms"],
            "整层完成时间 max(T_大,T_小) (ms)": r["compute_ms"],
            "有效 GFLOP/s": r["useful_GFLOP_s"],
            "阵列有效利用率 (%)": r["wall_peak_utilization_pct"],
            "相比最优单核加速比": single,
            "相比最优同构加速比": hom,
        })
    return table


def user_table2(rows):
    return [{
        "组织形态": r["organization"],
        "7-Batch GM 纯计算耗时 (ms)": r["compute_ms"],
        "唯一权重等效 Fetch Time @256GB/s (ms)": r["unique_weight_fetch_ms_at256"],
        "维持纯计算所需等效带宽 Equivalent Fetch BW (GB/s)": r["equivalent_unique_weight_BW_GB_s"],
        "真实物理瓶颈状态判断": (
            f"理想供数下纯计算；每窗冷取唯一权重时，256 GB/s服务量约为计算的{r['unique_weight_service_over_compute']:.2f}倍，"
            "外部供数成为必要瓶颈；驻留系统需另审片上供数。"),
    } for r in rows]


def table(rows):
    if not rows:
        return ""
    headers = list(rows[0])
    lines = ["| " + " | ".join(headers) + " |", "|" + "|".join("---" for _ in headers) + "|"]
    for row in rows:
        values = []
        for key in headers:
            value = row[key]
            if value is None:
                value = "—"
            elif isinstance(value, float):
                value = f"{value:,.6f}" if "ms" in key else f"{value:,.4f}"
            values.append(str(value).replace("|", "\\|"))
        lines.append("| " + " | ".join(values) + " |")
    return "\n".join(lines)


def alignment_table(rows):
    return [{"Batch": f"B{r['batch']}" if isinstance(r["batch"], int) else r["batch"],
             "Shared ms": r["shared_core_ms"], "Routed ms": r["routed_core_ms"],
             "Routed/Shared 时间": r["routed_over_shared_time"],
             "Core0工作占层时间 (%)": r["core0_duty_pct"],
             "Core1工作占层时间 (%)": r["core1_duty_pct"],
             "Core0平均等待 ms": r["core0_wait_ms"],
             "Core1平均等待 ms": r["core1_wait_ms"]} for r in rows]


def paired_bootstrap(rows, *, repetitions=2000, seed=655362026):
    """Paired resampling within each batch, followed by equal batch weighting."""
    import numpy as np
    windows = {(r["design"], r["split"], r["window_id"]): r for r in rows if r["row_type"] == "window"}
    result = []
    rng = np.random.default_rng(seed)
    for split in ("development", "heldout"):
        for prefix in ("strict", "extended"):
            for suffix in ("largeShared", "free"):
                candidate = prefix + "_dedicated_" + suffix
                selected = [r for r in rows if r["row_type"] == "window" and r["design"] == candidate and r["split"] == split]
                if not selected:
                    continue
                for family in ("single", "homogeneous"):
                    baseline = prefix + "_" + family
                    log_groups, counts = [], {"win": 0, "tie": 0, "loss": 0}
                    for batch in BATCHES:
                        group = [r for r in selected if r["batch"] == batch]
                        logs = []
                        for row in group:
                            ref = windows[(baseline, split, row["window_id"])]
                            logs.append(math.log(ref["compute_ms"] / row["compute_ms"]))
                            counts["win" if row["compute_cycles"] < ref["compute_cycles"] else
                                   "tie" if row["compute_cycles"] == ref["compute_cycles"] else "loss"] += 1
                        log_groups.append(np.asarray(logs))
                    samples = np.zeros(repetitions)
                    for logs in log_groups:
                        sampled = rng.integers(0, len(logs), size=(repetitions, len(logs)))
                        samples += logs[sampled].mean(axis=1) / len(BATCHES)
                    low, high = np.quantile(np.exp(samples), (.025, .975))
                    result.append({
                        "split": split, "baseline": baseline, "candidate": candidate,
                        "batch_balanced_GM_speedup": math.exp(math.fsum(float(logs.mean()) for logs in log_groups) / len(BATCHES)),
                        "bootstrap_95pct_lower": float(low), "bootstrap_95pct_upper": float(high),
                        "candidate_win_windows": counts["win"], "candidate_tie_windows": counts["tie"],
                        "candidate_loss_windows": counts["loss"], "repetitions": repetitions, "seed": seed,
                    })
    return result


def actual_worst_alignment(rows):
    result = []
    for name in dict.fromkeys(r["design"] for r in rows if r["family"] == "dedicated"):
        for split in ("development", "heldout"):
            selected = [r for r in rows if r["design"] == name and r["split"] == split and r["row_type"] == "window"]
            if not selected:
                continue
            for kind, key in (("max_slow_fast_ratio", lambda r: r["slow_fast_ratio"]),
                              ("max_absolute_wait", lambda r: r["core0_wait_ms"] + r["core1_wait_ms"])):
                row = max(selected, key=key)
                result.append({
                    "design": name, "split": split, "criterion": kind, "window_id": row["window_id"],
                    "batch": row["batch"], "shared_core_ms": row["shared_core_ms"],
                    "routed_core_ms": row["routed_core_ms"], "compute_ms": row["compute_ms"],
                    "slow_fast_ratio": row["slow_fast_ratio"], "core0_wait_ms": row["core0_wait_ms"],
                    "core1_wait_ms": row["core1_wait_ms"],
                    "slower_core": 1 if row["routed_core_ms"] > row["shared_core_ms"] else 0,
                })
    return result


def make_workbook(out, tables):
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    path = out / "PLENA65536_严格专职DSE_完整表.xlsx"
    book = Workbook()
    book.remove(book.active)
    for name, rows in tables.items():
        sheet = book.create_sheet(name[:31])
        headers = list(rows[0])
        sheet.append(headers)
        for r in rows:
            sheet.append([r[k] for k in headers])
        sheet.freeze_panes, sheet.auto_filter.ref = "A2", sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17365D")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[1].height = 46
        for column in sheet.columns:
            header = str(column[0].value)
            sheet.column_dimensions[column[0].column_letter].width = min(45, max(14, len(header) * 1.2))
            for cell in column[1:]:
                if isinstance(cell.value, float):
                    cell.number_format = "0.000000" if "ms" in header else "0.0000"
    book.save(path)
    read, cells = load_workbook(path, data_only=True, read_only=True), 0
    for name, rows in tables.items():
        expected = [list(rows[0])] + [[r[k] for k in rows[0]] for r in rows]
        actual = list(read[name[:31]].values)
        if len(actual) != len(expected):
            raise AssertionError("Workbook row count changed on readback")
        for arow, erow in zip(actual, expected):
            for a, e in zip(arow, erow):
                cells += 1
                if isinstance(e, (int, float)) and not isinstance(e, bool):
                    if a is None or not math.isclose(float(a), float(e), rel_tol=1e-12, abs_tol=1e-12):
                        raise AssertionError("Workbook numeric cell changed on readback")
                elif a != e and not (a is None and e == ""):
                    raise AssertionError("Workbook cell changed on readback")
    read.close()
    return path, cells


def make_plot(out, aggregates):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    names = ("original", "extended_single", "extended_homogeneous", "extended_dedicated_largeShared")
    labels = ("Paper single", "DSE single", "Dedicated homogeneous", "Dedicated heterogeneous")
    fig, axes = plt.subplots(1, 2, figsize=(12, 4.3))
    for name, label in zip(names, labels):
        rows = [r for r in aggregates if r["design"] == name and r["split"] == "heldout" and r["row_type"] == "batch_GM"]
        if rows:
            axes[0].plot(range(len(rows)), [r["compute_ms"] for r in rows], "o-", label=label)
    axes[0].set_yscale("log")
    axes[0].set_ylabel("Whole-layer compute makespan (ms)")
    axes[0].set_title("65,536 MACs, 1 GHz, ideal zero-latency supply")
    axes[0].legend(fontsize=8)
    rows = [r for r in aggregates if r["design"] == "extended_dedicated_largeShared" and r["split"] == "heldout" and r["row_type"] == "batch_GM"]
    if rows:
        axes[1].plot(range(len(rows)), [r["shared_core_ms"] for r in rows], "o-", label="Core 0: Shared only")
        axes[1].plot(range(len(rows)), [r["routed_core_ms"] for r in rows], "s-", label="Core 1: Routed only")
    axes[1].set_yscale("log")
    axes[1].set_ylabel("Dedicated core compute time (ms)")
    axes[1].set_title("Hard affinity, independent queues (BLEN=1 allowed)")
    axes[1].legend(fontsize=8)
    for ax in axes:
        ax.set_xticks(range(len(BATCHES)), [f"B{b}" for b in BATCHES])
        ax.grid(True, alpha=.25)
    fig.tight_layout()
    path = out / "dedicated_compute_and_balance.png"
    fig.savefig(path, dpi=180)
    fig.savefig(out / "dedicated_compute_and_balance.pdf")
    plt.close(fig)
    return path


def run(out, *, scheduler=None, core_factory=None):
    if scheduler is None or core_factory is None:
        from .geometry import Core
        from .model import schedule
        scheduler, core_factory = scheduler or schedule, core_factory or Core
    out = Path(out)
    selection_path = out / "FROZEN_SELECTION.json"
    selection = json.loads(selection_path.read_text())
    inputs = load_workloads()
    write_workload_artifacts(inputs, out)
    names = [n for n in LABELS if n in selection["families"]]
    designs = []
    for name in names:
        metadata = selection["families"][name]["cores"]
        cores = tuple(core_factory(**{k: m[k] for k in ("b", "r", "rm", "rn")}) for m in metadata)
        designs.append((name, cores))
    all_rows, aggregates, repeated_points, sensitivity = [], [], 0, []
    from research.moe_dispatch.plena65536_shared.timing import PROFILES
    detail_path = out / "execution_details.jsonl.gz"
    with gzip.open(detail_path, "wt", encoding="utf-8") as detail:
        for name, cores in designs:
            for split in ("development", "heldout"):
                window_rows = []
                for window in inputs[split]:
                    plan, observed, issued = checked_point(cores, window, scheduler)
                    repeated_points += 1
                    window_rows.append(window_row(name, cores, split, window, plan, observed, issued))
                    detail.write(canonical({"design": name, "split": split, "window_id": window["id"],
                                            "plan": plan, "window_payload_sha256": observed["window_payload_sha256"],
                                            "total_useful_macs": observed["total_useful_macs"], "issued_macs": issued,
                                            "point_sha256": hashlib.sha256(canonical(plan).encode()).hexdigest()}) + "\n")
                batch_rows = [aggregate([r for r in window_rows if r["batch"] == b], b, "batch_GM") for b in BATCHES]
                combined = aggregate(batch_rows, "7-Batch GM", "7batch_GM")
                all_rows.extend(window_rows + batch_rows + [combined])
                aggregates.extend(batch_rows + [combined])
                # Sensitivity retains each main-profile frozen geometry.
                # It is distinct from search.profile_optima's separate rerun winners.
                for profile in PROFILES:
                    profile_rows = []
                    for window in inputs[split]:
                        plan, observed, issued = checked_point(cores, window, scheduler, profile)
                        repeated_points += 1
                        profile_rows.append(window_row(name, cores, split, window, plan, observed, issued))
                    batch_values = [gm(r["compute_ms"] for r in profile_rows if r["batch"] == b) for b in BATCHES]
                    sensitivity.append({
                        "design": name, "scope": scope_for(name), "split": split, "profile": profile.name,
                        "physical_PM_PN_G": physical_string(cores), "native_frames": frame_string(cores),
                        "MAC_allocation": "+".join(str(c.macs) for c in cores),
                        "GM_compute_ms_frozen_geometry": gm(batch_values),
                        "sample_weighted_GM_compute_ms": gm(r["compute_ms"] for r in profile_rows),
                        "separately_reoptimized_GM_cycles": selection["profile_optima"][profile.name][name]["development_GM_cycles"]
                        if name != "original" else None,
                        "reoptimized_value_split": "development" if name != "original" else None,
                    })
    fill_speedups(all_rows)
    fill_speedups(aggregates)
    order = {n: i for i, n in enumerate(names)}
    heldout = sorted([r for r in aggregates if r["split"] == "heldout"],
                     key=lambda r: (BATCHES.index(r["batch"]) if r["batch"] in BATCHES else 7, order[r["design"]]))
    primary_names = ("original", "extended_single", "extended_homogeneous", "extended_dedicated_largeShared")
    strict_names = ("original", "strict_single", "strict_homogeneous", "strict_dedicated_largeShared")
    primary = [r for r in heldout if r["design"] in primary_names]
    strict = [r for r in heldout if r["design"] in strict_names]
    summary = [r for r in primary if r["row_type"] == "7batch_GM"]
    strict_summary = [r for r in strict if r["row_type"] == "7batch_GM"]
    diagnostic = [r for r in heldout if r["design"].endswith("_free")]
    aligned = [r for r in primary if r["design"] == "extended_dedicated_largeShared"]
    base = {}
    for r in primary:
        if r["family"] in ("single", "homogeneous"):
            base.setdefault((r["split"], str(r["batch"])), {})[r["family"]] = r["compute_ms"]
    table1 = user_table1(primary, original_baselines=base)
    table2 = user_table2(summary)
    bootstrap = paired_bootstrap(all_rows)
    worst_actual = actual_worst_alignment(all_rows)
    write_csv(out / "all_compute_results.csv", all_rows)
    write_csv(out / "compute_summary.csv", aggregates)
    write_csv(out / "table1_pure_compute.csv", table1)
    write_csv(out / "table2_equivalent_supply.csv", table2)
    write_csv(out / "paired_bootstrap.csv", bootstrap)
    write_csv(out / "worst_actual_alignment.csv", worst_actual)
    write_csv(out / "frozen_timing_sensitivity.csv", sensitivity)
    tables = {"表1严格专职纯计算": table1, "表2等效访存反推": table2,
              "异构负载对齐": alignment_table(aligned), "BLEN≥2保留域对照": user_table1(strict),
              "BLEN≥2保留域供数": user_table2(strict_summary), "所有窗口结果": all_rows, "所有Batch汇总": aggregates,
              "真实窗口最大失配": worst_actual, "配对Bootstrap": bootstrap, "冻结硬件时序敏感性": sensitivity}
    if diagnostic:
        tables["Routed大资源诊断"] = user_table1(diagnostic)
    workbook, cells = make_workbook(out, tables)
    plot = make_plot(out, aggregates)
    lines = ["# 65,536 MACs 严格专职绑定纯计算 DSE", "",
             "双核Core0只执行Shared，Core1只执行全部活跃Routed；两核从0周期独立推进，整层时间严格取逐窗口max(TShared,TRouted)。所有任务均禁止代算、偷取与跨核切分。单核执行全部专家。", "",
             "BF16输入/权重、FP32累加，1GHz；外部带宽无穷、供数零延迟，HBM与SRAM端口等待为0。时间包含阵列发射与声明的归约/提交流水线，采用论文拓扑一致的解析模型，未进行原生Rust/RTL或硅片校准。", "",
             "18个development窗口选型后冻结硬件，135个heldout真实Top-6窗口评估；全部专家执行Gate/Up/Down。Shared F2816，Routed F1408，H2048。Shared处理B个token，Routed处理6B次token分配，Routed有用MAC逐窗等于Shared的3倍。认证最优范围为搜索证书声明的有限几何域。", "",
             "PM×PN×G中G=MLEN/BLEN，第三维不是逻辑GEMM的K。主表允许请求中的PM=1，须引入BLEN=1的标量子阵列扩展；保留BLEN≥2几何域的结果也完整附后。所搜索的复制几何属于声明的解析架构空间，尚未逐一证明有原论文RTL实现。主目标额外强制Core0 MACs≥Core1 MACs，使‘Shared大核/Routed小核’按物理资源成立；解除该资源大小关系的固定归属结果另列诊断，仍无任何代算。", "",
             "每Batch先取逐窗口max的GM，再对七Batch等权取GM。两核各自时间GM只用于负载分析，禁止以max(两核时间GM)替代GM(逐窗max)。利用率主列为有用MAC/(65,536×整层周期)，包含padding、流水线开销和专职核闲置；原始CSV额外spatial列为有用MAC/发射MAC。速度倍数=最优基线时间/当前时间，低于1表示变慢。", "",
             "**表1：全Batch纯计算结果（允许PM=1扩展，严格Shared大核/Routed小核）**", "", table(table1), "",
             "**表2：唯一权重等效访存反推**", "", table(table2), "",
             "唯一BF16权重=33MiB Shared+16.5MiB×活跃Routed专家。Fetch Time为字节/256GB/s服务量，等效带宽为字节/纯计算时间。每窗冷取的外部必要服务量大于计算时间；驻留系统实际端到端瓶颈需另审片上供数和控制开销。此反推不计入本轮DSE时间。", "",
             "**专职异构负载对齐**", "", table(alignment_table(aligned)), ""]
    if aligned:
        cores = next(c for n, c in designs if n == "extended_dedicated_largeShared")
        divisor = math.gcd(cores[0].macs, cores[1].macs)
        worst = max((r for r in aligned if r["row_type"] == "batch_GM"), key=lambda r: r["slow_fast_ratio"])
        wait = max((r for r in aligned if r["row_type"] == "batch_GM"), key=lambda r: r["core0_wait_ms"] + r["core1_wait_ms"])
        final = next(r for r in aligned if r["row_type"] == "7batch_GM")
        vs_hom = next(r for r in bootstrap if r["split"] == "heldout" and r["candidate"] == "extended_dedicated_largeShared"
                      and r["baseline"] == "extended_homogeneous")
        worst_window = next(r for r in worst_actual if r["design"] == "extended_dedicated_largeShared"
                            and r["split"] == "heldout" and r["criterion"] == "max_slow_fast_ratio")
        wait_window = next(r for r in worst_actual if r["design"] == "extended_dedicated_largeShared"
                           and r["split"] == "heldout" and r["criterion"] == "max_absolute_wait")
        slower = "Core1（Routed）" if worst["routed_core_ms"] > worst["shared_core_ms"] else "Core0（Shared）"
        lines += [f"认证选中Core0={physical_string((cores[0],))}，Core1={physical_string((cores[1],))}；{frame_string(cores)}；资源{cores[0].macs:,}+{cores[1].macs:,}，配比{cores[0].macs//divisor}:{cores[1].macs//divisor}。", "",
                  f"相对失配最大为B{worst['batch']}，慢/快核GM比{worst['slow_fast_ratio']:.4f}，拖尾为{slower}。绝对等待最大为B{wait['batch']}，两核等待之和的逐窗口算术平均{wait['core0_wait_ms']+wait['core1_wait_ms']:.6f}ms。", "",
                  f"逐窗口最大相对失配为{worst_window['window_id']}（B{worst_window['batch']}），Shared={worst_window['shared_core_ms']:.6f}ms、Routed={worst_window['routed_core_ms']:.6f}ms、慢/快比={worst_window['slow_fast_ratio']:.4f}。绝对等待最大窗口为{wait_window['window_id']}（B{wait_window['batch']}），等待={wait_window['core0_wait_ms']+wait_window['core1_wait_ms']:.6f}ms。", "",
                  f"专职异构7-Batch GM相对最优单核{final['speedup_vs_optimal_single']:.6f}×，相对最优同构{final['speedup_vs_optimal_homogeneous']:.6f}×。Top-6×F1408使Routed有用工作为Shared的3倍；Shared核至少取得一半资源时，几何只能改善padding/发射，闲置Core0仍无法处理Routed工作。", ""]
        if final["speedup_vs_optimal_single"] < 1 and vs_hom["candidate_tie_windows"] == len(inputs["heldout"]):
            lines += [f"净收益判定：当前严格专职异构没有拉开纯计算优势；与最优同构在全部{len(inputs['heldout'])}个留出窗口完全持平，较最优单核GM耗时增加{100*(1/final['speedup_vs_optimal_single']-1):.4f}%。", ""]
        lines += ["在严格Shared资源≥Routed资源下，即使完全消除padding与流水线开销，T≥RoutedMAC/Core1MAC，整片有效利用率≤Core1MAC/(0.75×65,536)≤66.67%。这是当前工作量与资源归属的结构上界。", ""]
    lines += ["**保留BLEN≥2几何域：同样的专职约束，不使用标量扩展**", "", table(user_table1(strict)), "", table(user_table2(strict_summary)), ""]
    if diagnostic:
        lines += ["**允许Routed资源更大的固定归属诊断**", "", table(user_table1(diagnostic)), ""]
    lines += ["**配对窗口收益区间（Heldout）**", "",
              table([r for r in bootstrap if r["split"] == "heldout"]), "",
              "Bootstrap按Batch内配对重采样2000次，再对七Batch等权GM；固定seed=655362026。区间只描述已捕获窗口的波动，未覆盖时序模型误差与层/step相关性。", ""]
    lines += [f"![纯计算耗时与专职核负载对齐]({plot})", "",
              "执行明细保留每个真实窗口的专家归属、起止周期和输入哈希。CSV保留原始窗口、Batch GM、七Batch等权GM及有用/padding账本。每点独立重算两次，Excel逐单元格回读核验。"]
    (out / "REPORT.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    audit = {"status": "PASS", "contract": "Core0 Shared only; Core1 all Routed only; no stealing; no slicing",
             "physical_MAC_budget": BUDGET, "clock_GHz": 1, "data_supply_latency_cycles": 0,
             "external_bandwidth": "infinity", "development_windows": len(inputs["development"]),
             "heldout_windows": len(inputs["heldout"]), "designs": len(designs),
             "repeated_points": repeated_points, "repetitions_each": 2, "rows": len(all_rows),
             "main_profile_points": len(designs) * (len(inputs["development"]) + len(inputs["heldout"])),
             "frozen_sensitivity_points": len(PROFILES) * len(designs) * (len(inputs["development"]) + len(inputs["heldout"])),
             "bootstrap_repetitions": 2000, "bootstrap_seed": 655362026,
             "all_experts_executed_once": True, "all_dualcore_affinity_pass": True,
             "all_dualcore_time_equals_max": True, "useful_macs_routed_equals_3x_shared": True,
             "workbook": str(workbook), "workbook_readback_cells": cells,
             "selection_sha256": sha256(selection_path), "report_source_sha256": sha256(Path(__file__)),
             "details_sha256": sha256(detail_path), "aggregation": "GM(window max) then equal GM of seven batches",
             "analytical_model_not_calibrated_measurement": True}
    write_json(out / "REPORT_AUDIT.json", audit)
    print(canonical(audit), flush=True)
    return audit


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT)
    run(parser.parse_args().out)
