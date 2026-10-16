"""Readable workbook and standalone research figures from audited CSVs."""
from __future__ import annotations
import argparse, csv, collections, json, math, statistics
from pathlib import Path
import numpy as np

def read(path):
    if not path.exists():return []
    return list(csv.DictReader(path.open()))

def table(headers,rows):
    return '\n'.join(['| '+' | '.join(headers)+' |','| '+' | '.join(['---']*len(headers))+' |']+
        ['| '+' | '.join(map(str,r))+' |' for r in rows])

def main():
    p=argparse.ArgumentParser();p.add_argument('--report',type=Path,required=True);a=p.parse_args();d=a.report
    audit=json.loads((d/'AUDIT.json').read_text())
    if not (audit['capture_complete'] and audit['profile_complete']):raise ValueError('publish only complete campaign')
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.utils import get_column_letter
    batches=[1,2,4,8,16,32,64,96,128,256]
    precisions=['bf16','fp16','fp8_w8a8','int8_w8a8','int8_w8a16','int4_w4a16']
    labels=['BF16','FP16','FP8 W8A8','INT8 W8A8','INT8 W8A16','INT4 W4A16']
    workload={int(r['batch']):r for r in read(d/'workload_by_batch.csv') if r['precision']=='bf16'}
    perf={(r['precision'],r['path'],int(r['batch'])):r for r in read(d/'profile_by_batch.csv')}
    speed={(r['precision'],int(r['batch'])):float(r['speedup_vs_bf16_GM']) for r in read(d/'precision_speedups.csv') if r['path']=='grouped'}
    err={(r['precision'],int(r['batch'])):r for r in read(d/'precision_error.csv')}
    rows=[]
    for b in batches:
        r=workload[b];s=perf['bf16','split',b]
        rows.append([f'B{b}',int(r['observations']),f"{float(r['hot_experts']):.2f}",f"{float(r['cold_experts']):.2f}",
            '25.00',f"{float(r['hot_flops_pct']):.2f}",f"{float(r['cold_flops_pct']):.2f}",
            f"{float(s['shared_time_pct']):.2f}",f"{float(s['hot_time_pct']):.2f}",f"{float(s['cold_time_pct']):.2f}"])
    lines=['# 完整 Decode 与六精度结果','',
        '26个MoE层 × 10档真实并发请求数 × 32个逐token decode步；3类请求、57个请求组。六种路由采集合计284,544个层/步窗口。GPU回放使用BF16捕获输入与路由，采样第1/4/8/16/32步、每点50次、每种精度两遍，共177,840个测量点。',
        '', '## 专家数量、计算量与时间占比','',
        '数量与计算量：全部32步的BF16路由算术平均。时间：5个采样步、两遍GPU回放的拆分调用，按阶段平均时间/平均总时间计算。Shared为2个共享专家合并后的FFN；有效矩阵FLOPs恒占25%，不等于实际耗时恒占25%。时间占比还剩Combine，故三项不必加到100%。', '',
        table(['Batch','路由窗口','Hot数','Cold数','Shared FLOPs %','Hot FLOPs %','Cold FLOPs %','Shared时间 %','Hot时间 %','Cold时间 %'],rows),
        '', '## 六精度整层GPU回放时间（µs）','',
        '主表是未插入内部阶段计时点的grouped路径，包含Shared、Routed、SiLU和合并，Router/Attention不在计时范围内。数值是各窗口采样均值的算术平均，未把拆分Hot/Cold路径当作原生Grouped路径。', '',
        table(['Batch']+labels,[[f'B{b}']+[f"{float(perf[p,'grouped',b]['e2e_us']):.3f}" for p in precisions] for b in batches]),
        '', '## 相对BF16加速比（相同窗口配对GM）','',
        table(['Batch']+labels,[[f'B{b}']+[f'{speed[p,b]:.3f}×' for p in precisions] for b in batches]),
        '', '## 局部FFN输出误差：relative L2（%）','',
        '参照是完整BF16模型捕获的该层FFN输出。这里只评估固定输入、固定路由的局部误差；不能据此宣称任务准确率保持不变。', '',
        table(['Batch']+labels,[[f'B{b}']+[f"{100*float(err[p,b]['relative_l2_mean']):.3f}" for p in precisions] for b in batches]),
        '', '## 测量边界','',
        '- B表示同一步的不同请求，不是prefill片段。所有26层均来自真实自回归执行。提示词上限1024 token，截断情况逐请求保留。',
        '- 固定生成32步，EOS后继续执行；各Batch的EOS后请求比例在workload_by_batch.csv。结果不等同于自然停止的在线服务。',
        '- BF16/FP16路由来自相应完整浮点模型；其余精度仅量化MoE的Shared和Routed FFN，attention等仍为BF16，router为FP32。自由生成的跨精度路由差异同时包含生成前缀分歧。',
        '- FP8/INT8 W8A8含动态激活量化；W8A16/W4A16只压缩权重，激活BF16并在GEMM内解量化。RTN量化未做任务准确率验证，也不是优化过的AWQ/GPTQ发布模型。',
        '- 同型号B200上每种精度独占一张GPU，未锁频；两遍物理测时不会逐位一致，重复差异见repeat_stability_summary.csv。默认内核配置不是各精度的最优实现证明。',
        '- 未测硬件性能计数器，不根据本表直接标记Memory-bound或Compute-bound。',
        '', '逐层、逐步、逐数据集、原始窗口、全部重复样本及来源见CSV、原始目录和README。']
    (d/'SUMMARY_ZH.md').write_text('\n'.join(lines)+'\n')
    wb=Workbook();ws=wb.active;ws.title='Readme'
    for line in lines[:4]+lines[-10:]:ws.append([line])
    ws.column_dimensions['A'].width=130
    for row in ws:row[0].alignment=Alignment(wrap_text=True,vertical='top')
    files=['coverage','workload_by_batch','workload_by_layer','workload_by_step','profile_by_batch','profile_by_layer','profile_by_step','profile_by_dataset',
        'precision_speedups','precision_error','precision_route_comparison','repeat_stability_summary','request_provenance','profile_coverage']
    for name in files:
        data=read(d/(name+'.csv'))
        if not data:continue
        ws=wb.create_sheet(name[:31]);ws.append(list(data[0]))
        for r in data:
            values=[]
            for value in r.values():
                try:value=float(value);value=int(value) if value.is_integer() else value
                except (ValueError,AttributeError):pass
                values.append(value)
            ws.append(values)
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='18354A')
        for j in range(1,ws.max_column+1):ws.column_dimensions[get_column_letter(j)].width=min(40,max(14,len(str(ws.cell(1,j).value))+2))
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value,float):cell.number_format='0.0000'
    wb.save(d/'Decode_Precision_Profile.xlsx')
    figures=d/'figures';figures.mkdir(exist_ok=True)
    layer_rows=[r for r in read(d/'workload_by_layer.csv') if r['precision']=='bf16']
    # Equal groups within each dataset; combine by actual observation count.
    sums=collections.defaultdict(lambda:[0.,0.])
    for r in layer_rows:
        n=int(r['observations']);v=sums[int(r['layer']),int(r['batch'])];v[0]+=n*float(r['cold_flops_pct']);v[1]+=n
    grid=np.array([[sums[l,b][0]/sums[l,b][1] for b in batches] for l in range(1,27)])
    fig,axes=plt.subplots(1,2,figsize=(13,7),layout='constrained')
    im=axes[0].imshow(grid,aspect='auto',vmin=0,vmax=75,cmap='viridis');axes[0].set_title('Cold expert share of useful GEMM FLOPs (%)')
    axes[0].set_xticks(range(10),batches,rotation=45);axes[0].set_yticks(range(0,26,2),range(1,27,2));axes[0].set_xlabel('Concurrent decode requests');axes[0].set_ylabel('MoE layer');fig.colorbar(im,ax=axes[0])
    matrix=np.array([[speed[p,b] for b in batches] for p in precisions])
    im=axes[1].imshow(matrix,aspect='auto',cmap='RdYlGn',vmin=0,vmax=max(2,float(matrix.max())));axes[1].set_title('Measured grouped speedup vs BF16')
    axes[1].set_xticks(range(10),batches,rotation=45);axes[1].set_yticks(range(6),labels);axes[1].set_xlabel('Concurrent decode requests')
    for i in range(6):
        for j in range(10):axes[1].text(j,i,f'{matrix[i,j]:.2f}',ha='center',va='center',fontsize=8)
    fig.colorbar(im,ax=axes[1]);fig.savefig(figures/'decode_precision_overview.pdf');fig.savefig(figures/'decode_precision_overview.png',dpi=180);plt.close(fig)
    print('Published workbook, concise report and figures')

if __name__=='__main__':main()
