"""Aggregate GPU event timings without equating split diagnostics to fused MoE.

No CUDA dependency. The command requires complete frozen-route coverage by
 default and will not infer hardware bottlenecks from latency shares.
"""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
import csv
import hashlib
import json
import math
from pathlib import Path
import statistics
from typing import Any, Iterable

from .traces import BATCHES, EXPECTED_COUNTS, load_routes, population_metrics, sha256

COMPONENTS = ("shared", "dispatch", "hot", "cold", "combine")
COMPONENT_LABELS = {"shared": "Shared", "dispatch": "分发", "hot": "Hot", "cold": "Cold", "combine": "合并", "other": "其他"}
SCHEMA = "gpu_hotcold_report_v1"


def finite(value: Any, name: str, *, positive: bool = False) -> float:
    value = float(value)
    if not math.isfinite(value) or value < 0 or (positive and value <= 0):
        raise ValueError(f"{name}: expected finite {'positive' if positive else 'nonnegative'} measurement")
    return value


def geometric_mean(values: Iterable[float]) -> float:
    """Only positive latency/ratio observations; zero category times are not GM'd."""
    values = [finite(v, "GM", positive=True) for v in values]
    if not values:
        raise ValueError("Geometric mean requires observations")
    return math.exp(statistics.fmean(math.log(v) for v in values))


def sample_statistics(values: Iterable[float]) -> dict[str, float | int]:
    values = [finite(v, "event sample") for v in values]
    if not values:
        raise ValueError("An event series must not be empty")
    return {"samples": len(values), "mean_ms": statistics.fmean(values),
            "median_ms": statistics.median(values),
            "stdev_ms": statistics.pstdev(values), "min_ms": min(values), "max_ms": max(values)}


def component_accounting(total: list[float], components: dict[str, list[float]],
                         *, tolerance_ms: float = 0.00002) -> dict[str, list[float]]:
    """Close adjacent-event intervals; retain nonnegative unassigned event time.

    Very small negative residuals can arise because CUDA event elapsed values
    are independently rounded; reject anything beyond the explicit tolerance.
    """
    totals = [finite(t, "instrumented total", positive=True) for t in total]
    if not totals:
        raise ValueError("Missing instrumented total")
    result: dict[str, list[float]] = {}
    for key in COMPONENTS:
        if key not in components or len(components[key]) != len(totals):
            raise ValueError(f"Component {key} samples do not align with same-invocation total")
        result[key] = [finite(v, key) for v in components[key]]
    residual = [t - math.fsum(result[key][i] for key in COMPONENTS)
                for i, t in enumerate(totals)]
    if min(residual) < -tolerance_ms:
        raise ValueError(f"Component time exceeds same-invocation total: {min(residual)} ms")
    result["other"] = [max(0.0, v) for v in residual]
    # A tiny negative residual is carried independently in the audit, never
    # hidden as a negative 'Other' runtime component.
    result["closure_roundoff"] = [min(0.0, v) for v in residual]
    result["total"] = totals
    return result


def additive_summary(rows: list[dict[str, Any]]) -> dict[str, Any]:
    """Equal-weight window/run means, not sum of medians or geometric means."""
    if not rows:
        raise ValueError("Missing component rows")
    keys = (*COMPONENTS, "other", "closure_roundoff", "total")
    out = {f"{key}_ms": statistics.fmean(r[f"{key}_ms"] for r in rows) for key in keys}
    total = finite(out["total_ms"], "additive total", positive=True)
    out.update({f"{key}_pct": 100 * out[f"{key}_ms"] / total for key in (*COMPONENTS, "other")})
    out["closure_error_ms"] = math.fsum(out[f"{key}_ms"] for key in (*COMPONENTS, "other", "closure_roundoff")) - total
    return out


def dump_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    if not rows:
        path.write_text("")
        return
    fields = list(dict.fromkeys(k for row in rows for k in row))
    with path.open("w", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        for row in rows:
            writer.writerow({k: json.dumps(v, ensure_ascii=False, sort_keys=True) if isinstance(v, (dict, list)) else v for k, v in row.items()})


def markdown_table(rows: list[dict[str, Any]], columns: list[tuple[str, str]],
                   digits: int = 5) -> str:
    def cell(v: Any) -> str:
        if v is None:
            return "—"
        if isinstance(v, float):
            return f"{v:.{digits}f}"
        return str(v).replace("|", "\\|")
    lines = ["| " + " | ".join(title for _, title in columns) + " |",
             "| " + " | ".join("---" for _ in columns) + " |"]
    lines += ["| " + " | ".join(cell(row.get(key)) for key, _ in columns) + " |" for row in rows]
    return "\n".join(lines)


def discover_json(path: Path) -> tuple[list[dict[str, Any]], dict[str, Any], list[Path]]:
    """Input adapter is deliberately small: benchmark JSONL plus metadata JSON."""
    candidates = [path] if path.is_file() else sorted(path.glob("*.jsonl"))
    records: list[dict[str, Any]] = []
    files: list[Path] = []
    metadata: dict[str, Any] = {}
    for source in candidates:
        if source.suffix == ".jsonl":
            payloads = [json.loads(line) for line in source.read_text().splitlines() if line.strip()]
        else:
            parsed = json.loads(source.read_text())
            payloads = parsed.get("records", parsed.get("results", [])) if isinstance(parsed, dict) else parsed
            if isinstance(parsed, dict):
                metadata.update(parsed.get("metadata", {}))
        for payload in payloads:
            if isinstance(payload, dict) and ("window_id" in payload or "id" in payload):
                records.append(payload)
        files.append(source)
    parent = path if path.is_dir() else path.parent
    for name in ("environment.json", "metadata.json", "run_metadata.json", "manifest.json", "run_manifest.json"):
        source = parent / name
        if source.exists():
            metadata[name] = json.loads(source.read_text())
            files.append(source)
    if not records:
        raise ValueError(f"No benchmark records found in {path}")
    return records, metadata, files


def load_campaigns(inputs: list[Path], routes: dict[str, Any], routes_path: Path,
                   *, allow_partial: bool = False) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    expected = {w["id"]: w for w in routes["windows"]}
    rows, campaigns, correctness_all = [], [], []
    if len(inputs) != 2 and not allow_partial:
        raise ValueError("Full report requires two independently completed campaign directories")
    for repeat, source in enumerate(inputs, start=1):
        folder = source if source.is_dir() else source.parent
        records, metadata, files = discover_json(folder / "measurements.jsonl")
        env = metadata.get("metadata.json", {})
        if env.get("routes_sha256") != sha256(routes_path):
            raise ValueError(f"Route file digest mismatch in {folder}")
        receipt_path, check_path = folder / "COMPLETE.json", folder / "correctness.json"
        receipt = json.loads(receipt_path.read_text()) if receipt_path.exists() else {}
        checks = json.loads(check_path.read_text()) if check_path.exists() else []
        if not allow_partial and receipt.get("status") != "PASS":
            raise ValueError(f"Campaign lacks a passing completion receipt: {folder}")
        seen, paths, modes = set(), set(), set()
        for record in records:
            wid, mode, path = record["window_id"], record["mode"], record["path"]
            key = (wid, mode, path)
            if key in seen:
                raise ValueError(f"Duplicate timing point {repeat}/{key}")
            seen.add(key)
            if wid not in expected or record["batch"] != expected[wid]["batch"]:
                raise ValueError(f"Unknown or altered trace identity: {wid}")
            if mode not in ("eager", "graph"):
                raise ValueError(f"Unknown execution mode: {mode}")
            if path not in ("vllm_grouped", "triton_grouped", "triton_split", "vllm_split"):
                raise ValueError(f"Unknown execution path: {path}")
            if record["warmup"] < 10 or record["samples"] < 50:
                raise ValueError("Warmup/sample minimum not met")
            pop = population_metrics(expected[wid])
            for key2 in ("hot_experts", "cold_experts", "hot_token_assignments", "cold_token_assignments"):
                if record[key2] != pop[key2]:
                    raise ValueError(f"Population changed: {wid}/{key2}")
            for key2 in ("e2e_ms", "instrumented_total_ms"):
                if len(record[key2]) != record["samples"]:
                    raise ValueError(f"Wrong sample count: {wid}/{key2}")
                for val in record[key2]:
                    finite(val, key2, positive=True)
            for correctness_flag in ("raw_output_correct", "instrumented_output_correct"):
                if record.get(correctness_flag) is not True:
                    raise ValueError(f"Missing/failed measured-path numerical check: {wid}/{correctness_flag}")
            stages = record["stage_samples_ms"]
            for stage, values in stages.items():
                if len(values) != record["samples"]:
                    raise ValueError(f"Wrong component count: {wid}/{stage}")
                sample_statistics(values)
            if path.endswith("split"):
                for category in ("hot", "cold"):
                    if record[f"{category}_experts"] == 0 and any(stages[category]):
                        raise ValueError(f"Empty category carries event overhead: {wid}/{category}")
            paths.add(path)
            modes.add(mode)
            rows.append({**record, "repeat": repeat, "gpu": env.get("gpu", "unknown")})
        if not allow_partial:
            if not {"vllm_grouped", "triton_grouped", "triton_split"} <= paths or modes != {"eager", "graph"}:
                raise ValueError(f"Required execution paths/modes missing: {folder}")
            required = {(wid, mode, path) for wid in expected for mode in modes for path in paths}
            if seen != required:
                raise ValueError(f"Incomplete campaign coverage: missing {len(required-seen)}, unexpected {len(seen-required)}")
            if receipt.get("points") != len(records) or receipt.get("windows") != len(expected):
                raise ValueError("Completion receipt disagrees with raw timing coverage")
        check_seen = set()
        for check in checks:
            key = (check["window_id"], check["path"])
            if key in check_seen or check.get("passed") is not True:
                raise ValueError("Duplicate or failed numerical validation")
            check_seen.add(key)
            finite(check["max_abs_error"], "absolute error")
            finite(check["relative_l2_error"], "relative L2 error")
            correctness_all.append({**check, "repeat": repeat})
        if not allow_partial and check_seen != {(wid, path) for wid in expected for path in paths}:
            raise ValueError("Numerical checks do not cover every frozen window and execution path")
        files.extend(p for p in (receipt_path, check_path) if p.exists())
        campaigns.append({"repeat": repeat, "directory": str(folder), "metadata": env,
                          "receipt": receipt, "paths": sorted(paths), "modes": sorted(modes),
                          "source_sha256": {str(p): sha256(p) for p in dict.fromkeys(files)}})
    if not allow_partial and (len({tuple(c["paths"]) for c in campaigns}) != 1 or
                              len({c["metadata"].get("gpu") for c in campaigns}) != 1 or
                              len({json.dumps(c["metadata"].get("software"), sort_keys=True) for c in campaigns}) != 1):
        raise ValueError("Repeats use inconsistent paths/GPU models/software")
    audit = {"schema": SCHEMA, "status": "PARTIAL" if allow_partial else "PASS",
             "campaigns": campaigns, "timing_points": len(rows), "route_sha256": sha256(routes_path),
             "expected_windows": len(expected), "batch_window_counts": EXPECTED_COUNTS,
             "numeric_checks": correctness_all,
             "measured_path_checks": len(rows) * 2,
             "warmup_counts": sorted({r["warmup"] for r in rows}),
             "sample_counts": sorted({r["samples"] for r in rows}), "claim": "Real GPU trace-driven post-router layer replay; synthetic BF16 tensors"}
    return rows, audit


def aggregate(rows: list[dict[str, Any]]) -> dict[str, list[dict[str, Any]]]:
    point_rows, split_points, grouped_points = [], [], []
    by_group = defaultdict(list)
    for row in rows:
        key = {k: row[k] for k in ("window_id", "batch", "mode", "path", "repeat", "gpu")}
        plain = sample_statistics(row["e2e_ms"])
        instr = sample_statistics(row["instrumented_total_ms"])
        point = {**key, **plain, "instrumented_mean_ms": instr["mean_ms"],
                 "instrumentation_delta_pct": 100 * (instr["mean_ms"] / plain["mean_ms"] - 1)}
        point_rows.append(point)
        by_group[(row["mode"], row["path"], row["batch"])].append(row)
        stages = row["stage_samples_ms"]
        if row["path"].endswith("split"):
            # Dispatch is included in the category stage by the benchmark, not
            # independently measured. Internal zero is solely for accounting.
            complete = {**stages, "dispatch": stages.get("dispatch", [0.] * len(row["e2e_ms"]))}
            closed = component_accounting(row["instrumented_total_ms"], complete)
            sp = {**key, **{f"{k}_ms": statistics.fmean(v) for k, v in closed.items()},
                  "dispatch_scope": "included in Hot/Cold" if "dispatch" not in stages else "separate",
                  "uninstrumented_mean_ms": plain["mean_ms"]}
            split_points.append(sp)
        else:
            if set(stages) != {"shared", "routed", "combine"}:
                raise ValueError("Unexpected grouped stage schema")
            for i, total in enumerate(row["instrumented_total_ms"]):
                if abs(math.fsum(v[i] for v in stages.values()) - total) > .00002:
                    raise ValueError("Grouped same-invocation timing does not close")
            grouped_points.append({**key, **{f"{k}_ms": statistics.fmean(v) for k, v in stages.items()},
                                   "total_ms": instr["mean_ms"]})
    main_rows, overhead_rows, stage_statistics = [], [], []
    for (mode, path, batch), group in sorted(by_group.items()):
        means = [statistics.fmean(r["e2e_ms"]) for r in group]
        samples = [v for r in group for v in r["e2e_ms"]]
        inst = statistics.fmean(statistics.fmean(r["instrumented_total_ms"]) for r in group)
        stats = sample_statistics(samples)
        out = {"batch": batch, "mode": mode, "path": path, "gpu": group[0]["gpu"],
               "windows": len({r["window_id"] for r in group}), "repeats": len({r["repeat"] for r in group}),
               "mean_ms": statistics.fmean(means), "pooled_median_ms": stats["median_ms"],
               "pooled_stdev_ms": stats["stdev_ms"], "samples": len(samples)}
        main_rows.append(out)
        for stage in group[0]["stage_samples_ms"]:
            values = [v for r in group for v in r["stage_samples_ms"][stage]]
            st = sample_statistics(values)
            stage_statistics.append({"batch": batch, "mode": mode, "path": path, "stage": stage,
                                     **st, "scope": "instrumented same-invocation; pooled samples including window variation"})
        overhead_rows.append({"batch": batch, "mode": mode, "path": path,
                              "plain_mean_ms": out["mean_ms"], "instrumented_mean_ms": inst,
                              "instrumentation_delta_pct": (inst / out["mean_ms"] - 1) * 100})
    for mode, path in sorted({(r["mode"], r["path"]) for r in main_rows}):
        group = [r for r in main_rows if r["mode"] == mode and r["path"] == path]
        if {r["batch"] for r in group} == set(BATCHES):
            main_rows.append({"batch": "7-Batch GM", "mode": mode, "path": path, "gpu": group[0]["gpu"],
                              "windows": sum(r["windows"] for r in group), "repeats": group[0]["repeats"],
                              "mean_ms": geometric_mean(r["mean_ms"] for r in group),
                              "pooled_median_ms": None, "pooled_stdev_ms": None, "samples": sum(r["samples"] for r in group)})
    split_rows = []
    for mode, path in sorted({(r["mode"], r["path"]) for r in split_points}):
        for batch in BATCHES:
            group = [r for r in split_points if (r["mode"], r["path"], r["batch"]) == (mode, path, batch)]
            if not group:
                continue
            result = additive_summary(group)
            split_rows.append({"batch": batch, "mode": mode, "path": path, "gpu": group[0]["gpu"], **result,
                               "dispatch_scope": group[0]["dispatch_scope"],
                               "uninstrumented_mean_ms": statistics.fmean(r["uninstrumented_mean_ms"] for r in group),
                               "bottleneck": "未由硬件计数器判定"})
        group = [r for r in split_rows if r["mode"] == mode and r["path"] == path]
        if {r["batch"] for r in group} == set(BATCHES):
            result = additive_summary(group)
            # AM of Batch shares is requested separately from ratio of AM times.
            result.update({f"{k}_pct": statistics.fmean(r[f"{k}_pct"] for r in group)
                           for k in (*COMPONENTS, "other")})
            split_rows.append({"batch": "7-Batch 等权 AM", "mode": mode, "path": path, "gpu": group[0]["gpu"],
                               **result, "dispatch_scope": group[0]["dispatch_scope"],
                               "uninstrumented_mean_ms": statistics.fmean(r["uninstrumented_mean_ms"] for r in group),
                               "bottleneck": "时间和占比分别等权 AM；占比非 AM 时间之比"})
    ratios = []
    by_point = {(r["window_id"], r["mode"], r["repeat"], r["path"]): r for r in point_rows}
    pairs = [("triton_split", "triton_grouped"), ("triton_grouped", "vllm_grouped"), ("vllm_split", "vllm_grouped")]
    for numerator, denominator in pairs:
        for mode in ("eager", "graph"):
            grouped = defaultdict(list)
            for key, num in by_point.items():
                if num["path"] != numerator or num["mode"] != mode:
                    continue
                den = by_point.get((num["window_id"], mode, num["repeat"], denominator))
                if den:
                    grouped[num["batch"]].append(num["mean_ms"] / den["mean_ms"])
            for batch, values in sorted(grouped.items()):
                ratios.append({"batch": batch, "mode": mode, "numerator": numerator, "denominator": denominator,
                               "paired_ratio_gm": geometric_mean(values), "paired_ratio_am": statistics.fmean(values),
                               "meaning": "latency ratio; >1 means numerator slower"})
            if set(grouped) == set(BATCHES):
                ratios.append({"batch": "7-Batch GM", "mode": mode, "numerator": numerator, "denominator": denominator,
                               "paired_ratio_gm": geometric_mean(geometric_mean(v) for v in grouped.values()),
                               "paired_ratio_am": None, "meaning": "equal-Batch GM of paired latency ratios"})
    repeat_rows = []
    for mode, path, batch in sorted(by_group):
        points = [r for r in point_rows if (r["mode"], r["path"], r["batch"]) == (mode, path, batch)]
        windows = defaultdict(dict)
        for r in points:
            windows[r["window_id"]][r["repeat"]] = r["mean_ms"]
        pairs2 = [v for v in windows.values() if set(v) == {1, 2}]
        if pairs2:
            repeat_rows.append({"batch": batch, "mode": mode, "path": path,
                                "repeat2_over_repeat1_gm": geometric_mean(v[2] / v[1] for v in pairs2),
                                "mean_absolute_repeat_delta_pct": statistics.fmean(abs(v[2]/v[1]-1)*100 for v in pairs2),
                                "worst_absolute_repeat_delta_pct": max(abs(v[2]/v[1]-1)*100 for v in pairs2)})
    return {"point_summary": point_rows, "main_latency": main_rows, "split_breakdown": split_rows,
            "split_points": split_points, "grouped_points": grouped_points,
            "instrumentation": overhead_rows, "controlled_ratios": ratios, "repeat_stability": repeat_rows,
            "component_statistics": stage_statistics}


def workload_rows(routes: dict[str, Any]) -> list[dict[str, Any]]:
    output = []
    for batch in BATCHES:
        metrics = [population_metrics(w) for w in routes["windows"] if w["batch"] == batch]
        if not metrics:
            continue
        row = {"batch": batch, "windows": len(metrics)}
        for key in ("hot_experts", "cold_experts", "inactive_experts", "hot_token_assignments", "cold_token_assignments"):
            row[key] = statistics.fmean(m[key] for m in metrics)
        for category in ("shared", "hot", "cold"):
            row[f"{category}_flops_pct"] = statistics.fmean(100 * m[f"{category}_gemm_flops"] / m["total_gemm_flops"] for m in metrics)
        row["unique_weights_MiB"] = 33. + (row["hot_experts"] + row["cold_experts"]) * 16.5
        output.append(row)
    return output


def load_counter_tables(directory: Path) -> tuple[dict[str, list[dict[str, Any]]], dict[str, Any]]:
    """Include independently produced counter CSVs without rewriting them."""
    tables, provenance = {}, {}
    names = ("counter_summary", "routed_gemm_summary", "kineto_launch_summary", "counter_kernels")
    for name in names:
        source = directory / f"{name}.csv"
        if not source.exists():
            continue
        with source.open(newline="") as stream:
            reader = csv.DictReader(stream)
            if not reader.fieldnames or len(set(reader.fieldnames)) != len(reader.fieldnames):
                raise ValueError(f"Invalid counter table header: {source}")
            rows = list(reader)
        def typed(value):
            if value is None or value == "":
                return None
            if value in ("True", "False"):
                return value == "True"
            try:
                return int(value)
            except (ValueError, TypeError):
                try:
                    result = float(value)
                    return result if math.isfinite(result) else value
                except (ValueError, TypeError):
                    return value
        tables[name] = [{key: typed(value) for key, value in row.items()} for row in rows]
        provenance[name] = {"file": source.name, "sha256": sha256(source), "rows": len(rows),
                            "scope": "independent profiler replay; not primary elapsed-time samples"}
    return tables, provenance


def export_excel(path: Path, tables: dict[str, list[dict[str, Any]]]) -> bool:
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        return False
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for name, rows in tables.items():
        sheet = workbook.create_sheet(name[:31])
        if not rows:
            sheet.append(["No observations"])
            continue
        fields = list(dict.fromkeys(k for row in rows for k in row))
        sheet.append(fields)
        for row in rows:
            values = []
            for field in fields:
                v = row.get(field)
                if isinstance(v, (dict, list)):
                    v = json.dumps(v, ensure_ascii=False, sort_keys=True)
                if isinstance(v, str):
                    if len(v) > 32767:
                        v = v[:32700] + " [truncated in Excel; complete JSON retained]"
                    if v.startswith(("=", "+", "-", "@")):
                        v = "'" + v
                values.append(v)
            sheet.append(values)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="234E70")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for col in sheet.columns:
            width = min(48, max(14, max(len(str(c.value or "")) for c in list(col)[:50]) + 2))
            sheet.column_dimensions[col[0].column_letter].width = width
    workbook.save(path)
    check = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if len(check.sheetnames) != len(tables):
        raise AssertionError("Excel readback sheet count mismatch")
    check.close()
    return True


def make_report(tables: dict[str, list[dict[str, Any]]], audit: dict[str, Any],
                routes: dict[str, Any], *, excel: bool) -> str:
    env = audit["campaigns"][0]["metadata"]
    gpu = env.get("gpu", "unknown")
    soft = env.get("software", {})
    coverage = ("135 个冻结留出窗口，B2/4/8/16/64/96/128；两轮独立运行，Eager 与 CUDA Graph 分别测试。"
                if audit["status"] == "PASS" else "仅为部分点集调试报告，不是完整实验交付。")
    lines = ["# DeepSeek-V2-Lite：真实 GPU Shared / Hot / Cold 补测", "",
             f"状态：**{audit['status']}**。GPU：**{gpu}**。{coverage}",
             "", "本报告是**真实路由驱动的 post-router 单层微基准**：保留捕获的逐 Token Top-6 专家 ID 和路由权重，数值权重与激活使用固定种子的 BF16 随机张量。不是重新运行预训练模型，也不包含 Router GEMM/Top-K。",
             "", f"软件：PyTorch {soft.get('torch')}；Triton {soft.get('triton')}；vLLM {soft.get('vllm')}；CUDA {env.get('cuda_runtime')}。",
             "", "正常路径调用 vLLM 的 fused_experts，计入其 GPU 专家对齐、专家 GEMM、SiLU 与加权规约。Triton 路径是独立的 trace 专用原型，专家分组/块映射在 CPU 提前完成，GPU gather、GEMM、激活及合并计时。它不代表完整 vLLM 替代实现。",
             "", "**正常融合路径无法直接拆成互斥 Hot/Cold 时间。** 冷热表来自改变执行顺序后的受控诊断；同时报告拆分代价，不能把这些占比套回未拆分的生产路径。Shared/Hot/Cold 都按完整 FFN（含 SiLU）计时，Hot/Cold 包含相应分发，Combine 单列。",
             "", "## 1. 真实路由负载（计算量不是耗时）", "",
             markdown_table(tables["workload"], [("batch", "Batch"), ("windows", "窗口"), ("hot_experts", "Hot 数"),
                             ("cold_experts", "Cold 数"), ("shared_flops_pct", "Shared FLOPs %"), ("hot_flops_pct", "Hot FLOPs %"),
                             ("cold_flops_pct", "Cold FLOPs %"), ("unique_weights_MiB", "唯一权重 MiB")], 3),
             "", "Cold 指 1≤Me≤2，Me＝0 不计入 Cold；Hot 指 Me>2。Shared 中间宽度 2816，Routed 宽度 1408；每 Token 的六次 Routed 计算全部保留。Shared 有效 GEMM FLOPs 固定为总量的 25%，不预设其时间占比。",
             "", "## 2. 正常单层端到端延迟：主结果", "",
             f"下表无内部计时事件，使用 CUDA Events 包围整次调用。实际每点 warmup 次数为 {audit.get('warmup_counts', [])}，采样次数为 {audit.get('sample_counts', [])}；每 Batch 的 Mean 是窗口/轮次等权算术平均。Median、Std 为所有同 Batch 原始样本的合并统计，包含窗口差异与运行抖动，不是 Mean 的置信区间。7-Batch GM 只对七个 Batch Mean 取几何平均，不混合不同 Batch 的占比。"]
    for mode in ("graph", "eager"):
        for path in ("vllm_grouped", "triton_grouped"):
            selected = [r for r in tables["main_latency"] if r["mode"] == mode and r["path"] == path]
            if not selected:
                continue
            lines += ["", f"**{mode} / {path}**", "", markdown_table(selected,
                       [("batch", "Batch"), ("windows", "窗口"), ("mean_ms", "Mean ms"),
                        ("pooled_median_ms", "Median ms"), ("pooled_stdev_ms", "Std ms")])]
    lines += ["", "CUDA Events 记录 GPU 时间轴上的完整区间。Eager 模式可包含 GPU 等待主机提交造成的间隙；它不是各 kernel 执行时间之和。",
              "", "## 3. Shared / Hot / Cold 可加拆分：受控诊断", "",
              "各阶段事件边界位于同一次完整调用中，禁止把独立热缓存微测结果相加。表中 Total 为加内部事件后的同次调用时间；没有任务的类别严格记零。分发含在 Hot/Cold，不能重复加一份分发时间。阶段均值可加，阶段中位数不可加。",
              "", "末行“7-Batch 等权 AM”的时间、百分比分别取七 Batch 等权算术平均；该行的平均百分比不等于平均时间之比，且不是 GM。"]
    for mode in ("graph", "eager"):
        for path in ("vllm_split", "triton_split"):
            selected = [r for r in tables["split_breakdown"] if r["mode"] == mode and r["path"] == path]
            if not selected:
                continue
            lines += ["", f"**{mode} / {path}（条件拆分，不是正常路径归因）**", "",
                      markdown_table(selected, [("batch", "Batch"), ("shared_ms", "Shared ms"), ("hot_ms", "Hot ms"),
                          ("cold_ms", "Cold ms"), ("combine_ms", "合并 ms"), ("total_ms", "Total ms"),
                          ("shared_pct", "Shared %"), ("hot_pct", "Hot %"), ("cold_pct", "Cold %"), ("combine_pct", "合并 %")], 5)]
    lines += ["", "vLLM 冷热拆分分别调用两次带专家映射的 fused_experts；专家对齐与部分公共处理会重复。Triton 拆分也改变 kernel 数量与块排布。两者都是合法数值实现，但不保留原融合路径性能；需要结合下面的拆分倍率理解。",
              "", "## 4. 拆分代价与计时扰动", "",
              "倍率严格使用没有内部事件的完整调用时间，逐窗口/轮次配对后取 GM。大于 1 表示分子更慢。Triton/vLLM 的倍率仅比较当前实现，不证明某算法通用优劣。",
              "", markdown_table(tables["controlled_ratios"], [("batch", "Batch"), ("mode", "模式"),
                  ("numerator", "分子"), ("denominator", "分母"), ("paired_ratio_gm", "配对倍率 GM")], 4),
              "", "加入内部事件会改变图调度及调用节奏，扰动可能为正也可能为负，不能把差值认作精确计时开销。下表比较同一实现的两种完整调用。",
              "", markdown_table(tables["instrumentation"], [("batch", "Batch"), ("mode", "模式"), ("path", "路径"),
                  ("plain_mean_ms", "无内部事件 ms"), ("instrumented_mean_ms", "带内部事件 ms"),
                  ("instrumentation_delta_pct", "差异 %")], 5),
              "", "## 5. 瓶颈与结论边界", "",
              "本表不凭耗时占比填 Mem/Compute。Shared/Hot/Cold 是工作类型；显存带宽、Tensor Core 饱和程度是资源状态。Hot 耗时占比高不能推出 Compute-bound，Cold 多也不能推出显存带宽饱和。Eager 与 Graph 的差异仅表明对提交/调度方式敏感。",
              "", "硬件计数器必须来自单独的 Nsight/CUPTI 实测；PyTorch Profiler 的 kernel 时间不能替代 Tensor Core 占空比或 DRAM 带宽利用率。若同目录另有计数器报告，其 replay 时间也不能直接拼接成本报告端到端延迟。当前此聚合器不自动猜测 crossover。"]
    selected = [r for r in tables["split_breakdown"] if r["path"] == "vllm_split" and r["mode"] == "graph" and isinstance(r["batch"], int)]
    if selected:
        lo, hi = min(r["shared_pct"] for r in selected), max(r["shared_pct"] for r in selected)
        b2 = next(r for r in selected if r["batch"] == 2)
        lines += ["", f"**300 字内结果摘要：** 在 vLLM 的 CUDA Graph 冷热拆分诊断中，Shared 时间占比范围为 {lo:.2f}%–{hi:.2f}%；B2 的 Cold 占比为 {b2['cold_pct']:.2f}%。这些是改变融合计划后的实测比例，不能等同正常 Grouped 路径内部归因，也不能把 25% FLOPs 当成 25% 时间。正常路径耗时、拆分倍率与计时扰动已分别报告；是否访存/计算受限及其翻转 Batch，需要独立硬件计数器支撑，单靠此时间表不能确定。"]
    checks = audit["numeric_checks"]
    lines += ["", "## 6. 覆盖、校验与可复现性", "",
              f"- 原始计时点：{audit['timing_points']}；独立参考数值检查：{len(checks)}；无内部事件/带内部事件执行结果检查：{audit.get('measured_path_checks', 0)}；完整覆盖状态：{audit['status']}。",
              f"- 输入文件 SHA-256：`{audit['route_sha256']}`。",
              f"- 数值误差相对独立 BF16 参考逐窗口检查；最大绝对误差 {max((c['max_abs_error'] for c in checks), default=0):.8g}，最大相对 L2 误差 {max((c['relative_l2_error'] for c in checks), default=0):.8g}。完整容差与误差见 correctness.csv。",
              "- GPU 计时不会逐位重复；两轮差异逐 Batch/路径完整保留在 repeat_stability.csv，未以“完全相同”作为实测判据。",
              "- 窗口身份、Top-6、原始捕获分数、来源哈希与 6B 守恒均核验；大 Batch 为已有 captured_mixed 窗口，保留其来源，不伪称新生成的原生大 Batch 前向。",
              "- 测试集曾在此前 DSE 中查看；本次 GPU 数据是补测，不能把此集合描述为从未查看的新泛化测试集。",
              "", "[主延迟 CSV](main_latency.csv) · [冷热拆分 CSV](split_breakdown.csv) · [全部点摘要](point_summary.csv) · [两轮稳定性](repeat_stability.csv) · [审计与来源](REPORT_AUDIT.json)"]
    if excel:
        lines += ["", "[全部 Excel 数据表](GPU_HotCold_完整结果.xlsx)"]
    return "\n".join(lines) + "\n"


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, nargs="+", required=True, help="Two completed campaign directories; no smoke-run auto-discovery")
    parser.add_argument("--traces", type=Path, required=True)
    parser.add_argument("--out", type=Path, required=True)
    parser.add_argument("--allow-partial", action="store_true", help="For debugging only; marks report PARTIAL")
    args = parser.parse_args()
    routes = load_routes(args.traces)
    rows, audit = load_campaigns(args.input, routes, args.traces, allow_partial=args.allow_partial)
    tables = aggregate(rows)
    tables["workload"] = workload_rows(routes)
    tables["correctness"] = audit["numeric_checks"]
    tables["environment"] = [{"repeat": c["repeat"], "key": k, "value": v}
                             for c in audit["campaigns"] for k, v in c["metadata"].items()]
    args.out.mkdir(parents=True, exist_ok=True)
    for name, data in tables.items():
        dump_csv(args.out / f"{name}.csv", data)
    counter_tables, counter_provenance = load_counter_tables(args.out)
    excel = export_excel(args.out / "GPU_HotCold_完整结果.xlsx", {**tables, **counter_tables})
    audit["excel_counter_csvs"] = counter_provenance
    audit["excel_created"] = excel
    audit["report_source_sha256"] = sha256(Path(__file__))
    audit["aggregation"] = {"main_latency": "equal window/run AM; GM is equal GM of seven Batch means",
                            "median_stdev": "pooled raw samples, including window variation",
                            "split": "paired same-invocation stage means; dispatch included in category",
                            "split_summary": "equal Batch AM of times and separately equal Batch AM of shares",
                            "ratios": "within-window/run paired latency ratios; Batch GM then equal seven Batch GM"}
    text = make_report(tables, audit, routes, excel=excel)
    counter_candidates = [args.out / "GPU_COUNTERS.md", args.out.parent / "GPU_COUNTERS.md",
                          args.out.parent / "counters" / "GPU_COUNTERS.md"]
    for counter in counter_candidates:
        if counter.exists():
            import os
            relative = os.path.relpath(counter, args.out)
            text += f"\n[独立实测硬件计数器与启动行为报告]({relative})。计数器只覆盖所选代表窗口，不代表全部 135 窗口。\n"
            audit["counter_report"] = {"relative_path": relative, "sha256": sha256(counter)}
            break
    (args.out / "REPORT.md").write_text(text)
    (args.out / "REPORT_AUDIT.json").write_text(json.dumps(audit, indent=2, sort_keys=True, ensure_ascii=False, allow_nan=False) + "\n")
    print(json.dumps({"status": audit["status"], "points": len(rows), "output": str(args.out),
                      "excel": excel, "files": [p.name for p in sorted(args.out.iterdir())]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
